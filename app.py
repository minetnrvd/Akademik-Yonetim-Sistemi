import os
import datetime
import secrets
import socket
import hmac
import time
import re
import logging
import smtplib
import qrcode
import csv
import json
from collections import deque, defaultdict
from io import StringIO, BytesIO
from functools import wraps
from email.message import EmailMessage
from urllib.parse import urlparse, quote
from urllib.request import Request, urlopen
from logging.handlers import RotatingFileHandler
from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, Response, g, abort, send_from_directory
from sqlalchemy import text, func
from env_config import load_dotenv, resolve_database_url
from models import db, User, Student, StudentGroup, Message, Class, AttendanceSession, AttendanceRecord, Course, CourseEnrollment, GradeRecord, CourseRegistrationPolicy, CourseEnrollmentAudit, TermPlan, DEFAULT_COURSE_CREDIT, DEFAULT_COURSE_ECTS, Announcement, StudentCalendarEvent, UserCalendarEvent, PermissionAuditLog, AdminOperationLog


LANGUAGES = ('en', 'tr')
DEFAULT_LANGUAGE = 'en'
MAX_PERMISSION_AUDIT_EVENTS = 500
PERMISSION_AUDIT_EVENTS = deque(maxlen=MAX_PERMISSION_AUDIT_EVENTS)
RATE_LIMIT_EVENTS = {}
REQUEST_METRICS = {
    'since_utc': datetime.datetime.now(datetime.UTC).isoformat().replace('+00:00', 'Z'),
    'total_requests': 0,
    'error_4xx': 0,
    'error_5xx': 0,
    'by_endpoint': {},
}

TRANSLATIONS = {
    'en': {
        'language': 'Language',
        'site_title': 'QR Attendance System',
        'english': 'English',
        'turkish': 'Turkish',
        'language_updated': 'Language updated.',
        'skip_to_content': 'Skip to content',
        'login': 'Login',
        'register': 'Register',
        'email': 'Email',
        'password': 'Password',
        'remember_me': 'Remember Me',
        'dont_have_account': "Don't have an account?",
        'already_registered': 'Already registered?',
        'welcome_short': 'Welcome.',
        'student': 'Student',
        'teacher': 'Teacher',
        'name': 'Name',
        'student_number': 'Student Number',
        'student_home': 'Student Home',
        'welcome_user_continue': 'Welcome, {name}.',
        'my_account': 'My Account',
        'logout': 'Logout',
        'announcements': 'Announcements',
        'no_announcements': 'No announcements available.',
        'weather': 'Weather',
        'enter_city': 'Enter city',
        'update': 'Update',
        'use_current_location': 'Use My Current Location',
        'source': 'Source',
        'automatic_location': 'Automatic location',
        'manual_city': 'Manual city',
        'city': 'City',
        'condition': 'Condition',
        'temperature': 'Temperature',
        'humidity': 'Humidity',
        'weather_service_unavailable': 'Weather service is temporarily unavailable.',
        'calendar_and_upcoming': 'Calendar and Upcoming Items',
        'calendar_view': 'Calendar View',
        'previous_month': 'Previous',
        'next_month': 'Next',
        'select_day_prompt': 'Select a day from the calendar.',
        'selected_day': 'Selected Day',
        'no_events_this_day': 'No event for this day.',
        'today': 'Today',
        'events': 'Events',
        'upcoming_counts': 'Upcoming exams: {exams} | Upcoming activities: {activities}',
        'add_item': 'Add Item',
        'title': 'Title',
        'event_type': 'Event Type',
        'exam': 'Exam',
        'activity': 'Activity',
        'event_date': 'Event Date',
        'optional_note': 'Optional note',
        'add_to_calendar': 'Add to Calendar',
        'upcoming_exams_activities': 'Upcoming Exams and Activities',
        'date': 'Date',
        'delete': 'Delete',
        'no_upcoming_item': 'No upcoming item in your calendar.',
        'page_not_found': 'Page Not Found',
        'page_not_found_404': 'Page Not Found (404)',
        'requested_page_not_found': 'The page you requested could not be found.',
        'check_url_or_dashboard': 'Please check the URL or go back to your dashboard.',
        'teacher_dashboard': 'Teacher Dashboard',
        'go_to_login': 'Go to Login',
        'server_error': 'Server Error',
        'something_went_wrong_500': 'Something Went Wrong (500)',
        'unexpected_server_error': 'An unexpected server error occurred.',
        'try_again_moment': 'Please try again in a moment.',
        'attendance_successful': 'Attendance Successful',
        'attendance_recorded': 'Attendance Recorded',
        'class': 'Class',
        'week': 'Week',
        'your_attendance_confirmed': 'Your attendance has been successfully confirmed.',
        'return_to_student_dashboard': 'Return to Student Dashboard',
        'class_details': 'Class Details',
        'students': 'Students',
        'student_name': 'Student Name',
        'full_name': 'Full Name',
        'first_name': 'First Name',
        'last_name': 'Last Name',
        'gender': 'Gender',
        'birth_date': 'Birth Date',
        'birth_place': 'Birth Place',
        'citizenship': 'Citizenship',
        'identity_number': 'Identity Number',
        'nationality': 'Nationality',
        'registered_city': 'Registered City',
        'registered_district': 'Registered District',
        'passport_number': 'Passport Number',
        'active_box': 'Active Box',
        'passport_details': 'Passport Details',
        'issued_date': 'Issue Date',
        'issued_place': 'Issue Place',
        'valid_until': 'Validity Date',
        'active_status_box': 'Active',
        'other_information': 'Other Information',
        'marital_status': 'Marital Status',
        'blood_type': 'Blood Type',
        'veteran_martyr_relative_q': 'Veteran/Martyr Relative?',
        'disabled_q': 'Disabled?',
        'disability_type': 'Disability Type',
        'disability_rate': 'Disability Rate',
        'working_information': 'Working Information',
        'is_employed_q': 'Is Employed?',
        'is_group_company_q': 'Group Company?',
        'company_name': 'Company Name',
        'work_type': 'Work Type',
        'employment_start_date': 'Employment Start Date',
        'student_affairs_edit_note': 'If you want to make changes, please apply to Student Affairs.',
        'expand': 'Expand',
        'collapse': 'Collapse',
        'role_label': 'Role',
        'student_number_label': 'Student Number',
        'attended_sessions': 'Attended Sessions',
        'absences': 'Absences',
        'passed_q': 'Passed?',
        'yes': 'Yes',
        'no': 'No',
        'no_students_in_class': 'No students in this class yet.',
        'back': 'Back',
        'create_class': 'Create Class',


        'create_new_class': 'Create New Class',
        'class_name': 'Class Name',
        'create': 'Create',
        'join_class': 'Join Class',
        'search_class': 'Search class...',
        'join': 'Join',
        'absence_title': 'Absence',
        'absence_subtitle': 'All attendance operations for {name} are managed here.',
        'back_to_home': 'Back to Home',
        'enrolled_courses': 'Enrolled Courses',
        'active_attendance': 'Active Attendance',
        'total_sessions': 'Total Sessions',
        'average_absence': 'Average Absence',
        'search_course': 'Search course...',
        'search': 'Search',
        'course_based_absence': 'Course-Based Absence',
        'total': 'Total',
        'attended': 'Attended',
        'status': 'Status',
        'active': 'Active',
        'attendance_already_marked': 'Attendance already marked',
        'mark_attendance': 'Mark Attendance',
        'no_active_attendance': 'No active attendance',
        'absence_history': 'Absence History',
        'no_enrolled_course_found': 'No enrolled course found.',
        'active_attendance_alerts': 'Active Attendance Alerts',
        'ends': 'Ends',
        'no_active_attendance_alerts': 'No active attendance alerts.',
        'join_course': 'Join Course',
        'no_additional_course': 'No additional course available to join.',
        'my_attendance_history': 'My Attendance History',
        'present': 'Present',
        'absent': 'Absent',
        'no_confirmed_attendance': 'No confirmed attendance sessions for this class yet.',
        'account_settings': 'Account Settings',
        'change_password': 'Change Password',
        'current_password': 'Current Password',
        'new_password': 'New Password',
        'confirm_new_password': 'Confirm New Password',
        'account_deletion': 'Account Deletion',
        'student_delete_account_disabled': 'Account deletion is disabled for students. Your profile and attendance history are kept for record integrity.',
        'delete_account': 'Delete Account',
        'delete_account_warning': 'This action cannot be undone. Enter your password to delete your account.',
        'my_attendance_history_title': 'My Attendance History',
        'attendance_history': 'Attendance History',
        'back_to_dashboard': 'Back to Dashboard',
        'filters': 'Filters',
        'search_week': 'Search week',
        'start': 'Start',
        'end': 'End',
        'filter': 'Filter',
        'clear_filters': 'Clear filters',
        'download_csv': 'Download CSV',
        'summary': 'Summary',
        'detail': 'Detail',
        'no_recorded_attendance': 'No recorded attendance sessions yet.',
        'confirm_delete_attendance_session': 'Are you sure you want to delete this attendance session?',
        'attendance_detail': 'Attendance Detail',
        'absence_pct': 'Absence %',
        'total_classes': 'Total Classes',
        'absence_percent_label': 'Absence (%)',
        'no_students_found_class': 'No students found in this class yet.',
        'qr_code': 'QR Code',
        'token': 'Token',
        'attendance_link': 'Attendance Link',
        'qr_help_text': 'Make sure your phone is on the same network and this app is reachable via this URL.',
        'download': 'Download',
        'print_or_pdf': 'Print / Save as PDF',
        'welcome_teacher': 'Welcome, {name}',
        'start_attendance': 'Start Attendance',
        'week_placeholder': 'Week (e.g. 1)',
        'start_and_open_qr': 'Start and Open QR',
        'create_class_first': 'You need to create a class first.',
        'current_classes_and_qr': 'Current Classes and QR Codes',
        'students_count': 'Students: {count}',
        'session_active': 'Session Active',
        'no_active_session': 'No Active Session',
        'class_qr_code': 'Class QR Code',
        'fullscreen': 'Fullscreen',
        'qr_not_generated': 'QR code has not been generated.',
        'active_attendance_session': 'Active Attendance Session',
        'live_attendance': 'Live Attendance',
        'ends_at': 'Ends At',
        'finish_attendance': 'Finish Attendance',
        'no_active_session_note': 'There is no active attendance session yet. You can create one above using Start Attendance.',
        'you_have_not_created_class': 'You have not created any class yet.',
        'popup_blocked': 'New tab was blocked. Please allow pop-ups.',
        'pass': 'Pass',
        'fail': 'Fail',
        'academic_status': 'Academic Status',
        'active_value': 'Active',
        'program': 'Program',
        'university_undergraduate_program': 'University Undergraduate Program',
        'course_count': 'Course Count',
        'grading_system': 'Grading System',
        'letter_grade_400': 'Letter Grade (4.00)',
        'guardian_full_name': 'Guardian Full Name',
        'relationship': 'Relationship',
        'emergency_phone': 'Emergency Phone',
        'note': 'Note',
        'not_set_yet': 'Not set yet',
        'family_note': 'These fields can be edited via forms in the next step.',
        'document_status': 'Document Status',
        'registration_document': 'Registration Document',
        'student_certificate': 'Student Certificate',
        'documents_note': 'The documents module will be expanded on this screen.',
        'no_uploaded_document': 'No uploaded document yet',
        'pending': 'Pending',
        'can_be_requested': 'Can be requested',
        'phone': 'Phone',
        'address': 'Address',
        'contact_preference': 'Contact Preference',
        'total_debt': 'Total Debt',
        'total_paid': 'Total Paid',
        'current_balance': 'Current Balance',
        'current_account_note': 'The current account module will be detailed here.',
        'last_payment': 'Last Payment',
        'payment_method': 'Payment Method',
        'scheduled_payment': 'Scheduled Payment',
        'payment_note': 'The payment module can be expanded in the next step.',
        'no_record_found': 'No record found',
        'undefined': 'Undefined',
        'none': 'None',
        'current_term': 'Current Term',
        'registered_course_count': 'Registered Course Count',
        'registered_courses': 'Registered Courses',
        'status_label': 'Status',
        'no_registered_course': 'No registered course yet',
        'active_term': 'Active Term',
        'student_label': 'Student',
        'gpa_4': 'GPA (4.00)',
        'weighted_credits': 'Weighted ECTS',
        'no_graded_data': 'No graded course data yet.',
        'advisor_student_not_found': 'Student not found.',
        'advisor_must_be_teacher': 'The advisor must be a teacher.',
        'advisor_update_failed': 'Advisor update failed (database error).',
        'advisor_updated': 'Advisor updated.',
        'advisee_not_found': 'This student is not one of your advisees.',
        'view_as_read_only': 'Read-only student view: this action is not available.',
        'calculated_from_courses': 'Calculated from {count} graded course(s).',
        'transcript_table_not_initialized': 'Transcript table is not initialized yet.',
        'today': 'Today',
        'upcoming_academic_events': 'Upcoming Academic Events',
        'term_window': 'Term Window',
        'configured_by_admin': 'Configured by administration',
        'academic_calendar_note': 'Academic calendar events will be listed here in detail.',
        'upcoming_exams': 'Upcoming Exams',
        'completed_exams': 'Completed Exams',
        'next_exam_date': 'Next Exam Date',
        'not_scheduled': 'Not scheduled',
        'exam_module_note': 'Exam management module can be expanded in the next step.',
        'already_enrolled_named': "You are already enrolled in '{name}'.",
        'no_class_found_named': "No class found with the name '{name}'.",
        'identity_information': 'Identity Information',
        'identity_information_desc': 'Student identity and core profile details.',
        'education_information': 'Education Information',
        'education_information_desc': 'Department, program, term, and academic status.',
        'education_reset_note': 'This section has been cleared. We will rebuild it from scratch.',
        'university_information': 'University Information',
        'entry_place': 'Entry Place',
        'entry_type': 'Entry Type',
        'academic_year': 'Academic Year',
        'term': 'Term',
        'faculty': 'Faculty',
        'department': 'Department',
        'scholarship_type': 'Scholarship Type',
        'placement_type': 'Placement Type',
        'score_type': 'Score Type',
        'achievement_score': 'Achievement Score',
        'placement_score': 'Placement Score',
        'preference_order': 'Preference Order',
        'highschool_information': 'High School Information',
        'highschool_name': 'High School Name',
        'highschool_info': 'High School Info',
        'highschool_graduation_date': 'Graduation Date',
        'family_information': 'Family Information',
        'family_information_desc': 'Family contact and relative information.',
        'documents': 'Documents',
        'documents_desc': 'Document and certificate operations.',
        'contact': 'Contact',
        'contact_desc': 'Contact preferences and address details.',
        'current_account': 'Current Account',
        'current_account_desc': 'Balance, debt, and receivable summary.',
        'payments': 'Payments',
        'payments_desc': 'Payment history and new payment actions.',
        'absence': 'Absence',
        'absence_desc': 'Complete attendance workflow and absence analysis.',
        'term_courses': 'Term Courses',
        'term_courses_desc': 'Current term course list and load.',
        'transcript': 'Transcript',
        'transcript_desc': 'Grade summary and GPA overview.',
        'academic_calendar': 'Academic Calendar',
        'academic_calendar_desc': 'Important dates and academic events.',
        'exams': 'Exams',
        'exams_desc': 'Upcoming and completed exam information.',
        'modules': 'Modules',
    },
    'tr': {
        'language': 'Dil',
        'site_title': 'QR Yoklama Sistemi',
        'english': 'İngilizce',
        'turkish': 'Türkçe',
        'language_updated': 'Dil güncellendi.',
        'skip_to_content': 'İçeriğe geç',
        'login': 'Giriş',
        'register': 'Kayıt Ol',
        'email': 'E-posta',
        'password': 'Şifre',
        'remember_me': 'Beni Hatırla',
        'dont_have_account': 'Hesabın yok mu?',
        'already_registered': 'Zaten kayıtlı mısın?',
        'welcome_short': 'Hoş geldin.',
        'student': 'Öğrenci',
        'teacher': 'Öğretmen',
        'name': 'Ad Soyad',
        'student_number': 'Öğrenci Numarası',
        'student_home': 'Öğrenci Ana Sayfası',
        'welcome_user_continue': 'Hoş geldin, {name}.',
        'my_account': 'Hesabım',
        'logout': 'Çıkış',
        'announcements': 'Duyurular',
        'no_announcements': 'Henüz duyuru yok.',
        'weather': 'Hava Durumu',
        'enter_city': 'Şehir girin',
        'update': 'Güncelle',
        'use_current_location': 'Mevcut Konumumu Kullan',
        'source': 'Kaynak',
        'automatic_location': 'Otomatik konum',
        'manual_city': 'Manuel şehir',
        'city': 'Şehir',
        'condition': 'Durum',
        'temperature': 'Sıcaklık',
        'humidity': 'Nem',
        'weather_service_unavailable': 'Hava durumu servisi şu an kullanılamıyor.',
        'calendar_and_upcoming': 'Takvim ve Yaklaşan Öğeler',
        'calendar_view': 'Takvim Görünümü',
        'previous_month': 'Önceki',
        'next_month': 'Sonraki',
        'select_day_prompt': 'Takvimden bir gün seçin.',
        'selected_day': 'Seçilen Gün',
        'no_events_this_day': 'Bu gün için etkinlik yok.',
        'today': 'Bugün',
        'events': 'Etkinlikler',
        'upcoming_counts': 'Yaklaşan sınav: {exams} | Yaklaşan etkinlik: {activities}',
        'add_item': 'Öğe Ekle',
        'title': 'Başlık',
        'event_type': 'Etkinlik Türü',
        'exam': 'Sınav',
        'activity': 'Etkinlik',
        'event_date': 'Etkinlik Tarihi',
        'optional_note': 'İsteğe bağlı not',
        'add_to_calendar': 'Takvime Ekle',
        'upcoming_exams_activities': 'Yaklaşan Sınavlar ve Etkinlikler',
        'date': 'Tarih',
        'delete': 'Sil',
        'no_upcoming_item': 'Takviminde yaklaşan öğe yok.',
        'page_not_found': 'Sayfa Bulunamadı',
        'page_not_found_404': 'Sayfa Bulunamadı (404)',
        'requested_page_not_found': 'İstediğin sayfa bulunamadı.',
        'check_url_or_dashboard': 'Lütfen URL adresini kontrol et ya da paneline dön.',
        'teacher_dashboard': 'Öğretmen Paneli',
        'go_to_login': 'Giriş Sayfasına Git',
        'server_error': 'Sunucu Hatası',
        'something_went_wrong_500': 'Bir Şeyler Ters Gitti (500)',
        'unexpected_server_error': 'Beklenmeyen bir sunucu hatası oluştu.',
        'try_again_moment': 'Lütfen birazdan tekrar dene.',
        'attendance_successful': 'Yoklama Başarılı',
        'attendance_recorded': 'Yoklama Kaydedildi',
        'class': 'Ders',
        'week': 'Hafta',
        'your_attendance_confirmed': 'Yoklaman başarıyla onaylandı.',
        'return_to_student_dashboard': 'Öğrenci Paneline Dön',
        'class_details': 'Ders Detayları',
        'students': 'Öğrenciler',
        'student_name': 'Öğrenci Adı',
        'full_name': 'Ad Soyad',
        'first_name': 'Ad',
        'last_name': 'Soyad',
        'gender': 'Cinsiyet',
        'birth_date': 'Doğum Tarihi',
        'birth_place': 'Doğum Yeri',
        'citizenship': 'Vatandaşlık',
        'identity_number': 'Kimlik Numarası',
        'nationality': 'Uyruk',
        'registered_city': 'Kayıtlı Il',
        'registered_district': 'Kayıtlı İlçe',
        'passport_number': 'Pasaport No',
        'active': 'Aktif',
        'passport_details': 'Pasaport Detayları',
        'issued_date': 'Verilme Tarihi',
        'issued_place': 'Verilme Yeri',
        'valid_until': 'Geçerlilik Tarihi',
        'active_status_box': 'Aktif',
        'other_information': 'Diğer Bilgiler',
        'marital_status': 'Medeni Durum',
        'blood_type': 'Kan Grubu',
        'veteran_martyr_relative_q': 'Gazi/Şehit Yakını mi?',
        'disabled_q': 'Engelli mi?',
        'disability_type': 'Engel Türü',
        'disability_rate': 'Engel Oranı',
        'working_information': 'Çalışma Bilgileri',
        'is_employed_q': 'Çalışıyor mu?',
        'is_group_company_q': 'Grup Şirketi mi?',
        'company_name': 'Şirket Adı',
        'work_type': 'Çalışma Şekli',
        'employment_start_date': 'Ise Giriş Tarihi',
        'student_affairs_edit_note': 'Düzenleme Yapmak İstiyorsanız Lütfen Öğrenci İşlerine Başvurun',
        'expand': 'Aç',
        'collapse': 'Kapat',
        'role_label': 'Rol',
        'student_number_label': 'Öğrenci Numarası',
        'attended_sessions': 'Katıldığı Oturum',
        'absences': 'Devamsızlık',
        'passed_q': 'Geçti mi?',
        'yes': 'Evet',
        'no': 'Hayır',
        'no_students_in_class': 'Bu derste henüz öğrenci yok.',
        'back': 'Geri',
        'create_class': 'Ders Oluştur',
        'create_new_class': 'Yeni Ders Oluştur',
        'class_name': 'Ders Adı',
        'create': 'Oluştur',
        'join_class': 'Derse Katıl',
        'search_class': 'Ders ara...',
        'join': 'Katıl',
        'absence_title': 'Devamsızlık',
        'absence_subtitle': '{name} için tüm yoklama işlemleri burada yönetilir.',
        'back_to_home': 'Ana Sayfaya Dön',
        'enrolled_courses': 'Kayıtlı Dersler',
        'active_attendance': 'Aktif Yoklama',
        'total_sessions': 'Toplam Oturum',
        'average_absence': 'Ortalama Devamsızlık',
        'search_course': 'Ders ara...',
        'search': 'Ara',
        'course_based_absence': 'Ders Bazlı Devamsızlık',
        'total': 'Toplam',
        'attended': 'Katılım',
        'status': 'Durum',
        'attendance_already_marked': 'Yoklama zaten işaretli',
        'mark_attendance': 'Yoklama İşaretle',
        'no_active_attendance': 'Aktif yoklama yok',
        'absence_history': 'Devamsızlık Geçmişi',
        'no_enrolled_course_found': 'Kayıtlı ders bulunamadı.',
        'active_attendance_alerts': 'Aktif Yoklama Uyarıları',
        'ends': 'Biter',
        'no_active_attendance_alerts': 'Aktif yoklama uyarısı yok.',
        'join_course': 'Derse Katıl',
        'no_additional_course': 'Katılınabilecek ek ders yok.',
        'my_attendance_history': 'Yoklama Geçmişim',
        'present': 'Var',
        'absent': 'Yok',
        'no_confirmed_attendance': 'Bu ders için henüz onaylı yoklama oturumu yok.',
        'account_settings': 'Hesap Ayarları',
        'change_password': 'Şifre Değiştir',
        'current_password': 'Mevcut Şifre',
        'new_password': 'Yeni Şifre',
        'confirm_new_password': 'Yeni Şifreyi Onayla',
        'account_deletion': 'Hesap Silme',
        'student_delete_account_disabled': 'Öğrenciler için hesap silme kapalıdır. Profilin ve yoklama geçmişin kayıt bütünlüğü için saklanır.',
        'delete_account': 'Hesabı Sil',
        'delete_account_warning': 'Bu işlem geri alınamaz. Hesabını silmek için şifreni gir.',
        'my_attendance_history_title': 'Yoklama Geçmişim',
        'attendance_history': 'Yoklama Geçmişi',
        'back_to_dashboard': 'Panele Dön',
        'filters': 'Filtreler',
        'search_week': 'Hafta ara',
        'start': 'Başlangıç',
        'end': 'Bitiş',
        'filter': 'Filtrele',
        'clear_filters': 'Filtreleri temizle',
        'download_csv': 'CSV Indir',
        'summary': 'Özet',
        'detail': 'Detay',
        'no_recorded_attendance': 'Henüz kayıtlı yoklama oturumu yok.',
        'confirm_delete_attendance_session': 'Bu yoklama oturumunu silmek istediğine emin misin?',
        'attendance_detail': 'Yoklama Detayı',
        'absence_pct': 'Devamsızlık %',
        'total_classes': 'Toplam Ders',
        'absence_percent_label': 'Devamsızlık (%)',
        'no_students_found_class': 'Bu derste henüz öğrenci bulunamadı.',
        'qr_code': 'QR Kod',
        'token': 'Token',
        'attendance_link': 'Yoklama Bağlantısı',
        'qr_help_text': 'Telefonunun aynı ağda olduğundan ve uygulamanın bu URL üzerinden erişilebilir olduğundan emin ol.',
        'download': 'Indir',
        'print_or_pdf': 'Yazdır / PDF Kaydet',
        'welcome_teacher': 'Hoş geldin, {name}',
        'start_attendance': 'Yoklama Başlat',
        'week_placeholder': 'Hafta (örnek: 1)',
        'start_and_open_qr': 'Başlat ve QR Aç',
        'create_class_first': 'Önce bir ders oluşturman gerekiyor.',
        'current_classes_and_qr': 'Mevcut Dersler ve QR Kodları',
        'students_count': 'Öğrenci: {count}',
        'session_active': 'Oturum Aktif',
        'no_active_session': 'Aktif Oturum Yok',
        'class_qr_code': 'Ders QR Kodu',
        'fullscreen': 'Tam Ekran',
        'qr_not_generated': 'QR kodu henüz üretilmedi.',
        'active_attendance_session': 'Aktif Yoklama Oturumu',
        'live_attendance': 'Canlı Yoklama',
        'ends_at': 'Biter',
        'finish_attendance': 'Yoklamayı Bitir',
        'no_active_session_note': 'Henüz aktif yoklama oturumu yok. Yukarıdan Yoklama Başlat ile oluşturabilirsin.',
        'you_have_not_created_class': 'Henüz hiç ders oluşturmadın.',
        'popup_blocked': 'Yeni sekme engellendi. Lütfen açılır pencerelere izin ver.',
        'pass': 'Geçti',
        'fail': 'Kaldı',
        'academic_status': 'Akademik Durum',
        'active_value': 'Aktif',
        'program': 'Program',
        'university_undergraduate_program': 'Üniversite Lisans Programı',
        'course_count': 'Ders Sayısı',
        'grading_system': 'Notlandırma Sistemi',
        'letter_grade_400': 'Harf Notu (4.00)',
        'guardian_full_name': 'Veli Ad Soyad',
        'relationship': 'Yakınlık',
        'emergency_phone': 'Acil Durum Telefonu',
        'note': 'Not',
        'not_set_yet': 'Henüz ayarlanmadı',
        'family_note': 'Bu alanlar bir sonraki adımda formlar ile düzenlenebilir.',
        'document_status': 'Belge Durumu',
        'registration_document': 'Kayıt Belgesi',
        'student_certificate': 'Öğrenci Belgesi',
        'documents_note': 'Belgeler modülü bu ekranda genişletilecektir.',
        'no_uploaded_document': 'Henüz yüklenmiş belge yok',
        'pending': 'Beklemede',
        'can_be_requested': 'Talep edilebilir',
        'phone': 'Telefon',
        'address': 'Adres',
        'contact_preference': 'İletişim Tercihi',
        'total_debt': 'Toplam Borç',
        'total_paid': 'Toplam Ödenen',
        'current_balance': 'Güncel Bakiye',
        'current_account_note': 'Cari hesap modülü burada detaylandırılacaktır.',
        'last_payment': 'Son Ödeme',
        'payment_method': 'Ödeme Yöntemi',
        'scheduled_payment': 'Planlı Ödeme',
        'payment_note': 'Ödeme modülü bir sonraki adımda genişletilebilir.',
        'no_record_found': 'Kayıt bulunamadı',
        'undefined': 'Tanımsız',
        'none': 'Yok',
        'current_term': 'Güncel Dönem',
        'registered_course_count': 'Kayıtlı Ders Sayısı',
        'registered_courses': 'Kayıtlı Dersler',
        'status_label': 'Durum',
        'no_registered_course': 'Henüz kayıtlı ders yok',
        'active_term': 'Aktif Dönem',
        'student_label': 'Öğrenci',
        'gpa_4': 'GPA (4.00)',
        'weighted_credits': 'Ağırlıklı AKTS',
        'no_graded_data': 'Henüz notlandırılmış ders verisi yok.',
        'advisor_student_not_found': 'Öğrenci bulunamadı.',
        'advisor_must_be_teacher': 'Danışman bir öğretmen olmalı.',
        'advisor_update_failed': 'Danışman güncellenemedi (veritabanı hatası).',
        'advisor_updated': 'Danışman güncellendi.',
        'advisee_not_found': 'Bu öğrenci senin danışmanlık öğrencin değil.',
        'view_as_read_only': 'Salt okunur öğrenci görünümü: bu işlem yapılamaz.',
        'calculated_from_courses': '{count} notlandırılmış ders üzerinden hesaplandı.',
        'transcript_table_not_initialized': 'Transkript tablosu henüz hazır değil.',
        'today': 'Bugün',
        'upcoming_academic_events': 'Yaklaşan Akademik Etkinlikler',
        'term_window': 'Dönem Aralığı',
        'configured_by_admin': 'Yönetim tarafından tanımlanır',
        'academic_calendar_note': 'Akademik takvim etkinlikleri burada detaylı listelenecektir.',
        'upcoming_exams': 'Yaklaşan Sınavlar',
        'completed_exams': 'Tamamlanan Sınavlar',
        'next_exam_date': 'Sonraki Sınav Tarihi',
        'not_scheduled': 'Planlanmadı',
        'exam_module_note': 'Sınav yönetimi modülü bir sonraki adımda genişletilebilir.',
        'already_enrolled_named': "'{name}' dersine zaten kayıtlısın.",
        'no_class_found_named': "'{name}' adında bir ders bulunamadı.",
        'identity_information': 'Kimlik Bilgileri',
        'identity_information_desc': 'Öğrenci kimliği ve temel profil bilgileri.',
        'education_information': 'Eğitim Bilgileri',
        'education_information_desc': 'Bölüm, program, dönem ve akademik durum.',
        'education_reset_note': 'Bu bölüm temizlendi. Burayı sıfırdan birlikte oluşturacağız.',
        'university_information': 'Üniversite Bilgileri',
        'entry_place': 'Giriş Yeri',
        'entry_type': 'Giriş Tipi',
        'academic_year': 'Akademik Yıl',
        'term': 'Dönem',
        'faculty': 'Fakülte',
        'department': 'Bölüm',
        'scholarship_type': 'Burs Türü',
        'placement_type': 'Yerleşme Türü',
        'score_type': 'Puan Türü',
        'achievement_score': 'Başarı Puanı',
        'placement_score': 'Yerleşme Puanı',
        'preference_order': 'Tercih Sırası',
        'highschool_information': 'Lise Bilgileri',
        'highschool_name': 'Lise Adı',
        'highschool_info': 'Lise Bilgisi',
        'highschool_graduation_date': 'Mezuniyet Tarihi',
        'family_information': 'Aile Bilgileri',
        'family_information_desc': 'Aile iletişimi ve yakın bilgileri.',
        'documents': 'Belgeler',
        'documents_desc': 'Belge ve sertifika işlemleri.',
        'contact': 'İletişim',
        'contact_desc': 'İletişim tercihleri ve adres bilgileri.',
        'current_account': 'Cari Hesap',
        'current_account_desc': 'Bakiye, borç ve alacak özeti.',
        'payments': 'Ödemeler',
        'payments_desc': 'Ödeme geçmişi ve yeni ödeme işlemleri.',
        'absence': 'Devamsızlık',
        'absence_desc': 'Tam yoklama akışıyla devamsızlık analizi.',
        'term_courses': 'Dönem Dersleri',
        'term_courses_desc': 'Güncel dönem ders listesi ve yuk bilgisi.',
        'transcript': 'Transkript',
        'transcript_desc': 'Not özeti ve ortalama görünümü.',
        'academic_calendar': 'Akademik Takvim',
        'academic_calendar_desc': 'Önemli tarihler ve akademik etkinlikler.',
        'exams': 'Sınavlar',
        'exams_desc': 'Yaklaşan ve tamamlanan sınav bilgileri.',
        'modules': 'Modüller',
    },
}


def _detect_local_ip() -> str:
    """Best-effort local network IP detection for phone-accessible QR links."""
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        return s.getsockname()[0]
    except Exception:
        return "127.0.0.1"
    finally:
        s.close()


def _safe_int_env(name: str, default: int, minimum: int | None = None, maximum: int | None = None) -> int:
    raw = (os.getenv(name) or '').strip()
    value = default
    if raw:
        try:
            value = int(raw)
        except ValueError:
            value = default

    if minimum is not None:
        value = max(minimum, value)
    if maximum is not None:
        value = min(maximum, value)
    return value


def _configure_app_logging(flask_app: Flask):
    log_dir = (os.getenv('APP_LOG_DIR') or os.path.join(flask_app.root_path, 'logs')).strip()
    log_file = (os.getenv('APP_LOG_FILE') or 'app.log').strip() or 'app.log'
    max_bytes = _safe_int_env('APP_LOG_MAX_BYTES', 5 * 1024 * 1024, minimum=1024, maximum=200 * 1024 * 1024)
    backup_count = _safe_int_env('APP_LOG_BACKUP_COUNT', 7, minimum=1, maximum=90)
    level_name = (os.getenv('APP_LOG_LEVEL') or 'INFO').strip().upper()
    level = getattr(logging, level_name, logging.INFO)

    os.makedirs(log_dir, exist_ok=True)
    log_path = os.path.join(log_dir, log_file)
    formatter = logging.Formatter('%(asctime)s %(levelname)s %(name)s %(message)s')

    file_handler = RotatingFileHandler(
        log_path,
        maxBytes=max_bytes,
        backupCount=backup_count,
        encoding='utf-8',
    )
    file_handler.setLevel(level)
    file_handler.setFormatter(formatter)

    stream_handler = logging.StreamHandler()
    stream_handler.setLevel(level)
    stream_handler.setFormatter(formatter)

    flask_app.logger.handlers.clear()
    flask_app.logger.setLevel(level)
    flask_app.logger.addHandler(file_handler)
    flask_app.logger.addHandler(stream_handler)

    werkzeug_logger = logging.getLogger('werkzeug')
    werkzeug_logger.setLevel(level)
    if not any(isinstance(handler, RotatingFileHandler) for handler in werkzeug_logger.handlers):
        werkzeug_logger.addHandler(file_handler)


def _normalize_samesite(value: str | None) -> str:
    normalized = (value or 'Lax').strip().lower()
    if normalized == 'strict':
        return 'Strict'
    if normalized == 'none':
        return 'None'
    return 'Lax'


load_dotenv()

# ------------------ APP INIT ------------------
app = Flask(__name__)
# Stable fallback secret for local development so debug reloads do not invalidate sessions.
app.secret_key = os.getenv('FLASK_SECRET_KEY') or 'dev-local-secret-key-change-in-production'
app.config['SQLALCHEMY_DATABASE_URI'] = resolve_database_url()
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
# Base URL encoded in QR codes for phone scanning (example: http://192.168.1.25:5000)
app.config['QR_BASE_URL'] = os.getenv('QR_BASE_URL', f"http://{_detect_local_ip()}:5000").rstrip('/')
app.config['ATTENDANCE_WINDOW_MINUTES'] = _safe_int_env('ATTENDANCE_WINDOW_MINUTES', 180, minimum=1, maximum=24 * 60)
app.config['FORCE_CANONICAL_HOST'] = os.getenv('FORCE_CANONICAL_HOST', '0') == '1'
# Session/cookie security defaults (can be overridden by environment variables).
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_NAME'] = os.getenv('SESSION_COOKIE_NAME', 'qr_attendance_session')
app.config['SESSION_COOKIE_SAMESITE'] = _normalize_samesite(os.getenv('SESSION_COOKIE_SAMESITE', 'Lax'))
app.config['SESSION_COOKIE_SECURE'] = os.getenv('SESSION_COOKIE_SECURE', '0') == '1'
app.config['PERMANENT_SESSION_LIFETIME'] = datetime.timedelta(days=_safe_int_env('REMEMBER_ME_DAYS', 30, minimum=1, maximum=30))
app.config['SESSION_REFRESH_EACH_REQUEST'] = False
app.config['LOGIN_RATE_LIMIT_MAX_ATTEMPTS'] = _safe_int_env('LOGIN_RATE_LIMIT_MAX_ATTEMPTS', 8, minimum=1, maximum=100)
app.config['LOGIN_RATE_LIMIT_WINDOW_SECONDS'] = _safe_int_env('LOGIN_RATE_LIMIT_WINDOW_SECONDS', 300, minimum=10, maximum=3600)
app.config['LOGIN_LOCK_MAX_FAILURES'] = _safe_int_env('LOGIN_LOCK_MAX_FAILURES', 5, minimum=1, maximum=50)
app.config['LOGIN_LOCK_WINDOW_SECONDS'] = _safe_int_env('LOGIN_LOCK_WINDOW_SECONDS', 900, minimum=30, maximum=86400)
app.config['ADMIN_MUTATION_RATE_LIMIT_MAX'] = _safe_int_env('ADMIN_MUTATION_RATE_LIMIT_MAX', 30, minimum=1, maximum=200)
app.config['ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS'] = _safe_int_env('ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', 60, minimum=10, maximum=3600)
app.config['PASSWORD_MIN_LENGTH'] = _safe_int_env('PASSWORD_MIN_LENGTH', 10, minimum=8, maximum=128)
app.config['METRICS_SLOW_REQUEST_MS'] = _safe_int_env('METRICS_SLOW_REQUEST_MS', 750, minimum=50, maximum=60000)
app.config['HEALTH_WARN_ERROR_RATE_PCT'] = _safe_int_env('HEALTH_WARN_ERROR_RATE_PCT', 5, minimum=0, maximum=100)
app.config['HEALTH_WARN_P95_MS'] = _safe_int_env('HEALTH_WARN_P95_MS', 1000, minimum=50, maximum=120000)
app.config['EMAIL_DRY_RUN'] = os.getenv('EMAIL_DRY_RUN', '1') == '1'
app.config['EMAIL_SMTP_HOST'] = (os.getenv('EMAIL_SMTP_HOST') or '').strip()
app.config['EMAIL_SMTP_PORT'] = _safe_int_env('EMAIL_SMTP_PORT', 587, minimum=1, maximum=65535)
app.config['EMAIL_SMTP_USER'] = (os.getenv('EMAIL_SMTP_USER') or '').strip()
app.config['EMAIL_SMTP_PASSWORD'] = (os.getenv('EMAIL_SMTP_PASSWORD') or '').strip()
app.config['EMAIL_USE_TLS'] = os.getenv('EMAIL_USE_TLS', '1') == '1'
app.config['EMAIL_FROM'] = (os.getenv('EMAIL_FROM') or 'noreply@qr-attendance.local').strip()
app.config['EMAIL_TIMEOUT_SECONDS'] = _safe_int_env('EMAIL_TIMEOUT_SECONDS', 10, minimum=1, maximum=120)
db.init_app(app)
_configure_app_logging(app)


@app.before_request
def enforce_canonical_host():
    """Keep navigation on one host (LAN IP) to prevent mixed 127/192 session flow issues."""
    if not app.config.get('FORCE_CANONICAL_HOST'):
        return None

    base = urlparse(app.config['QR_BASE_URL'])
    canonical_host = base.netloc
    current_host = request.host

    if not canonical_host or current_host == canonical_host:
        return None

    if current_host.startswith('127.0.0.1') or current_host.startswith('localhost'):
        target = f"{base.scheme}://{canonical_host}{request.full_path}"
        return redirect(target.rstrip('?'))

    return None


@app.before_request
def mark_request_start_time():
    g.request_start_ts = time.perf_counter()


@app.before_request
def enforce_account_lock_policy():
    user_id = session.get('user_id')
    if not user_id:
        return None

    if request.endpoint in {'static', 'login', 'logout', 'set_language'}:
        return None

    user = db.session.get(User, user_id)
    if not user:
        session.clear()
        return redirect(url_for('login'))

    if not user.is_locked:
        return None

    lang = _get_language()
    session.clear()
    session['lang'] = lang
    flash('Hesabınız kilitlenmiş. Lütfen bir yönetici ile iletişime geçin.', 'danger')
    return redirect(url_for('login'))

# --- DB Migration Helpers ---
# Adds new columns if they don't exist yet.
def _ensure_column(table: str, column: str, definition: str):
    with app.app_context():
        from sqlalchemy import inspect as sa_inspect
        inspector = sa_inspect(db.engine)
        existing = {col['name'] for col in inspector.get_columns(table)}
        if column not in existing:
            normalized_definition = definition
            if db.engine.dialect.name == 'postgresql' and definition.strip().upper() == 'DATETIME':
                normalized_definition = 'TIMESTAMP'
            with db.engine.begin() as conn:
                conn.execute(text(f"ALTER TABLE {table} ADD COLUMN {column} {normalized_definition}"))


def _attendance_url(token: str) -> str:
    return f"{app.config['QR_BASE_URL']}/attendance/{token}"


def _session_deadline(session_obj: AttendanceSession) -> datetime.datetime:
    return session_obj.date + datetime.timedelta(minutes=app.config['ATTENDANCE_WINDOW_MINUTES'])


def _auto_finalize_expired_sessions(class_id: int | None = None):
    query = AttendanceSession.query.filter_by(active=True)
    if class_id is not None:
        query = query.filter_by(class_id=class_id)

    now = datetime.datetime.now()
    changed = False
    for sess in query.all():
        if now >= _session_deadline(sess):
            sess.active = False
            sess.confirmed = True
            changed = True

    if changed:
        db.session.commit()


def _is_safe_next_url(target: str) -> bool:
    if not target:
        return False
    ref = urlparse(request.host_url)
    test = urlparse(target)
    return test.scheme in ('http', 'https', '') and (not test.netloc or test.netloc == ref.netloc)


def _is_rate_limited(scope: str, key: str, limit: int, window_seconds: int, now_ts: float | None = None) -> bool:
    if limit <= 0 or window_seconds <= 0:
        return False

    bucket_key = f"{scope}:{key}"
    bucket = RATE_LIMIT_EVENTS.setdefault(bucket_key, deque())
    current_ts = now_ts if now_ts is not None else datetime.datetime.now(datetime.UTC).timestamp()
    cutoff_ts = current_ts - window_seconds

    while bucket and bucket[0] <= cutoff_ts:
        bucket.popleft()

    if len(bucket) >= limit:
        return True

    bucket.append(current_ts)
    return False


def _record_login_failure_and_lock(user: User | None, email: str, now_ts: float | None = None) -> bool:
    """Track failed login attempts by email and lock matching user if threshold is reached."""
    if not user:
        return False

    max_failures = int(app.config.get('LOGIN_LOCK_MAX_FAILURES', 0) or 0)
    window_seconds = int(app.config.get('LOGIN_LOCK_WINDOW_SECONDS', 0) or 0)
    if max_failures <= 0 or window_seconds <= 0:
        return False

    bucket_key = f"loginfail:{email.strip().lower()}"
    bucket = RATE_LIMIT_EVENTS.setdefault(bucket_key, deque())
    current_ts = now_ts if now_ts is not None else datetime.datetime.now(datetime.UTC).timestamp()
    cutoff_ts = current_ts - window_seconds

    while bucket and bucket[0] <= cutoff_ts:
        bucket.popleft()

    bucket.append(current_ts)

    if len(bucket) < max_failures or user.is_locked:
        return False

    user.is_locked = True
    db.session.commit()
    return True


def _clear_login_failure_history(email: str):
    RATE_LIMIT_EVENTS.pop(f"loginfail:{email.strip().lower()}", None)


def _admin_mutation_rate_key() -> str:
    actor_id = session.get('user_id') or 'anonymous'
    endpoint = request.endpoint or 'unknown'
    ip = request.remote_addr or 'unknown'
    return f"{actor_id}:{endpoint}:{ip}"


def _get_language() -> str:
    lang = (session.get('lang') or DEFAULT_LANGUAGE).lower()
    return lang if lang in LANGUAGES else DEFAULT_LANGUAGE


def _t(key: str, **kwargs) -> str:
    lang = _get_language()
    template = TRANSLATIONS.get(lang, {}).get(key, TRANSLATIONS[DEFAULT_LANGUAGE].get(key, key))
    if kwargs:
        try:
            return template.format(**kwargs)
        except Exception:
            return template
    return template


def _lang_text(tr_text: str, en_text: str) -> str:
    return tr_text if _get_language() == 'tr' else en_text


def _get_or_create_csrf_token() -> str:
    token = session.get('_csrf_token')
    if not token:
        token = secrets.token_urlsafe(32)
        session['_csrf_token'] = token
    return token


def _is_csrf_token_valid() -> bool:
    session_token = session.get('_csrf_token')
    form_token = (request.form.get('csrf_token') or '').strip()
    header_token = (request.headers.get('X-CSRF-Token') or '').strip()
    provided_token = form_token or header_token

    if not session_token or not provided_token:
        return False
    return hmac.compare_digest(session_token, provided_token)


@app.context_processor
def inject_i18n_helpers():
    return {
        't': _t,
        'lang_text': _lang_text,
        'current_lang': _get_language(),
        'csrf_token': _get_or_create_csrf_token,
    }


# ------------------ ANNOUNCEMENTS ------------------
def _announcements_for_user(user, role, include_expired=False):
    """Query of announcements addressed to this user.

    Students and teachers get announcements targeted at their role (or everyone) that are either general
    or tied to one of their courses. Admins oversee everything. Scheduled ones never show early; expired
    ones are kept only when include_expired is set (the inbox keeps history, dashboards do not).
    """
    now = datetime.datetime.now()
    query = Announcement.query.filter((Announcement.starts_at.is_(None)) | (Announcement.starts_at <= now))
    if not include_expired:
        query = query.filter((Announcement.ends_at.is_(None)) | (Announcement.ends_at >= now))
    if role in {'admin', 'registrar'}:  # staff oversee every announcement
        return query

    query = query.filter(
        (Announcement.target_role.is_(None)) | (Announcement.target_role == 'all') | (Announcement.target_role == role)
    )
    if role in ('teacher', 'assistant'):
        if role == 'teacher':
            course_ids = [course.id for course in user.courses_taught]
        else:
            codes = [f'CLS{c.id:04d}' for c in user.assisted_classes]
            course_ids = [c.id for c in Course.query.filter(Course.code.in_(codes or ['']))]
        query = query.filter(Announcement.group_id.is_(None))  # department groups address students
    else:
        student = db.session.get(Student, user.id)
        course_ids = [e.course_id for e in student.course_enrollments if e.approved] if student else []
        group_ids = [group.id for group in student.groups] if student else []
        query = query.filter((Announcement.group_id.is_(None)) | (Announcement.group_id.in_(group_ids or [-1])))
    if course_ids:
        return query.filter((Announcement.course_id.is_(None)) | (Announcement.course_id.in_(course_ids)))
    return query.filter(Announcement.course_id.is_(None))


def _unread_announcement_count(user, role):
    """Bell badge: announcements newer than the last inbox visit + unread personal messages."""
    query = _announcements_for_user(user, role, include_expired=True).filter(Announcement.author_id != user.id)
    if user.announcements_seen_at:
        query = query.filter(Announcement.created_at > user.announcements_seen_at)
    unread_messages = Message.query.filter_by(recipient_id=user.id, is_read=False).count()
    return query.count() + unread_messages


def _mailbox_item(obj):
    """Normalise an Announcement or a personal Message into one mailbox row."""
    if isinstance(obj, Message):
        return {
            'key': f'm{obj.id}', 'kind': 'message', 'obj': obj,
            'title': obj.subject or _lang_text('(Konusuz)', '(No subject)'), 'body': obj.body,
            'sender': obj.sender.name if obj.sender else '-',
            'recipient': obj.recipient.name if obj.recipient else '-',
            'created_at': obj.sent_at, 'course': None, 'group': None, 'ends_at': None,
        }
    return {
        'key': f'a{obj.id}', 'kind': 'announcement', 'obj': obj,
        'title': obj.title, 'body': obj.body,
        'sender': obj.author.name if obj.author else '-',
        'recipient': None,  # rendered from target_role / group in the template
        'created_at': obj.created_at, 'course': obj.course, 'group': obj.group, 'ends_at': obj.ends_at,
    }


def _resolve_recipient(raw):
    """Find a user from the composer's person field: 'Name <email>', an e-mail, or a unique name."""
    raw = (raw or '').strip()
    match = re.search(r'<([^<>@\s]+@[^<>\s]+)>', raw)
    email = (match.group(1) if match else raw).strip().lower()
    if '@' in email:
        return User.query.filter(func.lower(User.email) == email).first()
    candidates = User.query.filter(func.lower(User.name) == raw.lower()).limit(2).all()
    return candidates[0] if len(candidates) == 1 else None


# ------------------ TOP NAVIGATION ------------------
@app.context_processor
def inject_top_nav():
    """Signed-in user's name, role and unread announcement count for the shared top navigation."""
    origin = session.get('view_as_origin')
    user_id = origin['user_id'] if origin else session.get('user_id')
    role = origin['role'] if origin else session.get('role')
    empty = {'nav_role': None, 'nav_user_name': None, 'nav_unread_announcements': 0}
    if not user_id or role not in ROLE_DASHBOARDS:
        return empty
    user = db.session.get(User, user_id)
    if not user:
        return empty
    try:
        unread = _unread_announcement_count(user, role)
    except Exception:
        db.session.rollback()
        unread = 0
    return {'nav_role': role, 'nav_user_name': user.name, 'nav_unread_announcements': unread}


def csrf_protect(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if request.method in ('POST', 'PUT', 'PATCH', 'DELETE') and not _is_csrf_token_valid():
            flash('Güvenlik doğrulaması başarısız. Lütfen tekrar deneyin.', 'danger')
            fallback = request.referrer or url_for('login')
            return redirect(fallback)
        return f(*args, **kwargs)
    return decorated_function


def rate_limit_protect(scope: str, limit_config_key: str, window_config_key: str, key_builder):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            limit = int(app.config.get(limit_config_key, 0) or 0)
            window_seconds = int(app.config.get(window_config_key, 0) or 0)
            key = key_builder() if callable(key_builder) else str(key_builder)

            if _is_rate_limited(scope, key, limit, window_seconds):
                flash('Çok fazla istek gönderildi. Lütfen bekleyip tekrar deneyin.', 'danger')
                fallback = request.referrer or url_for('login')
                return redirect(fallback)
            return f(*args, **kwargs)
        return decorated_function
    return decorator


def _percentile(values, pct):
    if not values:
        return 0.0
    sorted_values = sorted(values)
    index = int(round((pct / 100.0) * (len(sorted_values) - 1)))
    index = max(0, min(index, len(sorted_values) - 1))
    return float(sorted_values[index])


def _build_health_snapshot():
    db_ok = True
    db_error = None
    try:
        db.session.execute(text('SELECT 1'))
    except Exception as exc:
        db.session.rollback()
        db_ok = False
        db_error = str(exc)

    total_requests = int(REQUEST_METRICS['total_requests'])
    error_total = int(REQUEST_METRICS['error_4xx']) + int(REQUEST_METRICS['error_5xx'])
    error_rate_pct = (error_total / total_requests * 100.0) if total_requests else 0.0

    avg_values = []
    p95_source = []
    for bucket in REQUEST_METRICS['by_endpoint'].values():
        if bucket['count'] > 0:
            avg = bucket['total_ms'] / bucket['count']
            avg_values.append(avg)
            p95_source.append(avg)

    global_avg_ms = (sum(avg_values) / len(avg_values)) if avg_values else 0.0
    p95_ms = _percentile(p95_source, 95)

    warn_error_rate = float(app.config.get('HEALTH_WARN_ERROR_RATE_PCT', 5))
    warn_p95 = float(app.config.get('HEALTH_WARN_P95_MS', 1000))

    status = 'healthy'
    reasons = []
    if not db_ok:
        status = 'unhealthy'
        reasons.append('database_check_failed')
    elif error_rate_pct >= warn_error_rate or p95_ms >= warn_p95:
        status = 'degraded'
        if error_rate_pct >= warn_error_rate:
            reasons.append('error_rate_threshold_exceeded')
        if p95_ms >= warn_p95:
            reasons.append('latency_threshold_exceeded')

    return {
        'status': status,
        'reasons': reasons,
        'generated_at_utc': datetime.datetime.now(datetime.UTC).isoformat().replace('+00:00', 'Z'),
        'since_utc': REQUEST_METRICS['since_utc'],
        'database_ok': db_ok,
        'database_error': db_error,
        'totals': {
            'requests': total_requests,
            'error_4xx': int(REQUEST_METRICS['error_4xx']),
            'error_5xx': int(REQUEST_METRICS['error_5xx']),
            'error_rate_pct': round(error_rate_pct, 2),
        },
        'latency_ms': {
            'global_avg': round(global_avg_ms, 2),
            'p95_estimate': round(p95_ms, 2),
            'slow_threshold': int(app.config.get('METRICS_SLOW_REQUEST_MS', 750)),
            'warn_threshold': int(warn_p95),
        },
        'thresholds': {
            'warn_error_rate_pct': warn_error_rate,
            'warn_p95_ms': warn_p95,
        },
    }


@app.after_request
def collect_request_metrics(response):
    started_at = getattr(g, 'request_start_ts', None)
    if started_at is None:
        return response

    duration_ms = (time.perf_counter() - started_at) * 1000.0
    endpoint = request.endpoint or 'unknown'
    method = request.method
    status_code = int(response.status_code)

    REQUEST_METRICS['total_requests'] += 1
    if 400 <= status_code < 500:
        REQUEST_METRICS['error_4xx'] += 1
    if status_code >= 500:
        REQUEST_METRICS['error_5xx'] += 1

    key = f"{method} {endpoint}"
    bucket = REQUEST_METRICS['by_endpoint'].get(key)
    if not bucket:
        bucket = {
            'count': 0,
            'total_ms': 0.0,
            'max_ms': 0.0,
            'min_ms': duration_ms,
            'last_status': status_code,
            'last_seen_utc': datetime.datetime.now(datetime.UTC).isoformat().replace('+00:00', 'Z'),
        }
        REQUEST_METRICS['by_endpoint'][key] = bucket

    bucket['count'] += 1
    bucket['total_ms'] += duration_ms
    bucket['max_ms'] = max(bucket['max_ms'], duration_ms)
    bucket['min_ms'] = min(bucket['min_ms'], duration_ms)
    bucket['last_status'] = status_code
    bucket['last_seen_utc'] = datetime.datetime.now(datetime.UTC).isoformat().replace('+00:00', 'Z')

    slow_ms = int(app.config.get('METRICS_SLOW_REQUEST_MS', 750) or 750)
    if duration_ms >= slow_ms:
        app.logger.warning(
            'Slow request detected: endpoint=%s method=%s status=%s duration_ms=%.2f',
            endpoint,
            method,
            status_code,
            duration_ms,
        )

    response.headers.setdefault('X-Content-Type-Options', 'nosniff')
    response.headers.setdefault('X-Frame-Options', 'DENY')
    response.headers.setdefault('Referrer-Policy', 'strict-origin-when-cross-origin')
    response.headers.setdefault('Permissions-Policy', 'camera=(), microphone=(), geolocation=()')
    response.headers.setdefault(
        'Content-Security-Policy',
        "default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; script-src 'self' 'unsafe-inline'; font-src 'self' data:",
    )

    return response


# wttr.in weather codes -> Turkish (its own lang=tr text is often still English).
WEATHER_CODES_TR = {
    '113': 'Açık', '116': 'Parçalı bulutlu', '119': 'Bulutlu', '122': 'Kapalı',
    '143': 'Puslu', '248': 'Sisli', '260': 'Dondurucu sis',
    '176': 'Yer yer yağmur', '263': 'Çisenti', '266': 'Çisenti', '281': 'Dondurucu çisenti', '284': 'Dondurucu çisenti',
    '293': 'Hafif yağmur', '296': 'Hafif yağmur', '299': 'Yağmurlu', '302': 'Yağmurlu', '305': 'Şiddetli yağmur',
    '308': 'Şiddetli yağmur', '311': 'Dondurucu yağmur', '314': 'Dondurucu yağmur', '353': 'Hafif sağanak',
    '356': 'Sağanak yağış', '359': 'Kuvvetli sağanak', '200': 'Yakında gök gürültülü sağanak',
    '386': 'Gök gürültülü yağmur', '389': 'Gök gürültülü kuvvetli yağmur', '392': 'Gök gürültülü kar',
    '395': 'Gök gürültülü yoğun kar', '179': 'Yer yer kar', '182': 'Karla karışık yağmur', '185': 'Karla karışık yağmur',
    '317': 'Karla karışık yağmur', '320': 'Karla karışık yağmur', '362': 'Karla karışık sağanak',
    '365': 'Karla karışık sağanak', '227': 'Tipi', '230': 'Kar fırtınası', '323': 'Hafif kar', '326': 'Hafif kar',
    '329': 'Kar yağışlı', '332': 'Kar yağışlı', '335': 'Yoğun kar', '338': 'Yoğun kar', '368': 'Kar sağanağı',
    '371': 'Kuvvetli kar sağanağı', '350': 'Dolu', '374': 'Dolu', '377': 'Dolu',
}


def _fetch_weather(city: str | None = None, lat: float | None = None, lon: float | None = None) -> dict | None:
    """Fetch current weather from wttr.in using coordinates when available, with city fallback."""
    query = None
    display_city = city

    if lat is not None and lon is not None:
        query = f"{lat:.4f},{lon:.4f}"
        display_city = f"{lat:.4f}, {lon:.4f}"
    elif city:
        query = quote(city.strip())

    if not query:
        return None

    try:
        req = Request(
            f"https://wttr.in/{query}?format=j1&lang=tr",
            headers={"User-Agent": "qr-attendance-portal/1.0"},
        )
        with urlopen(req, timeout=3.0) as resp:
            payload = json.loads(resp.read().decode('utf-8'))

        current = (payload.get('current_condition') or [{}])[0]
        desc_list = current.get('weatherDesc') or [{}]
        nearest = (payload.get('nearest_area') or [{}])[0]
        area_name = ((nearest.get('areaName') or [{}])[0]).get('value') if nearest else None
        today = (payload.get('weather') or [{}])[0]
        condition_tr = ((current.get('lang_tr') or [{}])[0]).get('value')
        condition_en = desc_list[0].get('value', 'Unknown')
        return {
            'city': area_name or display_city or 'Unknown',
            'temperature_c': current.get('temp_C', '-'),
            'min_c': today.get('mintempC', '-'),
            'max_c': today.get('maxtempC', '-'),
            'code': current.get('weatherCode', ''),
            'condition': _lang_text(WEATHER_CODES_TR.get(str(current.get('weatherCode', ''))) or condition_tr or condition_en,
                                    condition_en),
            'humidity': current.get('humidity', '-'),
            'wind_kmph': current.get('windspeedKmph', '-'),
            'source': 'location' if lat is not None and lon is not None else 'city',
            'ok': True,
        }
    except Exception:
        return {
            'city': display_city or 'Unknown',
            'temperature_c': '-',
            'min_c': '-',
            'max_c': '-',
            'code': '',
            'condition': _lang_text('Alınamadı', 'Unavailable'),
            'humidity': '-',
            'wind_kmph': '-',
            'source': 'location' if lat is not None and lon is not None else 'city',
            'ok': False,
        }

with app.app_context():
    db.create_all()

    # Add columns if they do not exist yet.
    _ensure_column('classes', 'qr_token', 'TEXT')
    _ensure_column('classes', 'qr_filename', 'TEXT')
    _ensure_column('users', 'is_locked', 'BOOLEAN DEFAULT 0')
    _ensure_column('users', 'announcements_seen_at', 'TIMESTAMP')
    _ensure_column('announcements', 'group_id', 'INTEGER REFERENCES student_groups(id)')
    _ensure_column('grade_records', 'entered_by_id', 'INTEGER REFERENCES users(id)')
    _ensure_column('attendance_sessions', 'active', 'BOOLEAN DEFAULT 0')
    _ensure_column('attendance_sessions', 'confirmed', 'BOOLEAN DEFAULT 0')
    _ensure_column('courses', 'capacity', 'INTEGER DEFAULT 60')
    _ensure_column('courses', 'schedule_slot', 'TEXT')
    _ensure_column('courses', 'ects', 'INTEGER NOT NULL DEFAULT 6')
    _ensure_column('course_enrollments', 'approved', 'BOOLEAN NOT NULL DEFAULT TRUE')  # existing rows stay official
    _ensure_column('announcements', 'starts_at', 'DATETIME')
    _ensure_column('announcements', 'ends_at', 'DATETIME')
    _ensure_column('students', 'first_name', 'TEXT')
    _ensure_column('students', 'last_name', 'TEXT')
    _ensure_column('students', 'gender', 'TEXT')
    _ensure_column('students', 'birth_date', 'DATE')
    _ensure_column('students', 'birth_place', 'TEXT')
    _ensure_column('students', 'identity_number', 'TEXT')
    _ensure_column('students', 'nationality', 'TEXT')
    _ensure_column('students', 'registered_city', 'TEXT')
    _ensure_column('students', 'registered_district', 'TEXT')
    _ensure_column('students', 'passport_number', 'TEXT')
    _ensure_column('students', 'passport_active', 'BOOLEAN DEFAULT 0')
    _ensure_column('students', 'passport_issue_date', 'DATE')
    _ensure_column('students', 'passport_issue_place', 'TEXT')
    _ensure_column('students', 'passport_expiry_date', 'DATE')
    _ensure_column('students', 'marital_status', 'TEXT')
    _ensure_column('students', 'blood_type', 'TEXT')
    _ensure_column('students', 'is_veteran_martyr_relative', 'BOOLEAN DEFAULT 0')
    _ensure_column('students', 'is_disabled', 'BOOLEAN DEFAULT 0')
    _ensure_column('students', 'disability_type', 'TEXT')
    _ensure_column('students', 'disability_rate', 'TEXT')
    _ensure_column('students', 'is_employed', 'BOOLEAN DEFAULT 0')
    _ensure_column('students', 'is_group_company', 'BOOLEAN DEFAULT 0')
    _ensure_column('students', 'company_name', 'TEXT')
    _ensure_column('students', 'work_type', 'TEXT')
    _ensure_column('students', 'employment_start_date', 'DATE')
    _ensure_column('students', 'university_entry_place', 'TEXT')
    _ensure_column('students', 'university_entry_type', 'TEXT')
    _ensure_column('students', 'university_academic_year', 'TEXT')
    _ensure_column('students', 'university_term', 'TEXT')
    _ensure_column('students', 'university_faculty', 'TEXT')
    _ensure_column('students', 'university_department', 'TEXT')
    _ensure_column('students', 'university_scholarship_type', 'TEXT')
    _ensure_column('students', 'university_placement_type', 'TEXT')
    _ensure_column('students', 'university_score_type', 'TEXT')
    _ensure_column('students', 'university_achievement_score', 'TEXT')
    _ensure_column('students', 'university_placement_score', 'TEXT')
    _ensure_column('students', 'university_preference_order', 'TEXT')
    _ensure_column('students', 'highschool_name', 'TEXT')
    _ensure_column('students', 'highschool_info', 'TEXT')
    _ensure_column('students', 'highschool_graduation_date', 'DATE')
    _ensure_column('students', 'advisor_id', 'INTEGER REFERENCES users(id) ON DELETE SET NULL')
    _ensure_column('student_groups', 'advisor_id', 'INTEGER REFERENCES users(id) ON DELETE SET NULL')
    _ensure_column('students', 'photo_filename', 'VARCHAR(120)')
    _ensure_column('users', 'academic_title', 'VARCHAR(60)')
    _ensure_column('users', 'faculty', 'VARCHAR(140)')
    _ensure_column('users', 'department', 'VARCHAR(140)')
    _ensure_column('users', 'office', 'VARCHAR(80)')
    _ensure_column('users', 'office_hours', 'VARCHAR(160)')
    _ensure_column('users', 'photo_filename', 'VARCHAR(120)')

    # Generate/update class QR files in URL format for all classes.
    qr_folder = os.path.join(app.root_path, "static", "qrcodes")
    if not os.path.exists(qr_folder):
        os.makedirs(qr_folder)

    for cls in Class.query.all():
        if not cls.qr_token:
            cls.qr_token = secrets.token_hex(4)
        if not cls.qr_filename:
            cls.qr_filename = f"{cls.name}_qr_{cls.qr_token}.png"

        qr_path = os.path.join(qr_folder, cls.qr_filename)
        img = qrcode.make(_attendance_url(cls.qr_token))
        img.save(qr_path)
    db.session.commit()

# ------------------ LOGIN REQUIRED DECORATOR ------------------
def login_required_session(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('Lütfen önce giriş yapın!', 'danger')
            return redirect(url_for('login', next=request.url))
        return f(*args, **kwargs)
    return decorated_function


ROLE_DASHBOARDS = {
    'admin': 'admin_dashboard',
    'registrar': 'registrar_dashboard',
    'teacher': 'teacher_dashboard',
    'assistant': 'assistant_dashboard',
    'student': 'student_dashboard',
}


# Announcement audiences: "all" plus every role. A new role in ROLE_DASHBOARDS shows up automatically;
# add a label here to give it a proper plural name.
AUDIENCE_LABELS = {
    'all': ('Herkes', 'Everyone'),
    'student': ('Öğrenciler', 'Students'),
    'teacher': ('Öğretmenler', 'Teachers'),
    'assistant': ('Asistanlar', 'Assistants'),
    'registrar': ('Öğrenci İşleri', 'Student Affairs'),
    'admin': ('Adminler', 'Admins'),
}
ANNOUNCEMENT_TARGETS = {'all', *ROLE_DASHBOARDS}


@app.template_global()
def audience_options():
    """[(value, label)] for announcement recipient pickers, 'all' first."""
    ordered = [r for r in AUDIENCE_LABELS if r in ANNOUNCEMENT_TARGETS]
    ordered += sorted(r for r in ANNOUNCEMENT_TARGETS if r not in AUDIENCE_LABELS)
    return [(r, _lang_text(*AUDIENCE_LABELS[r]) if r in AUDIENCE_LABELS else r.title()) for r in ordered]


def _dashboard_endpoint_for(role):
    """Home page endpoint for a role (teacher dashboard for anything unknown, as before)."""
    return ROLE_DASHBOARDS.get(role, 'teacher_dashboard')


def role_required(required_role):
    """Allow a single role ('admin') or several (('admin', 'registrar'))."""
    allowed_roles = (required_role,) if isinstance(required_role, str) else tuple(required_role)

    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'user_id' not in session:
                flash('Lütfen önce giriş yapın!', 'danger')
                return redirect(url_for('login', next=request.url))
            if session.get('role') not in allowed_roles:
                return redirect(url_for('login'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator


# Day 3: Minimal permission layer on top of existing role guards.
PERMISSIONS = {
    'TEACHER_DASHBOARD_VIEW': 'teacher.dashboard.view',
    'TEACHER_CLASS_HISTORY_READ': 'teacher.class.history.read',
    'TEACHER_CLASS_HISTORY_EXPORT': 'teacher.class.history.export',
    'TEACHER_SESSION_STATS_READ': 'teacher.session.stats.read',
    'TEACHER_SESSION_READ': 'teacher.session.read',
    'TEACHER_SESSION_DELETE': 'teacher.session.delete',
    'TEACHER_ACCOUNT_READ': 'teacher.account.read',
    'TEACHER_ACCOUNT_UPDATE': 'teacher.account.update',
    'TEACHER_CLASS_CREATE': 'teacher.class.create',
    'TEACHER_QR_VIEW': 'teacher.qr.view',
    'TEACHER_CLASS_READ': 'teacher.class.read',
    'TEACHER_SESSION_CREATE': 'teacher.session.create',
    'TEACHER_SESSION_UPDATE_ATTENDANCE': 'teacher.session.update_attendance',
    'TEACHER_SESSION_STOP': 'teacher.session.stop',
    'TEACHER_REPORT_READ': 'teacher.report.read',
    'TEACHER_DASHBOARD_CALENDAR_UPDATE': 'teacher.dashboard.calendar.update',
    'STUDENT_ACCOUNT_READ': 'student.account.read',
    'STUDENT_ACCOUNT_UPDATE': 'student.account.update',
    'STUDENT_DASHBOARD_VIEW': 'student.dashboard.view',
    'STUDENT_DASHBOARD_CALENDAR_UPDATE': 'student.dashboard.calendar.update',
    'STUDENT_ABSENCE_READ': 'student.absence.read',
    'STUDENT_ABSENCE_UPDATE': 'student.absence.update',
    'STUDENT_IDENTITY_READ': 'student.identity.read',
    'STUDENT_EDUCATION_READ': 'student.education.read',
    'STUDENT_FAMILY_READ': 'student.family.read',
    'STUDENT_DOCUMENTS_READ': 'student.documents.read',
    'STUDENT_CONTACT_READ': 'student.contact.read',
    'STUDENT_CURRENT_ACCOUNT_READ': 'student.current_account.read',
    'STUDENT_PAYMENT_READ': 'student.payment.read',
    'STUDENT_TERM_COURSES_READ': 'student.term_courses.read',
    'STUDENT_TERM_COURSES_UPDATE': 'student.term_courses.update',
    'STUDENT_TRANSCRIPT_READ': 'student.transcript.read',
    'STUDENT_ACADEMIC_CALENDAR_READ': 'student.academic_calendar.read',
    'STUDENT_EXAMS_READ': 'student.exams.read',
    'STUDENT_CLASS_HISTORY_READ': 'student.class.history.read',
    'STUDENT_ATTENDANCE_MARK': 'student.attendance.mark',
    'STUDENT_REPORT_READ': 'student.report.read',
    'ADMIN_DASHBOARD_VIEW': 'admin.dashboard.view',
    'ADMIN_DASHBOARD_CALENDAR_UPDATE': 'admin.dashboard.calendar.update',
    'ADMIN_USER_READ': 'admin.user.read',
    'ADMIN_USER_UPDATE_ROLE': 'admin.user.update_role',
    'ADMIN_USER_LOCK_TOGGLE': 'admin.user.lock_toggle',
    'ADMIN_USER_PASSWORD_RESET': 'admin.user.password_reset',
    'ADMIN_CLASS_ASSIGN_TEACHER': 'admin.class.assign_teacher',
    'ADMIN_AUDIT_READ': 'admin.audit.read',
    'ADMIN_METRICS_READ': 'admin.metrics.read',
    'ADMIN_PERMISSION_MATRIX_READ': 'admin.permission_matrix.read',
    'ADMIN_NOTIFICATION_SEND': 'admin.notification.send',
    'ADMIN_GROUP_MANAGE': 'admin.group.manage',
    'ADMIN_ACCOUNT_UPDATE': 'admin.account.update',
    'REGISTRAR_DASHBOARD_VIEW': 'registrar.dashboard.view',
    'REGISTRAR_DASHBOARD_CALENDAR_UPDATE': 'registrar.dashboard.calendar.update',
    'REGISTRAR_ACCOUNT_UPDATE': 'registrar.account.update',
    'TEACHER_ASSISTANT_MANAGE': 'teacher.assistant.manage',
    'ASSISTANT_DASHBOARD_VIEW': 'assistant.dashboard.view',
    'ASSISTANT_DASHBOARD_CALENDAR_UPDATE': 'assistant.dashboard.calendar.update',
    'ASSISTANT_ACCOUNT_UPDATE': 'assistant.account.update',
    'ADMIN_COURSE_REGISTRATION_MANAGE': 'admin.course_registration.manage',
    'ADMIN_COURSE_REGISTRATION_OVERRIDE': 'admin.course_registration.override',
    'ADMIN_STUDENT_ADVISOR_ASSIGN': 'admin.student.advisor_assign',
    'TEACHER_ADVISEE_READ': 'teacher.advisee.read',
    'TEACHER_STUDENT_VIEW': 'teacher.student.view',
    'TEACHER_TERM_PLAN_DECIDE': 'teacher.term_plan.decide',
}

PERMISSION_MAP = {
    'teacher_dashboard': PERMISSIONS['TEACHER_DASHBOARD_VIEW'],
    'teacher_attendance_hub': PERMISSIONS['TEACHER_DASHBOARD_VIEW'],
    'teacher_grade_entry': PERMISSIONS['TEACHER_CLASS_READ'],
    'teacher_course_roster': PERMISSIONS['TEACHER_CLASS_READ'],
    'teacher_class_history': PERMISSIONS['TEACHER_CLASS_HISTORY_READ'],
    'export_teacher_class_history': PERMISSIONS['TEACHER_CLASS_HISTORY_EXPORT'],
    'session_stats': PERMISSIONS['TEACHER_SESSION_STATS_READ'],
    'teacher_history_redirect': PERMISSIONS['TEACHER_CLASS_HISTORY_READ'],
    'session_detail': PERMISSIONS['TEACHER_SESSION_READ'],
    'delete_session': PERMISSIONS['TEACHER_SESSION_DELETE'],
    'teacher_account': PERMISSIONS['TEACHER_ACCOUNT_READ'],
    'create_class': PERMISSIONS['TEACHER_CLASS_READ'],
    'view_qr': PERMISSIONS['TEACHER_QR_VIEW'],
    'class_detail': PERMISSIONS['TEACHER_CLASS_READ'],
    'create_session': PERMISSIONS['TEACHER_SESSION_CREATE'],
    'update_attendance': PERMISSIONS['TEACHER_SESSION_UPDATE_ATTENDANCE'],
    'stop_session': PERMISSIONS['TEACHER_SESSION_STOP'],
    'student_account': PERMISSIONS['STUDENT_ACCOUNT_READ'],
    'student_dashboard': PERMISSIONS['STUDENT_DASHBOARD_VIEW'],
    'student_absence': PERMISSIONS['STUDENT_ABSENCE_READ'],
    'student_identity_info': PERMISSIONS['STUDENT_IDENTITY_READ'],
    'student_education_info': PERMISSIONS['STUDENT_EDUCATION_READ'],
    'student_family_info': PERMISSIONS['STUDENT_FAMILY_READ'],
    'student_documents_info': PERMISSIONS['STUDENT_DOCUMENTS_READ'],
    'student_contact_info': PERMISSIONS['STUDENT_CONTACT_READ'],
    'student_current_account': PERMISSIONS['STUDENT_CURRENT_ACCOUNT_READ'],
    'student_payment_info': PERMISSIONS['STUDENT_PAYMENT_READ'],
    'student_term_courses': PERMISSIONS['STUDENT_TERM_COURSES_READ'],
    'student_term_courses_update': PERMISSIONS['STUDENT_TERM_COURSES_UPDATE'],
    'student_transcript': PERMISSIONS['STUDENT_TRANSCRIPT_READ'],
    'student_academic_calendar': PERMISSIONS['STUDENT_ACADEMIC_CALENDAR_READ'],
    'student_exams': PERMISSIONS['STUDENT_EXAMS_READ'],
    'student_class_history': PERMISSIONS['STUDENT_CLASS_HISTORY_READ'],
    'mark_attendance': PERMISSIONS['STUDENT_ATTENDANCE_MARK'],
    'admin_dashboard': PERMISSIONS['ADMIN_DASHBOARD_VIEW'],
    'admin_user_inventory': PERMISSIONS['ADMIN_USER_READ'],
    'admin_update_user_role': PERMISSIONS['ADMIN_USER_UPDATE_ROLE'],
    'admin_toggle_user_lock': PERMISSIONS['ADMIN_USER_LOCK_TOGGLE'],
    'admin_reset_user_password': PERMISSIONS['ADMIN_USER_PASSWORD_RESET'],
    'admin_class_assignments': PERMISSIONS['ADMIN_CLASS_ASSIGN_TEACHER'],
    'admin_assign_class_teacher': PERMISSIONS['ADMIN_CLASS_ASSIGN_TEACHER'],
    'admin_permission_audit_report': PERMISSIONS['ADMIN_AUDIT_READ'],
    'admin_operation_audit_report': PERMISSIONS['ADMIN_AUDIT_READ'],
    'admin_health_status_report': PERMISSIONS['ADMIN_METRICS_READ'],
    'admin_permission_matrix': PERMISSIONS['ADMIN_PERMISSION_MATRIX_READ'],
    'admin_notifications': PERMISSIONS['ADMIN_NOTIFICATION_SEND'],
    'admin_groups': PERMISSIONS['ADMIN_GROUP_MANAGE'],
    'admin_account': PERMISSIONS['ADMIN_ACCOUNT_UPDATE'],
    'registrar_dashboard': PERMISSIONS['REGISTRAR_DASHBOARD_VIEW'],
    'registrar_account': PERMISSIONS['REGISTRAR_ACCOUNT_UPDATE'],
    'teacher_class_assistants': PERMISSIONS['TEACHER_ASSISTANT_MANAGE'],
    'assistant_dashboard': PERMISSIONS['ASSISTANT_DASHBOARD_VIEW'],
    'assistant_account': PERMISSIONS['ASSISTANT_ACCOUNT_UPDATE'],
    'admin_course_registration_window': PERMISSIONS['ADMIN_COURSE_REGISTRATION_MANAGE'],
    'admin_course_registration_override': PERMISSIONS['ADMIN_COURSE_REGISTRATION_OVERRIDE'],
    'admin_set_student_advisor': PERMISSIONS['ADMIN_STUDENT_ADVISOR_ASSIGN'],
    'teacher_advisees': PERMISSIONS['TEACHER_ADVISEE_READ'],
    'teacher_advisee_detail': PERMISSIONS['TEACHER_ADVISEE_READ'],
    'teacher_view_student': PERMISSIONS['TEACHER_STUDENT_VIEW'],
    'teacher_term_plan_decision': PERMISSIONS['TEACHER_TERM_PLAN_DECIDE'],
    'student_term_plan_submit': PERMISSIONS['STUDENT_TERM_COURSES_UPDATE'],
}

ROLE_PERMISSIONS = {
    'teacher': {
        PERMISSIONS['TEACHER_DASHBOARD_VIEW'],
        PERMISSIONS['TEACHER_CLASS_HISTORY_EXPORT'],
        PERMISSIONS['TEACHER_SESSION_STATS_READ'],
        PERMISSIONS['TEACHER_CLASS_HISTORY_READ'],
        PERMISSIONS['TEACHER_SESSION_READ'],
        PERMISSIONS['TEACHER_SESSION_DELETE'],
        PERMISSIONS['TEACHER_ACCOUNT_READ'],
        PERMISSIONS['TEACHER_ACCOUNT_UPDATE'],
        PERMISSIONS['TEACHER_CLASS_CREATE'],
        PERMISSIONS['TEACHER_QR_VIEW'],
        PERMISSIONS['TEACHER_CLASS_READ'],
        PERMISSIONS['TEACHER_SESSION_CREATE'],
        PERMISSIONS['TEACHER_SESSION_UPDATE_ATTENDANCE'],
        PERMISSIONS['TEACHER_SESSION_STOP'],
        PERMISSIONS['TEACHER_REPORT_READ'],
        PERMISSIONS['TEACHER_DASHBOARD_CALENDAR_UPDATE'],
        PERMISSIONS['TEACHER_ASSISTANT_MANAGE'],
        PERMISSIONS['TEACHER_ADVISEE_READ'],
        PERMISSIONS['TEACHER_STUDENT_VIEW'],
        PERMISSIONS['TEACHER_TERM_PLAN_DECIDE'],
    },
    # Teaching assistant: attendance, roster and grade entry, only in classes the teacher assigned them to
    # (scope enforced by ensure_teacher_class_ownership / _accessible_classes). No session delete,
    # class creation, teacher dashboard or reports.
    'assistant': {
        PERMISSIONS['ASSISTANT_DASHBOARD_VIEW'],
        PERMISSIONS['ASSISTANT_DASHBOARD_CALENDAR_UPDATE'],
        PERMISSIONS['ASSISTANT_ACCOUNT_UPDATE'],
        PERMISSIONS['TEACHER_DASHBOARD_VIEW'],
        PERMISSIONS['TEACHER_CLASS_READ'],
        PERMISSIONS['TEACHER_CLASS_HISTORY_READ'],
        PERMISSIONS['TEACHER_CLASS_HISTORY_EXPORT'],
        PERMISSIONS['TEACHER_SESSION_STATS_READ'],
        PERMISSIONS['TEACHER_SESSION_READ'],
        PERMISSIONS['TEACHER_QR_VIEW'],
        PERMISSIONS['TEACHER_SESSION_CREATE'],
        PERMISSIONS['TEACHER_SESSION_UPDATE_ATTENDANCE'],
        PERMISSIONS['TEACHER_SESSION_STOP'],
    },
    'student': {
        PERMISSIONS['STUDENT_ACCOUNT_READ'],
        PERMISSIONS['STUDENT_ACCOUNT_UPDATE'],
        PERMISSIONS['STUDENT_DASHBOARD_VIEW'],
        PERMISSIONS['STUDENT_DASHBOARD_CALENDAR_UPDATE'],
        PERMISSIONS['STUDENT_ABSENCE_READ'],
        PERMISSIONS['STUDENT_ABSENCE_UPDATE'],
        PERMISSIONS['STUDENT_IDENTITY_READ'],
        PERMISSIONS['STUDENT_EDUCATION_READ'],
        PERMISSIONS['STUDENT_FAMILY_READ'],
        PERMISSIONS['STUDENT_DOCUMENTS_READ'],
        PERMISSIONS['STUDENT_CONTACT_READ'],
        PERMISSIONS['STUDENT_CURRENT_ACCOUNT_READ'],
        PERMISSIONS['STUDENT_PAYMENT_READ'],
        PERMISSIONS['STUDENT_TERM_COURSES_READ'],
        PERMISSIONS['STUDENT_TERM_COURSES_UPDATE'],
        PERMISSIONS['STUDENT_TRANSCRIPT_READ'],
        PERMISSIONS['STUDENT_ACADEMIC_CALENDAR_READ'],
        PERMISSIONS['STUDENT_EXAMS_READ'],
        PERMISSIONS['STUDENT_CLASS_HISTORY_READ'],
        PERMISSIONS['STUDENT_ATTENDANCE_MARK'],
        PERMISSIONS['STUDENT_REPORT_READ'],
    },
    # Student Affairs (Öğrenci İşleri): department groups, announcements, user list read-only
    # plus password reset for student accounts only (enforced in admin_reset_user_password).
    'registrar': {
        PERMISSIONS['REGISTRAR_DASHBOARD_VIEW'],
        PERMISSIONS['REGISTRAR_DASHBOARD_CALENDAR_UPDATE'],
        PERMISSIONS['REGISTRAR_ACCOUNT_UPDATE'],
        PERMISSIONS['ADMIN_GROUP_MANAGE'],
        PERMISSIONS['ADMIN_NOTIFICATION_SEND'],
        PERMISSIONS['ADMIN_USER_READ'],
        PERMISSIONS['ADMIN_USER_PASSWORD_RESET'],
    },
    'admin': set(PERMISSIONS.values()),
}


def has_permission(role, permission_key):
    if not permission_key:
        return True
    role_permissions = ROLE_PERMISSIONS.get(role, set())
    return permission_key in role_permissions


def _log_permission_denied(permission_key):
    event = {
        'timestamp': datetime.datetime.now(datetime.UTC).isoformat().replace('+00:00', 'Z'),
        'user_id': session.get('user_id'),
        'role': session.get('role'),
        'endpoint': request.endpoint,
        'permission': permission_key,
        'method': request.method,
        'path': request.path,
        'ip': request.remote_addr,
    }
    PERMISSION_AUDIT_EVENTS.appendleft(event)

    try:
        db.session.add(
            PermissionAuditLog(
                user_id=event['user_id'],
                role=event['role'],
                endpoint=event['endpoint'],
                permission=event['permission'],
                method=event['method'],
                path=event['path'],
                ip=event['ip'],
            )
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        app.logger.exception('Failed to persist permission audit event')

    app.logger.warning(
        'Permission denied: user_id=%s role=%s endpoint=%s permission=%s method=%s path=%s ip=%s',
        session.get('user_id'),
        session.get('role'),
        request.endpoint,
        permission_key,
        request.method,
        request.path,
        request.remote_addr,
    )


def ensure_permission(permission_key, redirect_endpoint='login'):
    if has_permission(session.get('role'), permission_key):
        return None

    _log_permission_denied(permission_key)
    flash('Bu işlemi gerçekleştirme yetkiniz yok.', 'danger')
    return redirect(url_for(redirect_endpoint))


def _accessible_classes(user):
    """Classes the signed-in teacher owns, or the classes an assistant was assigned to."""
    if session.get('role') == 'assistant':
        return sorted(user.assisted_classes, key=lambda c: c.name or '')
    return Class.query.filter_by(teacher_id=user.id).all()


def _accessible_courses(user):
    """Course records the teacher owns, or those linked (CLS<class id>) to an assistant's classes."""
    if session.get('role') == 'assistant':
        codes = [f'CLS{c.id:04d}' for c in user.assisted_classes]
        return Course.query.filter(Course.code.in_(codes or [''])).order_by(Course.code.asc()).all()
    return Course.query.filter_by(teacher_id=user.id).order_by(Course.code.asc()).all()


def _accessible_class_or_404(class_id):
    class_obj = db.session.get(Class, class_id)
    if not _can_access_class(class_obj):
        abort(404)
    return class_obj


def _can_access_class(class_obj):
    user_id = session.get('user_id')
    if not class_obj or not user_id:
        return False
    if class_obj.teacher_id == user_id:
        return True
    return session.get('role') == 'assistant' and any(a.id == user_id for a in class_obj.assistants)


def ensure_teacher_class_ownership(class_obj, on_fail='login', fail_message=None):
    if _can_access_class(class_obj):
        return None

    app.logger.warning(
        'Teacher ownership denied: user_id=%s class_id=%s endpoint=%s',
        session.get('user_id'),
        getattr(class_obj, 'id', None),
        request.endpoint,
    )

    if on_fail == 'json':
        return jsonify({'error': 'forbidden'}), 403

    if fail_message:
        flash(fail_message, 'danger')
    return redirect(url_for(on_fail))


def ensure_teacher_session_ownership(session_obj, on_fail='login', fail_message=None):
    if not session_obj:
        return redirect(url_for(on_fail))
    return ensure_teacher_class_ownership(session_obj.class_obj, on_fail=on_fail, fail_message=fail_message)


def ensure_student_class_membership(student_obj, class_obj, on_fail='student_absence', fail_message=None):
    if student_obj and class_obj and class_obj in student_obj.classes:
        return None

    app.logger.warning(
        'Student membership denied: user_id=%s student_id=%s class_id=%s endpoint=%s',
        session.get('user_id'),
        getattr(student_obj, 'id', None),
        getattr(class_obj, 'id', None),
        request.endpoint,
    )

    if fail_message:
        flash(fail_message, 'danger')
    return redirect(url_for(on_fail))


def ensure_student_event_ownership(student_obj, event_obj, on_fail='student_dashboard', fail_message=None):
    if student_obj and event_obj and event_obj.student_id == student_obj.id:
        return None

    app.logger.warning(
        'Student event ownership denied: user_id=%s student_id=%s event_id=%s endpoint=%s',
        session.get('user_id'),
        getattr(student_obj, 'id', None),
        getattr(event_obj, 'id', None),
        request.endpoint,
    )

    if fail_message:
        flash(fail_message, 'warning')
    return redirect(url_for(on_fail))


def ensure_user_calendar_event_ownership(user_obj, event_obj, on_fail='login', fail_message=None):
    if user_obj and event_obj and event_obj.user_id == user_obj.id:
        return None

    app.logger.warning(
        'User calendar event ownership denied: user_id=%s event_id=%s endpoint=%s',
        getattr(user_obj, 'id', None),
        getattr(event_obj, 'id', None),
        request.endpoint,
    )

    if fail_message:
        flash(fail_message, 'warning')
    return redirect(url_for(on_fail))


def _handle_user_calendar_dashboard_post(user_obj, permission_key, redirect_endpoint):
    permission_error = ensure_permission(permission_key, redirect_endpoint)
    if permission_error:
        return permission_error

    action = _normalize_form_text(request.form.get('action'), max_length=40).lower()
    if action not in {'add_calendar_event', 'delete_calendar_event'}:
        flash(_lang_text('Geçersiz panel işlemi.', 'Invalid dashboard action.'), 'warning')
        return redirect(url_for(redirect_endpoint))

    if action == 'add_calendar_event':
        title = (request.form.get('title') or '').strip()
        event_type = (request.form.get('event_type') or 'activity').strip().lower()
        event_date_raw = (request.form.get('event_date') or '').strip()
        note = (request.form.get('note') or '').strip()

        if len(title) > 140:
            flash(_lang_text('Etkinlik başlığı çok uzun.', 'Event title is too long.'), 'warning')
            return redirect(url_for(redirect_endpoint))

        if not title or not event_date_raw:
            flash(_lang_text('Etkinlik başlığı ve tarihi zorunludur.', 'Event title and date are required.'), 'warning')
            return redirect(url_for(redirect_endpoint))

        if event_type not in {'exam', 'activity'}:
            event_type = 'activity'

        try:
            event_date = datetime.datetime.strptime(event_date_raw, '%Y-%m-%d').date()
        except ValueError:
            flash(_lang_text('Geçersiz etkinlik tarihi formatı.', 'Invalid event date format.'), 'warning')
            return redirect(url_for(redirect_endpoint))

        event = UserCalendarEvent(
            user_id=user_obj.id,
            title=title,
            event_type=event_type,
            event_date=event_date,
            note=note[:240] if note else None,
        )
        db.session.add(event)
        db.session.commit()
        flash(_lang_text('Takvim etkinliği eklendi.', 'Calendar event added.'), 'success')
        return redirect(url_for(redirect_endpoint))

    event_id_raw = _normalize_form_text(request.form.get('event_id'), max_length=20)
    if not event_id_raw.isdigit():
        flash(_lang_text('Geçersiz etkinlik kimliği.', 'Invalid event id.'), 'warning')
        return redirect(url_for(redirect_endpoint))

    event = db.session.get(UserCalendarEvent, int(event_id_raw))
    denied = ensure_user_calendar_event_ownership(
        user_obj,
        event,
        on_fail=redirect_endpoint,
        fail_message=_lang_text('Etkinlik bulunamadı.', 'Event not found.'),
    )
    if denied:
        return denied

    db.session.delete(event)
    db.session.commit()
    flash(_lang_text('Takvim etkinliği silindi.', 'Calendar event deleted.'), 'success')
    return redirect(url_for(redirect_endpoint))


def _serialize_user_calendar_events(user_id):
    today = datetime.date.today()
    all_events = UserCalendarEvent.query.filter_by(user_id=user_id).order_by(UserCalendarEvent.event_date.asc()).all()
    upcoming_events = [event for event in all_events if event.event_date >= today][:12]
    serialized_events = [
        {
            'id': event.id,
            'title': event.title,
            'event_type': event.event_type,
            'event_date': event.event_date.isoformat(),
            'note': event.note or '',
        }
        for event in all_events
    ]
    return all_events, upcoming_events, serialized_events


def _dashboard_weather():
    """Home-page weather: ?city= or ?lat&lon from the page, remembered in the session."""
    city_arg = (request.args.get('city') or '').strip()[:50]
    lat_arg = (request.args.get('lat') or '').strip()
    lon_arg = (request.args.get('lon') or '').strip()

    # Manual city search takes precedence and clears the location cache.
    if city_arg:
        session['weather_city'] = city_arg
        session.pop('weather_lat', None)
        session.pop('weather_lon', None)

    if lat_arg and lon_arg:
        try:
            parsed_lat, parsed_lon = float(lat_arg), float(lon_arg)
            if -90 <= parsed_lat <= 90 and -180 <= parsed_lon <= 180:
                session['weather_lat'] = parsed_lat
                session['weather_lon'] = parsed_lon
        except ValueError:
            pass

    weather_lat, weather_lon = session.get('weather_lat'), session.get('weather_lon')
    if weather_lat is not None and weather_lon is not None:
        return _fetch_weather(lat=float(weather_lat), lon=float(weather_lon))
    return _fetch_weather(city=session.get('weather_city') or 'Lefkosa')


def _name_initials(name):
    return ''.join(part[0] for part in (name or '?').split()[:2]).upper()


def _render_home_dashboard(user, role, endpoint, page_title, profile, calendar_events, upcoming_event_count):
    """Shared home page for every role: profile card | calendar | weather + announcements."""
    weather = _dashboard_weather()
    try:
        announcements = _announcements_for_user(user, role).order_by(Announcement.created_at.desc()).limit(6).all()
    except Exception:
        db.session.rollback()
        announcements = []
    profile.setdefault('initials', _name_initials(profile.get('name')))
    return render_template(
        'home_dashboard.html',
        page_title=page_title,
        user_name=user.name,
        profile=profile,
        dashboard_url=url_for(endpoint),
        calendar_events=calendar_events,
        upcoming_event_count=upcoming_event_count,
        weather=weather,
        weather_using_location=bool(weather and weather.get('source') == 'location'),
        announcements=announcements,
    )


def _staff_profile(user, role_label, edit_endpoint, stats):
    """Profile card for staff without an academic profile (admin, registrar, assistant)."""
    return {
        'name': user.name,
        'facts': [
            (_lang_text('E-posta', 'E-mail'), user.email),
            (_lang_text('Rol', 'Role'), role_label),
        ],
        'stats': stats,
        'edit_url': url_for(edit_endpoint),
    }


def permission_required(permission_key=None):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'user_id' not in session:
                flash('Lütfen önce giriş yapın!', 'danger')
                return redirect(url_for('login', next=request.url))

            role = session.get('role')
            resolved_permission = permission_key or PERMISSION_MAP.get(request.endpoint)
            if has_permission(role, resolved_permission):
                return f(*args, **kwargs)

            _log_permission_denied(resolved_permission)
            if request.endpoint == 'session_stats':
                return jsonify({'error': 'forbidden'}), 403

            flash('Bu işlemi gerçekleştirme yetkiniz yok.', 'danger')
            return redirect(url_for('login'))
        return decorated_function
    return decorator


def validate_admin_role_update(actor_user_id, actor_role, target_user_id, target_user_role, new_role, admin_count):
    allowed_roles = {'student', 'teacher', 'assistant', 'registrar', 'admin'}
    if actor_role != 'admin':
        return 'Only admin can update roles.'
    if new_role not in allowed_roles:
        return 'Invalid target role.'
    if actor_user_id == target_user_id:
        return 'You cannot change your own role.'
    if target_user_role == 'admin' and new_role != 'admin' and admin_count <= 1:
        return 'At least one admin account must remain in the system.'
    return None


def validate_admin_lock_update(actor_user_id, actor_role, target_user_id, target_user_role, target_user_locked, lock_state, unlocked_admin_count):
    if actor_role != 'admin':
        return 'Only admin can update lock status.'
    if actor_user_id == target_user_id and lock_state:
        return 'You cannot lock your own account.'
    if target_user_role == 'admin' and lock_state and not target_user_locked and unlocked_admin_count <= 1:
        return 'At least one unlocked admin account must remain in the system.'
    return None


def validate_admin_teacher_assignment(actor_role, class_obj, teacher_user, current_teacher_id, target_teacher_id):
    if actor_role != 'admin':
        return 'Only admin can assign class teacher.'
    if not class_obj:
        return 'Sınıf bulunamadı.'
    if not teacher_user or teacher_user.role != 'teacher':
        return 'Target teacher user is invalid.'
    if teacher_user.is_locked:
        return 'Locked teacher accounts cannot be assigned.'
    if current_teacher_id == target_teacher_id:
        return 'Class is already assigned to selected teacher.'
    return None


def validate_password_policy(password: str, user_name: str | None = None, user_email: str | None = None):
    if password is None:
        return 'Şifre zorunludur.'

    minimum_length = int(app.config.get('PASSWORD_MIN_LENGTH', 10) or 10)
    if len(password) < minimum_length:
        return f'Password must be at least {minimum_length} characters long.'

    has_upper = any(ch.isupper() for ch in password)
    has_lower = any(ch.islower() for ch in password)
    has_digit = any(ch.isdigit() for ch in password)
    has_special = any(not ch.isalnum() for ch in password)
    if not (has_upper and has_lower and has_digit and has_special):
        return 'Password must include uppercase, lowercase, digit, and special character.'

    lowered_password = password.lower()
    if user_name and user_name.strip():
        compact_name = ''.join(ch for ch in user_name.strip().lower() if ch.isalnum())
        if compact_name and compact_name in ''.join(ch for ch in lowered_password if ch.isalnum()):
            return 'Password cannot contain your name.'

    if user_email and '@' in user_email:
        local_part = user_email.split('@', 1)[0].strip().lower()
        if local_part and local_part in lowered_password:
            return 'Password cannot contain your email username.'

    return None


def _normalize_form_text(value, max_length: int | None = None) -> str:
    text = (value or '').strip()
    if max_length is not None and len(text) > max_length:
        return text[:max_length]
    return text


def _parse_filter_datetime(raw_value: str | None):
    raw = (raw_value or '').strip()
    if not raw:
        return None
    for fmt in ('%Y-%m-%d', '%Y-%m-%dT%H:%M', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _parse_optional_int(raw_value: str | None):
    raw = (raw_value or '').strip()
    if raw.isdigit():
        return int(raw)
    return None


def _is_valid_email(email: str) -> bool:
    if not email or len(email) > 254:
        return False
    return re.match(r'^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$', email) is not None


def _send_email_notification(to_email: str, subject: str, body: str) -> tuple[bool, str | None]:
    email_value = _normalize_form_text(to_email, max_length=254).lower()
    if not _is_valid_email(email_value):
        return False, 'invalid_email'

    if app.config.get('EMAIL_DRY_RUN', True):
        app.logger.info('EMAIL_DRY_RUN to=%s subject=%s', email_value, subject)
        return True, 'dry_run'

    smtp_host = (app.config.get('EMAIL_SMTP_HOST') or '').strip()
    smtp_port = int(app.config.get('EMAIL_SMTP_PORT') or 587)
    smtp_user = (app.config.get('EMAIL_SMTP_USER') or '').strip()
    smtp_password = app.config.get('EMAIL_SMTP_PASSWORD') or ''
    use_tls = bool(app.config.get('EMAIL_USE_TLS', True))
    email_from = (app.config.get('EMAIL_FROM') or '').strip()
    timeout_seconds = int(app.config.get('EMAIL_TIMEOUT_SECONDS') or 10)

    if not smtp_host:
        return False, 'smtp_host_missing'
    if not _is_valid_email(email_from):
        return False, 'invalid_sender_email'

    msg = EmailMessage()
    msg['Subject'] = subject
    msg['From'] = email_from
    msg['To'] = email_value
    msg.set_content(body)

    try:
        with smtplib.SMTP(smtp_host, smtp_port, timeout=timeout_seconds) as smtp:
            if use_tls:
                smtp.starttls()
            if smtp_user:
                smtp.login(smtp_user, smtp_password)
            smtp.send_message(msg)
        return True, None
    except Exception as exc:
        app.logger.exception('Email send failed: to=%s subject=%s', email_value, subject)
        return False, str(exc)


def _log_admin_operation(action, target_user_id, old_value, new_value, status='ok', detail=None):
    try:
        db.session.add(
            AdminOperationLog(
                actor_user_id=session.get('user_id'),
                target_user_id=target_user_id,
                action=action,
                old_value=old_value,
                new_value=new_value,
                status=status,
                detail=detail,
                ip=request.remote_addr,
            )
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        app.logger.exception('Failed to persist admin operation log: action=%s target=%s', action, target_user_id)


@app.route('/set-language', methods=['POST'])
def set_language():
    selected = (request.form.get('lang') or '').strip().lower()
    if selected in LANGUAGES:
        session['lang'] = selected
        flash(_t('language_updated'), 'success')

    next_url = request.form.get('next') or request.referrer or url_for('login')
    if not _is_safe_next_url(next_url):
        next_url = url_for('login')
    return redirect(next_url)

# ------------------ REGISTER ------------------
@app.route('/register', methods=['GET', 'POST'])
def register():
    # If already logged in, redirect to the dashboard.
    if 'user_id' in session:
        return redirect(url_for(_dashboard_endpoint_for(session.get('role'))))

    if request.method == 'POST':
        name = _normalize_form_text(request.form.get('name'), max_length=120)
        email = _normalize_form_text(request.form.get('email'), max_length=254).lower()
        password = request.form.get('password') or ''
        role = _normalize_form_text(request.form.get('role'), max_length=20).lower()
        student_number = _normalize_form_text(request.form.get('student_number'), max_length=32)

        if not name:
            flash('Ad alanı zorunludur!', 'danger')
            return redirect(url_for('register'))

        if not _is_valid_email(email):
            flash('Lütfen geçerli bir e-posta adresi girin.', 'danger')
            return redirect(url_for('register'))

        if role not in {'student', 'teacher'}:
            flash('Geçersiz rol seçimi.', 'danger')
            return redirect(url_for('register'))

        password_error = validate_password_policy(password, user_name=name, user_email=email)
        if password_error:
            flash(password_error, 'danger')
            return redirect(url_for('register'))

        if role == 'student' and not student_number:
            flash('Öğrenci numarası zorunludur!', 'danger')
            return redirect(url_for('register'))

        if student_number and not re.match(r'^[A-Za-z0-9_-]+$', student_number):
            flash('Öğrenci numarası geçersiz karakterler içeriyor.', 'danger')
            return redirect(url_for('register'))

        if User.query.filter_by(email=email).first():
            flash('Bu e-posta adresi zaten kayıtlı!', 'danger')
            return redirect(url_for('register'))

        user = User(name=name, email=email, role=role)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()

        # Create student profile for student role.
        if role == 'student':
            name_parts = [part for part in name.strip().split(' ') if part]
            first_name = name_parts[0] if name_parts else None
            last_name = ' '.join(name_parts[1:]) if len(name_parts) > 1 else None
            student = Student(
                id=user.id,
                student_number=student_number,
                first_name=first_name,
                last_name=last_name,
            )
            db.session.add(student)
            db.session.commit()

        # Auto-login after registration.
        session['user_id'] = user.id
        session['role'] = user.role

        flash('Kayıt başarılı! Hoş geldiniz.', 'success')
        if user.role == 'student':
            return redirect(url_for('student_dashboard'))
        if user.role == 'admin':
            return redirect(url_for('admin_dashboard'))
        else:
            return redirect(url_for('teacher_dashboard'))

    return render_template('register.html')

# ------------------ LOGIN ------------------
@app.route('/login', methods=['GET', 'POST'])
def login():
    next_url = request.args.get('next') or request.form.get('next')

    # If already logged in, redirect to the dashboard.
    if 'user_id' in session:
        if _is_safe_next_url(next_url):
            return redirect(next_url)
        return redirect(url_for(_dashboard_endpoint_for(session.get('role'))))

    if request.method == 'POST':
        email = _normalize_form_text(request.form.get('email'), max_length=254).lower()
        password = request.form.get('password') or ''
        remember_me = request.form.get('remember_me') == 'on'

        if not email or not password:
            flash('E-posta veya şifre hatalı!', 'danger')
            return redirect(url_for('login', next=next_url) if next_url else url_for('login'))

        if not _is_valid_email(email):
            flash('E-posta veya şifre hatalı!', 'danger')
            return redirect(url_for('login', next=next_url) if next_url else url_for('login'))

        rate_key = f"{(request.remote_addr or 'unknown').strip()}:{email.strip().lower()}"
        if _is_rate_limited(
            'login',
            rate_key,
            app.config['LOGIN_RATE_LIMIT_MAX_ATTEMPTS'],
            app.config['LOGIN_RATE_LIMIT_WINDOW_SECONDS'],
        ):
            flash('Çok fazla giriş denemesi. Lütfen bekleyip tekrar deneyin.', 'danger')
            return redirect(url_for('login', next=next_url) if next_url else url_for('login'))

        user = User.query.filter_by(email=email).first()

        if user and user.check_password(password):
            if user.is_locked:
                flash('Hesabınız kilitlenmiş. Lütfen bir yönetici ile iletişime geçin.', 'danger')
                return redirect(url_for('login', next=next_url) if next_url else url_for('login'))

            # Reset previous session data before establishing a new login session.
            lang = _get_language()
            session.clear()
            session['lang'] = lang
            session['user_id'] = user.id
            session['role'] = user.role
            session.permanent = remember_me
            RATE_LIMIT_EVENTS.pop(f"login:{rate_key}", None)
            _clear_login_failure_history(email)
            flash('Giriş başarılı!', 'success')
            if _is_safe_next_url(next_url):
                return redirect(next_url)
            return redirect(url_for(_dashboard_endpoint_for(user.role)))

        if _record_login_failure_and_lock(user, email):
            flash('Hesabınız kilitlenmiş. Lütfen bir yönetici ile iletişime geçin.', 'danger')
            return redirect(url_for('login', next=next_url) if next_url else url_for('login'))

        flash('E-posta veya şifre hatalı!', 'danger')
        return redirect(url_for('login', next=next_url) if next_url else url_for('login'))

    return render_template('login.html', next_url=next_url)


@app.route('/')
def index():
    return redirect(url_for(ROLE_DASHBOARDS.get(session.get('role'), 'login')))

# ------------------ LOGOUT ------------------
@app.route('/logout')
def logout():
    lang = _get_language()
    session.clear()
    session['lang'] = lang
    flash('Çıkış yapıldı.', 'success')
    return redirect(url_for('login'))


@app.route('/admin/dashboard', methods=['GET', 'POST'])
@role_required('admin')
@permission_required()
def admin_dashboard():
    admin_user = db.session.get(User, session['user_id'])
    if request.method == 'POST':
        return _handle_user_calendar_dashboard_post(admin_user, PERMISSIONS['ADMIN_DASHBOARD_CALENDAR_UPDATE'], 'admin_dashboard')

    profile = _staff_profile(admin_user, _lang_text('Yönetici', 'Administrator'), 'admin_account', [
        (User.query.filter_by(role='teacher').count(), _lang_text('Öğretmen', 'Teachers')),
        (User.query.filter_by(role='student').count(), _lang_text('Öğrenci', 'Students')),
        (Class.query.count(), _lang_text('Ders', 'Classes')),
        (AttendanceSession.query.filter_by(active=True).count(), _lang_text('Aktif oturum', 'Active sessions')),
    ])
    _, upcoming, events = _serialize_user_calendar_events(admin_user.id)
    return _render_home_dashboard(admin_user, 'admin', 'admin_dashboard', _lang_text('Yönetici Paneli', 'Admin Dashboard'),
                                  profile, events, len(upcoming))


@app.route('/admin/class-assignments')
@role_required('admin')
@permission_required()
def admin_class_assignments():
    q = (request.args.get('q') or '').strip()
    limit_raw = (request.args.get('limit') or '200').strip()
    limit = int(limit_raw) if limit_raw.isdigit() else 200
    limit = max(1, min(limit, 1000))

    class_query = Class.query
    if q:
        class_query = class_query.filter(Class.name.ilike(f"%{q}%"))

    classes = class_query.order_by(Class.id.asc()).limit(limit).all()
    teachers = User.query.filter_by(role='teacher').order_by(User.name.asc(), User.id.asc()).all()

    class_rows = []
    for class_obj in classes:
        teacher_user = db.session.get(User, class_obj.teacher_id)
        class_rows.append(
            {
                'id': class_obj.id,
                'name': class_obj.name,
                'teacher_id': class_obj.teacher_id,
                'teacher_name': teacher_user.name if teacher_user else 'Unknown',
                'teacher_email': teacher_user.email if teacher_user else '-',
            }
        )

    teacher_options = [
        {
            'id': teacher.id,
            'name': teacher.name,
            'email': teacher.email,
            'is_locked': bool(teacher.is_locked),
        }
        for teacher in teachers
    ]

    summary = {
        'total_loaded_classes': len(class_rows),
        'total_teachers': len(teachers),
        'active_teachers': sum(1 for teacher in teachers if not teacher.is_locked),
        'locked_teachers': sum(1 for teacher in teachers if teacher.is_locked),
    }

    return render_template(
        'admin_class_assignments.html',
        classes=class_rows,
        teachers=teacher_options,
        summary=summary,
        filters={'q': q, 'limit': limit},
    )


@app.route('/admin/users')
@role_required(('admin', 'registrar'))
@permission_required()
def admin_user_inventory():
    role_filter = (request.args.get('role') or '').strip().lower()
    query_text = (request.args.get('q') or '').strip()
    limit_raw = (request.args.get('limit') or '200').strip()
    limit = int(limit_raw) if limit_raw.isdigit() else 200
    limit = max(1, min(limit, 1000))

    query = User.query
    if role_filter in ('student', 'teacher', 'assistant', 'registrar', 'admin'):
        query = query.filter(User.role == role_filter)
    if query_text:
        query = query.filter(
            (User.name.ilike(f"%{query_text}%")) |
            (User.email.ilike(f"%{query_text}%"))
        )

    users = query.order_by(User.id.asc()).limit(limit).all()
    user_rows = []
    for user in users:
        teacher_class_count = len(user.teacher_classes) if user.role == 'teacher' else 0
        student_profile_exists = bool(user.students) if user.role == 'student' else False
        advisor, advisor_source, advisor_group = user.students[0].advisor_info() if student_profile_exists else (None, None, None)
        user_rows.append({
            'own_advisor_id': user.students[0].advisor_id if student_profile_exists else None,
            'advisor_name': advisor.name if advisor else None,
            'advisor_source': advisor_source,
            'advisor_group': advisor_group,
            'id': user.id,
            'name': user.name,
            'email': user.email,
            'role': user.role,
            'is_locked': bool(user.is_locked),
            'teacher_class_count': teacher_class_count,
            'student_profile_exists': student_profile_exists,
        })

    summary = {
        'total_loaded': len(user_rows),
        'teacher_count': sum(1 for u in user_rows if u['role'] == 'teacher'),
        'student_count': sum(1 for u in user_rows if u['role'] == 'student'),
        'admin_count': sum(1 for u in user_rows if u['role'] == 'admin'),
    }

    filters = {
        'role': role_filter,
        'q': query_text,
        'limit': limit,
    }
    return render_template(
        'admin_user_inventory.html',
        users=user_rows,
        filters=filters,
        summary=summary,
        role_options=('student', 'teacher', 'assistant', 'registrar', 'admin'),
        current_admin_user_id=session.get('user_id'),
        advisor_options=User.query.filter_by(role='teacher').order_by(User.name.asc()).all(),
        # Student Affairs sees the list read-only and may reset student passwords only.
        can_manage_accounts=session.get('role') == 'admin',
    )


@app.route('/admin/users/<int:user_id>/role', methods=['POST'])
@role_required('admin')
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_update_user_role(user_id):
    target_user = db.session.get(User, user_id)
    if not target_user:
        flash('Hedef kullanıcı bulunamadı.', 'warning')
        return redirect(url_for('admin_user_inventory'))

    new_role = (request.form.get('role') or '').strip().lower()
    actor_user_id = session.get('user_id')
    actor_role = session.get('role')
    old_role = (target_user.role or '').strip().lower()
    admin_count = User.query.filter_by(role='admin').count()

    validation_error = validate_admin_role_update(
        actor_user_id=actor_user_id,
        actor_role=actor_role,
        target_user_id=target_user.id,
        target_user_role=old_role,
        new_role=new_role,
        admin_count=admin_count,
    )
    if validation_error:
        _log_admin_operation(
            action='role_update',
            target_user_id=target_user.id,
            old_value=old_role,
            new_value=new_role,
            status='rejected',
            detail=validation_error,
        )
        flash(validation_error, 'danger')
        return redirect(url_for('admin_user_inventory'))

    if old_role == new_role:
        _log_admin_operation(
            action='role_update',
            target_user_id=target_user.id,
            old_value=old_role,
            new_value=new_role,
            status='noop',
            detail='Role is already set to requested value.',
        )
        flash('Rol zaten bu değere ayarlı.', 'info')
        return redirect(url_for('admin_user_inventory'))

    target_user.role = new_role
    if new_role == 'student' and not target_user.students:
        auto_number = f'AUTO-{target_user.id:05d}'
        name_parts = [part for part in (target_user.name or '').strip().split(' ') if part]
        first_name = name_parts[0] if name_parts else None
        last_name = ' '.join(name_parts[1:]) if len(name_parts) > 1 else None
        db.session.add(
            Student(
                id=target_user.id,
                student_number=auto_number,
                first_name=first_name,
                last_name=last_name,
            )
        )

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        _log_admin_operation(
            action='role_update',
            target_user_id=target_user.id,
            old_value=old_role,
            new_value=new_role,
            status='error',
            detail='Database commit failed.',
        )
        flash('Rol güncelleme başarısız (veritabanı hatası).', 'danger')
        return redirect(url_for('admin_user_inventory'))

    _log_admin_operation(
        action='role_update',
        target_user_id=target_user.id,
        old_value=old_role,
        new_value=new_role,
        status='updated',
        detail='Role updated successfully.',
    )
    flash(f'Role updated: {old_role} -> {new_role}', 'success')
    return redirect(url_for('admin_user_inventory'))


@app.route('/admin/users/<int:user_id>/advisor', methods=['POST'])
@role_required('admin')
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_set_student_advisor(user_id):
    target_user = db.session.get(User, user_id)
    student = target_user.students[0] if target_user and target_user.role == 'student' and target_user.students else None
    if not student:
        flash(_t('advisor_student_not_found'), 'warning')
        return redirect(url_for('admin_user_inventory'))

    raw = _normalize_form_text(request.form.get('advisor_id'), max_length=20)
    advisor = None
    if raw:
        advisor = db.session.get(User, int(raw)) if raw.isdigit() else None
        if not advisor or advisor.role != 'teacher':
            flash(_t('advisor_must_be_teacher'), 'danger')
            return redirect(url_for('admin_user_inventory'))

    old_value = str(student.advisor_id or '')
    student.advisor_id = advisor.id if advisor else None
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash(_t('advisor_update_failed'), 'danger')
        return redirect(url_for('admin_user_inventory'))

    _log_admin_operation(
        action='advisor_update',
        target_user_id=target_user.id,
        old_value=old_value,
        new_value=str(student.advisor_id or ''),
        status='updated',
        detail='Advisor updated.',
    )
    flash(_t('advisor_updated'), 'success')
    return redirect(url_for('admin_user_inventory'))


@app.route('/admin/users/<int:user_id>/lock', methods=['POST'])
@role_required('admin')
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_toggle_user_lock(user_id):
    target_user = db.session.get(User, user_id)
    if not target_user:
        flash('Hedef kullanıcı bulunamadı.', 'warning')
        return redirect(url_for('admin_user_inventory'))

    action = (request.form.get('action') or '').strip().lower()
    if action not in {'lock', 'unlock'}:
        flash('Geçersiz kilit işlemi.', 'danger')
        return redirect(url_for('admin_user_inventory'))

    lock_state = action == 'lock'
    actor_user_id = session.get('user_id')
    actor_role = session.get('role')
    target_locked = bool(target_user.is_locked)
    unlocked_admin_count = User.query.filter_by(role='admin', is_locked=False).count()

    validation_error = validate_admin_lock_update(
        actor_user_id=actor_user_id,
        actor_role=actor_role,
        target_user_id=target_user.id,
        target_user_role=(target_user.role or '').strip().lower(),
        target_user_locked=target_locked,
        lock_state=lock_state,
        unlocked_admin_count=unlocked_admin_count,
    )
    if validation_error:
        _log_admin_operation(
            action='account_lock_toggle',
            target_user_id=target_user.id,
            old_value='locked' if target_locked else 'active',
            new_value='locked' if lock_state else 'active',
            status='rejected',
            detail=validation_error,
        )
        flash(validation_error, 'danger')
        return redirect(url_for('admin_user_inventory'))

    if target_locked == lock_state:
        _log_admin_operation(
            action='account_lock_toggle',
            target_user_id=target_user.id,
            old_value='locked' if target_locked else 'active',
            new_value='locked' if lock_state else 'active',
            status='noop',
            detail='Lock status is already set to requested value.',
        )
        flash('Kilit durumu zaten bu değere ayarlı.', 'info')
        return redirect(url_for('admin_user_inventory'))

    target_user.is_locked = lock_state
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        _log_admin_operation(
            action='account_lock_toggle',
            target_user_id=target_user.id,
            old_value='locked' if target_locked else 'active',
            new_value='locked' if lock_state else 'active',
            status='error',
            detail='Database commit failed.',
        )
        flash('Kilit güncelleme başarısız (veritabanı hatası).', 'danger')
        return redirect(url_for('admin_user_inventory'))

    _log_admin_operation(
        action='account_lock_toggle',
        target_user_id=target_user.id,
        old_value='locked' if target_locked else 'active',
        new_value='locked' if lock_state else 'active',
        status='updated',
        detail='Account lock status updated successfully.',
    )
    flash('Hesap kilit durumu güncellendi.', 'success')
    return redirect(url_for('admin_user_inventory'))


@app.route('/admin/users/<int:user_id>/password', methods=['POST'])
@role_required(('admin', 'registrar'))
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_reset_user_password(user_id):
    target_user = db.session.get(User, user_id)
    if not target_user:
        flash('Hedef kullanıcı bulunamadı.', 'warning')
        return redirect(url_for('admin_user_inventory'))

    actor_user_id = session.get('user_id')
    if actor_user_id == target_user.id:
        flash('Kendi şifrenizi bu panelden sıfırlayamazsınız.', 'danger')
        return redirect(url_for('admin_user_inventory'))

    if session.get('role') == 'registrar' and target_user.role != 'student':
        _log_admin_operation('password_reset', target_user.id, None, None, status='denied', detail='registrar_non_student_target')
        flash(_lang_text('Öğrenci İşleri yalnızca öğrenci hesaplarının şifresini sıfırlayabilir.',
                         'Student Affairs can only reset passwords of student accounts.'), 'danger')
        return redirect(url_for('admin_user_inventory'))

    new_password = request.form.get('new_password') or ''
    confirm_password = request.form.get('confirm_password') or ''
    if not new_password or not confirm_password:
        flash('Her iki şifre alanı da zorunludur.', 'warning')
        return redirect(url_for('admin_user_inventory'))
    if new_password != confirm_password:
        flash('Şifre onayı eşleşmiyor.', 'warning')
        return redirect(url_for('admin_user_inventory'))

    policy_error = validate_password_policy(new_password, user_name=target_user.name, user_email=target_user.email)
    if policy_error:
        _log_admin_operation(
            action='password_reset',
            target_user_id=target_user.id,
            old_value='set',
            new_value='set',
            status='rejected',
            detail=policy_error,
        )
        flash(policy_error, 'danger')
        return redirect(url_for('admin_user_inventory'))

    target_user.set_password(new_password)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        _log_admin_operation(
            action='password_reset',
            target_user_id=target_user.id,
            old_value='set',
            new_value='set',
            status='error',
            detail='Database commit failed.',
        )
        flash('Şifre sıfırlama başarısız (veritabanı hatası).', 'danger')
        return redirect(url_for('admin_user_inventory'))

    _log_admin_operation(
        action='password_reset',
        target_user_id=target_user.id,
        old_value='set',
        new_value='set',
        status='updated',
        detail='Password reset by admin.',
    )
    flash(f'Kullanıcı #{target_user.id} için şifre sıfırlandı.', 'success')
    return redirect(url_for('admin_user_inventory'))


@app.route('/admin/classes/<int:class_id>/assign-teacher', methods=['POST'])
@role_required('admin')
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_assign_class_teacher(class_id):
    class_obj = db.session.get(Class, class_id)
    teacher_id_raw = (request.form.get('teacher_id') or '').strip()
    if not teacher_id_raw.isdigit():
        flash('Geçersiz öğretmen kimliği.', 'danger')
        return redirect(url_for('admin_class_assignments'))

    target_teacher_id = int(teacher_id_raw)
    teacher_user = db.session.get(User, target_teacher_id)
    current_teacher_id = class_obj.teacher_id if class_obj else None

    validation_error = validate_admin_teacher_assignment(
        actor_role=session.get('role'),
        class_obj=class_obj,
        teacher_user=teacher_user,
        current_teacher_id=current_teacher_id,
        target_teacher_id=target_teacher_id,
    )
    if validation_error:
        _log_admin_operation(
            action='class_teacher_assign',
            target_user_id=target_teacher_id if teacher_user else session.get('user_id'),
            old_value=str(current_teacher_id) if current_teacher_id is not None else None,
            new_value=str(target_teacher_id),
            status='rejected',
            detail=validation_error,
        )
        flash(validation_error, 'danger')
        return redirect(url_for('admin_class_assignments'))

    class_obj.teacher_id = target_teacher_id
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        _log_admin_operation(
            action='class_teacher_assign',
            target_user_id=target_teacher_id,
            old_value=str(current_teacher_id),
            new_value=str(target_teacher_id),
            status='error',
            detail=f'Database commit failed for class_id={class_obj.id}.',
        )
        flash('Öğretmen ataması başarısız (veritabanı hatası).', 'danger')
        return redirect(url_for('admin_class_assignments'))

    _log_admin_operation(
        action='class_teacher_assign',
        target_user_id=target_teacher_id,
        old_value=str(current_teacher_id),
        new_value=str(target_teacher_id),
        status='updated',
        detail=f'Class teacher updated for class_id={class_obj.id}.',
    )
    flash('Sınıf öğretmeni güncellendi.', 'success')
    return redirect(url_for('admin_class_assignments'))


@app.route('/admin/security/permission-audit')
@role_required('admin')
@permission_required()
def admin_permission_audit_report():
    limit_raw = (request.args.get('limit') or '100').strip()
    limit = int(limit_raw) if limit_raw.isdigit() else 100
    limit = max(1, min(limit, MAX_PERMISSION_AUDIT_EVENTS))

    role_filter = (request.args.get('role') or '').strip().lower()
    endpoint_filter = (request.args.get('endpoint') or '').strip().lower()
    permission_filter = (request.args.get('permission') or '').strip().lower()
    from_raw = (request.args.get('from') or '').strip()
    to_raw = (request.args.get('to') or '').strip()

    from_dt = _parse_filter_datetime(from_raw)
    to_dt = _parse_filter_datetime(to_raw)

    events = []
    source = 'database'
    try:
        query = PermissionAuditLog.query
        if role_filter:
            query = query.filter(PermissionAuditLog.role.ilike(f"%{role_filter}%"))
        if endpoint_filter:
            query = query.filter(PermissionAuditLog.endpoint.ilike(f"%{endpoint_filter}%"))
        if permission_filter:
            query = query.filter(PermissionAuditLog.permission.ilike(f"%{permission_filter}%"))
        if from_dt:
            query = query.filter(PermissionAuditLog.created_at >= from_dt)
        if to_dt:
            query = query.filter(PermissionAuditLog.created_at <= to_dt)

        rows = query.order_by(PermissionAuditLog.created_at.desc()).limit(limit).all()
        events = [
            {
                'timestamp': f"{row.created_at.isoformat()}Z" if row.created_at else '-',
                'user_id': row.user_id,
                'role': row.role,
                'endpoint': row.endpoint,
                'permission': row.permission,
                'method': row.method,
                'path': row.path,
                'ip': row.ip,
            }
            for row in rows
        ]
    except Exception:
        db.session.rollback()
        source = 'buffer'
        buffer_events = list(PERMISSION_AUDIT_EVENTS)

        def _buffer_matches(event):
            event_role = (event.get('role') or '').lower()
            event_endpoint = (event.get('endpoint') or '').lower()
            event_permission = (event.get('permission') or '').lower()

            if role_filter and role_filter not in event_role:
                return False
            if endpoint_filter and endpoint_filter not in event_endpoint:
                return False
            if permission_filter and permission_filter not in event_permission:
                return False

            timestamp_raw = (event.get('timestamp') or '').replace('Z', '')
            event_time = None
            try:
                event_time = datetime.datetime.fromisoformat(timestamp_raw)
            except ValueError:
                event_time = None

            if from_dt and event_time and event_time < from_dt:
                return False
            if to_dt and event_time and event_time > to_dt:
                return False
            return True

        events = [event for event in buffer_events if _buffer_matches(event)][:limit]

    filters = {
        'role': role_filter,
        'endpoint': endpoint_filter,
        'permission': permission_filter,
        'from': from_raw,
        'to': to_raw,
    }
    return render_template('admin_permission_audit.html', events=events, limit=limit, source=source, filters=filters)


@app.route('/admin/security/admin-operations')
@role_required('admin')
@permission_required()
def admin_operation_audit_report():
    limit_raw = (request.args.get('limit') or '100').strip()
    limit = int(limit_raw) if limit_raw.isdigit() else 100
    limit = max(1, min(limit, 500))

    action_filter = (request.args.get('action') or '').strip().lower()
    status_filter = (request.args.get('status') or '').strip().lower()
    actor_filter_raw = (request.args.get('actor_user_id') or '').strip()
    target_filter_raw = (request.args.get('target_user_id') or '').strip()
    from_raw = (request.args.get('from') or '').strip()
    to_raw = (request.args.get('to') or '').strip()

    actor_user_id = _parse_optional_int(actor_filter_raw)
    target_user_id = _parse_optional_int(target_filter_raw)
    from_dt = _parse_filter_datetime(from_raw)
    to_dt = _parse_filter_datetime(to_raw)

    query = AdminOperationLog.query
    if action_filter:
        query = query.filter(AdminOperationLog.action.ilike(f"%{action_filter}%"))
    if status_filter:
        query = query.filter(AdminOperationLog.status.ilike(f"%{status_filter}%"))
    if actor_user_id is not None:
        query = query.filter(AdminOperationLog.actor_user_id == actor_user_id)
    if target_user_id is not None:
        query = query.filter(AdminOperationLog.target_user_id == target_user_id)
    if from_dt:
        query = query.filter(AdminOperationLog.created_at >= from_dt)
    if to_dt:
        query = query.filter(AdminOperationLog.created_at <= to_dt)

    rows = query.order_by(AdminOperationLog.created_at.desc()).limit(limit).all()
    events = []
    for row in rows:
        actor_user = db.session.get(User, row.actor_user_id)
        target_user = db.session.get(User, row.target_user_id)
        events.append(
            {
                'timestamp': f"{row.created_at.isoformat()}Z" if row.created_at else '-',
                'action': row.action,
                'status': row.status,
                'actor_user_id': row.actor_user_id,
                'actor_user_email': actor_user.email if actor_user else '-',
                'target_user_id': row.target_user_id,
                'target_user_email': target_user.email if target_user else '-',
                'old_value': row.old_value,
                'new_value': row.new_value,
                'detail': row.detail,
                'ip': row.ip,
            }
        )

    filters = {
        'action': action_filter,
        'status': status_filter,
        'actor_user_id': actor_filter_raw,
        'target_user_id': target_filter_raw,
        'from': from_raw,
        'to': to_raw,
    }
    summary = {
        'loaded_events': len(events),
        'updated_count': sum(1 for event in events if event['status'] == 'updated'),
        'rejected_count': sum(1 for event in events if event['status'] == 'rejected'),
        'error_count': sum(1 for event in events if event['status'] == 'error'),
    }
    return render_template('admin_operation_audit.html', events=events, limit=limit, filters=filters, summary=summary)


@app.route('/health')
def health_check():
    snapshot = _build_health_snapshot()
    status_code = 200 if snapshot['status'] in {'healthy', 'degraded'} else 503
    return jsonify(snapshot), status_code


@app.route('/admin/security/health-status')
@role_required('admin')
@permission_required(PERMISSIONS['ADMIN_METRICS_READ'])
def admin_health_status_report():
    snapshot = _build_health_snapshot()
    return render_template('admin_health_status.html', health=snapshot)


@app.route('/admin/security/permission-matrix')
@role_required('admin')
@permission_required()
def admin_permission_matrix():
    all_permissions = sorted(set(PERMISSIONS.values()))
    matrix_rows = []
    for role in sorted(ROLE_PERMISSIONS.keys()):
        assigned = set(ROLE_PERMISSIONS.get(role, set()))
        matrix_rows.append(
            {
                'role': role,
                'permission_count': len(assigned),
                'permissions': [
                    {
                        'name': perm,
                        'enabled': perm in assigned,
                    }
                    for perm in all_permissions
                ],
            }
        )

    return render_template(
        'admin_permission_matrix.html',
        rows=matrix_rows,
        permissions=all_permissions,
        endpoint_map=sorted(PERMISSION_MAP.items()),
    )


@app.route('/admin/notifications', methods=['GET', 'POST'])
@role_required(('admin', 'registrar'))
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_notifications():
    """Compose announcements and/or e-mails; edit or delete an existing announcement (?edit=<id>)."""
    admin_user = db.session.get(User, session.get('user_id'))
    roles = tuple(ROLE_DASHBOARDS)
    role_counts = {role: User.query.filter_by(role=role).count() for role in roles}
    role_counts['all'] = sum(role_counts.values())
    invalid_emails = [u.email for u in User.query.all() if not _is_valid_email((u.email or '').strip().lower())]
    stats = {
        'role_counts': role_counts,
        'invalid_email_count': len(invalid_emails),
        'sample_invalid_emails': invalid_emails[:10],
        'dry_run': bool(app.config.get('EMAIL_DRY_RUN', True)),
    }
    courses = Course.query.order_by(Course.code.asc()).limit(500).all()
    groups = StudentGroup.query.order_by(StudentGroup.code.asc()).all()

    def _can_modify(announcement):
        # Admins manage every announcement; Student Affairs only the ones they wrote.
        return session.get('role') == 'admin' or announcement.author_id == session.get('user_id')

    not_yours = _lang_text('Yalnızca kendi gönderdiğiniz duyuruları değiştirebilirsiniz.',
                           'You can only change announcements you sent.')

    edit_announcement = None
    edit_id = request.args.get('edit', type=int)
    if edit_id:
        edit_announcement = db.session.get(Announcement, edit_id)
        if not edit_announcement:
            flash(_lang_text('Duyuru bulunamadı.', 'Announcement not found.'), 'warning')
        elif not _can_modify(edit_announcement):
            flash(not_yours, 'danger')
            edit_announcement = None

    def _render():
        people = User.query.filter(User.id != session.get('user_id')).order_by(User.name.asc()).all()
        return render_template('admin_notifications.html', stats=stats, courses=courses, groups=groups,
                               people=people, edit_announcement=edit_announcement)

    if request.method != 'POST':
        return _render()

    def _parse_dt_local(raw_value: str):
        raw_value = (raw_value or '').strip()
        if not raw_value:
            return None
        try:
            return datetime.datetime.fromisoformat(raw_value)
        except ValueError:
            return None

    def _announcement_payload():
        target_role = _normalize_form_text(request.form.get('announcement_target_role'), max_length=20).lower()
        if target_role not in ANNOUNCEMENT_TARGETS:
            target_role = 'all'
        course_id_raw = _normalize_form_text(request.form.get('announcement_course_id'), max_length=20)
        course = db.session.get(Course, int(course_id_raw)) if course_id_raw.isdigit() else None
        group_id_raw = _normalize_form_text(request.form.get('announcement_group_id'), max_length=20)
        group = db.session.get(StudentGroup, int(group_id_raw)) if group_id_raw.isdigit() else None
        return {
            'group': group,
            'title': _normalize_form_text(request.form.get('announcement_title'), max_length=150),
            'body': _normalize_form_text(request.form.get('announcement_body'), max_length=2000),
            'target_role': target_role,
            'course': course,
            'starts_at': _parse_dt_local(request.form.get('announcement_starts_at') or ''),
            'ends_at': _parse_dt_local(request.form.get('announcement_ends_at') or ''),
        }

    def _payload_error(payload):
        if not payload['title'] or not payload['body']:
            return _lang_text('Duyuru başlığı ve içeriği zorunludur.', 'Announcement title and body are required.')
        if payload['ends_at'] and payload['starts_at'] and payload['ends_at'] < payload['starts_at']:
            return _lang_text('Bitiş zamanı başlangıç zamanından önce olamaz.', 'End time cannot be earlier than start time.')
        return None

    def _max_recipients():
        raw = _normalize_form_text(request.form.get('max_recipients'), max_length=10)
        return max(1, min(int(raw) if raw.isdigit() else 200, 2000))

    def _email_recipients(target_role, course=None, group=None):
        query = User.query
        if target_role != 'all':
            query = query.filter_by(role=target_role)
        users = query.order_by(User.id.asc()).all()
        if group is not None:
            members = {student.id for student in group.students}
            users = [u for u in users if u.id in members]
        if course is not None:
            enrolled = {e.student_id for e in course.enrollments if e.approved}
            users = [u for u in users if u.id in enrolled or u.id == course.teacher_id]
        return users

    def _send_emails(users, subject, body):
        sent = failed = invalid = 0
        for user in users:
            recipient = (user.email or '').strip().lower()
            if not _is_valid_email(recipient):
                invalid += 1
                continue
            ok, _ = _send_email_notification(recipient, subject, body)
            if ok:
                sent += 1
            else:
                failed += 1
        flash(
            _lang_text(
                f'E-posta gönderimi tamamlandı. Gönderilen: {sent}, başarısız: {failed}, geçersiz adres: {invalid}, hedef: {len(users)}.',
                f'Email delivery completed. Sent={sent}, Failed={failed}, Invalid Emails={invalid}, Targeted={len(users)}.',
            ),
            'success' if failed == 0 else 'warning',
        )

    def _publish(payload):
        try:
            announcement = Announcement(
                title=payload['title'],
                body=payload['body'],
                author_id=admin_user.id if admin_user else session.get('user_id'),
                target_role=payload['target_role'],
                course_id=payload['course'].id if payload['course'] else None,
                group_id=payload['group'].id if payload['group'] else None,
                starts_at=payload['starts_at'],
                ends_at=payload['ends_at'],
            )
            db.session.add(announcement)
            db.session.commit()
            flash(_lang_text(f'Duyuru yayınlandı: {payload["title"]}', f'Announcement published: {payload["title"]}'), 'success')
            return announcement
        except Exception:
            db.session.rollback()
            flash(_lang_text('Duyuru yayınlanırken hata oluştu.', 'Failed to publish announcement.'), 'danger')
            return None

    def _back_to_inbox(announcement_id=None):
        return redirect(url_for('announcements_inbox', box='sent', id=f'a{announcement_id}') if announcement_id
                        else url_for('announcements_inbox', box='sent'))

    action = _normalize_form_text(request.form.get('action'), max_length=30).lower()
    return_to_inbox = request.form.get('return_to') == 'inbox'

    if action == 'compose' and request.form.get('announcement_target_role') == 'person':
        # Personal message to one user: in-app message, e-mail, or both.
        channel = _normalize_form_text(request.form.get('channel'), max_length=20).lower()
        if channel not in {'announcement', 'both', 'email'}:
            channel = 'announcement'
        recipient = _resolve_recipient(request.form.get('recipient'))
        title = _normalize_form_text(request.form.get('announcement_title'), max_length=150)
        body = _normalize_form_text(request.form.get('announcement_body'), max_length=2000)
        if recipient is None:
            flash(_lang_text('Kişi bulunamadı. Listeden bir kişi seçin.', 'Person not found. Pick someone from the list.'), 'warning')
            return _render()
        if not title or not body:
            flash(_lang_text('Konu ve içerik zorunludur.', 'Subject and message are required.'), 'warning')
            return _render()
        message = None
        if channel in {'announcement', 'both'}:
            message = Message(sender_id=session.get('user_id'), recipient_id=recipient.id, subject=title, body=body)
            db.session.add(message)
            db.session.commit()
            flash(_lang_text(f'Mesaj gönderildi: {recipient.name}', f'Message sent to {recipient.name}'), 'success')
        if channel in {'email', 'both'}:
            _send_emails([recipient], title, body)
        if message is not None:
            return redirect(url_for('announcements_inbox', box='sent', id=f'm{message.id}'))
        return _render()

    if action in {'compose', 'publish_announcement'}:
        channel = _normalize_form_text(request.form.get('channel'), max_length=20).lower() if action == 'compose' else 'announcement'
        if channel not in {'announcement', 'both', 'email'}:
            channel = 'announcement'
        payload = _announcement_payload()
        error = _payload_error(payload)
        if error:
            flash(error, 'warning')
            return _render()
        announcement = _publish(payload) if channel in {'announcement', 'both'} else None
        if channel in {'email', 'both'}:
            recipients = _email_recipients(payload['target_role'], payload['course'], payload['group'])[:_max_recipients()]
            _send_emails(recipients, payload['title'], payload['body'])
        if action == 'compose' and announcement is not None:
            return _back_to_inbox(announcement.id)
        return _render()

    if action == 'send_broadcast':
        target_role = _normalize_form_text(request.form.get('target_role'), max_length=20).lower()
        if target_role not in ANNOUNCEMENT_TARGETS:
            target_role = 'all'
        subject = _normalize_form_text(request.form.get('subject'), max_length=150)
        body = _normalize_form_text(request.form.get('body'), max_length=2000)
        if not subject or not body:
            flash(_lang_text('Toplu gönderim için konu ve içerik zorunludur.', 'Subject and body are required for a broadcast.'), 'warning')
            return _render()
        _send_emails(_email_recipients(target_role)[:_max_recipients()], subject, body)
        return _render()

    if action in {'update_announcement', 'delete_announcement'}:
        announcement_id_raw = _normalize_form_text(request.form.get('announcement_id'), max_length=20)
        announcement = db.session.get(Announcement, int(announcement_id_raw)) if announcement_id_raw.isdigit() else None
        if not announcement:
            flash(_lang_text('Duyuru bulunamadı.', 'Announcement not found.'), 'warning')
            return _back_to_inbox() if return_to_inbox else _render()
        if not _can_modify(announcement):
            flash(not_yours, 'danger')
            return _back_to_inbox() if return_to_inbox else _render()

        if action == 'delete_announcement':
            try:
                db.session.delete(announcement)
                db.session.commit()
                flash(_lang_text('Duyuru silindi.', 'Announcement deleted.'), 'success')
            except Exception:
                db.session.rollback()
                flash(_lang_text('Duyuru silinirken hata oluştu.', 'Failed to delete announcement.'), 'danger')
            return _back_to_inbox() if return_to_inbox else _render()

        payload = _announcement_payload()
        error = _payload_error(payload)
        if error:
            flash(error, 'warning')
            edit_announcement = announcement
            return _render()
        try:
            announcement.title = payload['title']
            announcement.body = payload['body']
            announcement.target_role = payload['target_role']
            announcement.course_id = payload['course'].id if payload['course'] else None
            announcement.group_id = payload['group'].id if payload['group'] else None
            announcement.starts_at = payload['starts_at']
            announcement.ends_at = payload['ends_at']
            db.session.commit()
            flash(_lang_text(f'Duyuru güncellendi: {payload["title"]}', f'Announcement updated: {payload["title"]}'), 'success')
        except Exception:
            db.session.rollback()
            flash(_lang_text('Duyuru güncellenirken hata oluştu.', 'Failed to update announcement.'), 'danger')
        return _back_to_inbox(announcement.id) if return_to_inbox else _render()

    flash(_lang_text('Geçersiz duyuru işlemi.', 'Invalid announcement action.'), 'warning')
    return _render()


GROUP_IMPORT_MAX_BYTES = 2 * 1024 * 1024
GROUP_IMPORT_MAX_ROWS = 5000
GROUP_IMPORT_HEADERS = {
    'öğrenci no', 'ogrenci no', 'öğrenci numarası', 'ogrenci numarasi', 'numara', 'no', 'student number',
    'student_number', 'student no', 'e-posta', 'eposta', 'email', 'e-mail', 'mail',
}


def _read_group_import_values(file_storage):
    """Return (values, error) from an uploaded .xlsx/.csv: one student number or e-mail per row.

    Uses the column whose header looks like a student number / e-mail; without a header row,
    the first non-empty cell of each row.
    """
    filename = (file_storage.filename or '').lower()
    data = file_storage.read(GROUP_IMPORT_MAX_BYTES + 1)
    if not data:
        return [], _lang_text('Dosya boş.', 'The file is empty.')
    if len(data) > GROUP_IMPORT_MAX_BYTES:
        return [], _lang_text('Dosya 2 MB sınırını aşıyor.', 'The file exceeds the 2 MB limit.')

    rows = []
    if filename.endswith(('.xlsx', '.xlsm')):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            for row in workbook.active.iter_rows(values_only=True):
                rows.append(row)
                if len(rows) > GROUP_IMPORT_MAX_ROWS:
                    break
            workbook.close()
        except Exception:
            return [], _lang_text('Excel dosyası okunamadı.', 'Could not read the Excel file.')
    elif filename.endswith(('.csv', '.txt')):
        try:
            text_data = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text_data = data.decode('cp1254', errors='replace')
        delimiter = ';' if text_data.count(';') > text_data.count(',') else ','
        rows = list(csv.reader(StringIO(text_data), delimiter=delimiter))[:GROUP_IMPORT_MAX_ROWS + 1]
    else:
        return [], _lang_text('Yalnızca .xlsx veya .csv dosyası yükleyin.', 'Upload an .xlsx or .csv file.')

    def _cell(value):
        if value is None:
            return ''
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    rows = [[_cell(c) for c in row] for row in rows if row and any(_cell(c) for c in row)]
    if not rows:
        return [], _lang_text('Dosyada veri bulunamadı.', 'No data found in the file.')

    header = [c.lower() for c in rows[0]]
    columns = [i for i, name in enumerate(header) if name in GROUP_IMPORT_HEADERS]
    if columns:
        values = [next((row[i] for i in columns if i < len(row) and row[i]), '') for row in rows[1:]]
    else:
        values = [next((c for c in row if c), '') for row in rows]
    return [v for v in values if v][:GROUP_IMPORT_MAX_ROWS], None


def _find_students_by_identifiers(values):
    """Map student numbers or e-mails to Student rows. Returns (found_students, not_found_values)."""
    numbers = {v.upper() for v in values if '@' not in v}
    emails = {v.lower() for v in values if '@' in v}
    by_number = {}
    if numbers:
        for student in Student.query.filter(func.upper(Student.student_number).in_(numbers)).all():
            by_number[(student.student_number or '').upper()] = student
    by_email = {}
    if emails:
        rows = db.session.query(Student, User).join(User, User.id == Student.id).filter(func.lower(User.email).in_(emails)).all()
        by_email = {user.email.lower(): student for student, user in rows}
    found, not_found, seen = [], [], set()
    for value in values:
        student = by_email.get(value.lower()) if '@' in value else by_number.get(value.upper())
        if student is None:
            not_found.append(value)
        elif student.id not in seen:
            seen.add(student.id)
            found.append(student)
    return found, not_found


@app.route('/admin/groups', methods=['GET', 'POST'])
@role_required(('admin', 'registrar'))
@permission_required()
@csrf_protect
@rate_limit_protect('admin_mutation', 'ADMIN_MUTATION_RATE_LIMIT_MAX', 'ADMIN_MUTATION_RATE_LIMIT_WINDOW_SECONDS', _admin_mutation_rate_key)
def admin_groups():
    """Department groups (CMP, SWE, ...): create, rename, delete, manage members, import from Excel/CSV."""
    actor_id = session.get('user_id')

    def _group_from_form():
        raw = _normalize_form_text(request.form.get('group_id'), max_length=20)
        return db.session.get(StudentGroup, int(raw)) if raw.isdigit() else None

    def _clean_code(raw):
        return re.sub(r'[^A-Z0-9_-]', '', (raw or '').strip().upper().replace('İ', 'I'))[:20]

    if request.method == 'POST':
        action = _normalize_form_text(request.form.get('action'), max_length=30).lower()

        if action == 'create_group':
            code = _clean_code(request.form.get('code'))
            name = _normalize_form_text(request.form.get('name'), max_length=120)
            if not code or not name:
                flash(_lang_text('Grup kodu ve adı zorunludur.', 'Group code and name are required.'), 'warning')
                return redirect(url_for('admin_groups'))
            if StudentGroup.query.filter_by(code=code).first():
                flash(_lang_text(f'{code} kodlu bir grup zaten var.', f'A group with code {code} already exists.'), 'warning')
                return redirect(url_for('admin_groups'))
            group = StudentGroup(code=code, name=name)
            db.session.add(group)
            db.session.commit()
            _log_admin_operation('group_create', actor_id, None, code)
            flash(_lang_text(f'{code} grubu oluşturuldu.', f'Group {code} created.'), 'success')
            return redirect(url_for('admin_groups', id=group.id))

        group = _group_from_form()
        if group is None:
            flash(_lang_text('Grup bulunamadı.', 'Group not found.'), 'warning')
            return redirect(url_for('admin_groups'))

        if action == 'update_group':
            code = _clean_code(request.form.get('code'))
            name = _normalize_form_text(request.form.get('name'), max_length=120)
            if not code or not name:
                flash(_lang_text('Grup kodu ve adı zorunludur.', 'Group code and name are required.'), 'warning')
            elif StudentGroup.query.filter(StudentGroup.code == code, StudentGroup.id != group.id).first():
                flash(_lang_text(f'{code} kodlu bir grup zaten var.', f'A group with code {code} already exists.'), 'warning')
            else:
                old_code = group.code
                group.code, group.name = code, name
                db.session.commit()
                _log_admin_operation('group_update', actor_id, old_code, code)
                flash(_lang_text('Grup güncellendi.', 'Group updated.'), 'success')
            return redirect(url_for('admin_groups', id=group.id))

        if action == 'set_advisor':
            raw = _normalize_form_text(request.form.get('advisor_id'), max_length=20)
            advisor = db.session.get(User, int(raw)) if raw.isdigit() else None
            if raw and (advisor is None or advisor.role != 'teacher'):
                flash(_t('advisor_must_be_teacher'), 'danger')
                return redirect(url_for('admin_groups', id=group.id))
            old_value = str(group.advisor_id or '')
            group.advisor_id = advisor.id if advisor else None
            db.session.commit()
            _log_admin_operation('group_advisor_update', actor_id, old_value, str(group.advisor_id or ''), detail=group.code)
            flash(_t('advisor_updated'), 'success')
            return redirect(url_for('admin_groups', id=group.id))

        if action == 'delete_group':
            code = group.code
            Announcement.query.filter_by(group_id=group.id).update({'group_id': None})
            group.students = []
            db.session.delete(group)
            db.session.commit()
            _log_admin_operation('group_delete', actor_id, code, None)
            flash(_lang_text(f'{code} grubu silindi. Bu gruba gönderilmiş duyurular genel duyuruya dönüştü.',
                             f'Group {code} deleted. Its announcements are now general announcements.'), 'success')
            return redirect(url_for('admin_groups'))

        if action == 'add_member':
            raw = _normalize_form_text(request.form.get('student'), max_length=200)
            identifier = raw.split()[0] if raw else ''
            found, _ = _find_students_by_identifiers([identifier]) if identifier else ([], [])
            if not found:
                flash(_lang_text('Öğrenci bulunamadı. Öğrenci numarası veya e-posta girin.',
                                 'Student not found. Enter a student number or e-mail.'), 'warning')
            elif found[0] in group.students:
                flash(_lang_text('Öğrenci zaten bu grupta.', 'The student is already in this group.'), 'info')
            else:
                group.students.append(found[0])
                db.session.commit()
                _log_admin_operation('group_member_add', found[0].id, None, group.code)
                flash(_lang_text('Öğrenci gruba eklendi.', 'Student added to the group.'), 'success')
            return redirect(url_for('admin_groups', id=group.id))

        if action == 'remove_member':
            raw = _normalize_form_text(request.form.get('student_id'), max_length=20)
            student = db.session.get(Student, int(raw)) if raw.isdigit() else None
            if student is not None and student in group.students:
                group.students.remove(student)
                db.session.commit()
                _log_admin_operation('group_member_remove', student.id, group.code, None)
                flash(_lang_text('Öğrenci gruptan çıkarıldı.', 'Student removed from the group.'), 'success')
            return redirect(url_for('admin_groups', id=group.id))

        if action == 'import_members':
            upload = request.files.get('file')
            if upload is None or not upload.filename:
                flash(_lang_text('Bir dosya seçin.', 'Choose a file.'), 'warning')
                return redirect(url_for('admin_groups', id=group.id))
            values, error = _read_group_import_values(upload)
            if error:
                flash(error, 'warning')
                return redirect(url_for('admin_groups', id=group.id))
            found, not_found = _find_students_by_identifiers(values)
            existing = {s.id for s in group.students}
            added = [s for s in found if s.id not in existing]
            group.students.extend(added)
            db.session.commit()
            _log_admin_operation('group_import', actor_id, None, group.code,
                                 detail=f'added={len(added)} existing={len(found) - len(added)} not_found={len(not_found)}')
            message = _lang_text(
                f'İçe aktarma tamamlandı: {len(added)} öğrenci eklendi, {len(found) - len(added)} zaten üyeydi, {len(not_found)} kayıt bulunamadı.',
                f'Import finished: {len(added)} added, {len(found) - len(added)} already members, {len(not_found)} not found.',
            )
            if not_found:
                message += ' ' + _lang_text('Bulunamayanlar', 'Not found') + ': ' + ', '.join(not_found[:15]) + ('…' if len(not_found) > 15 else '')
            flash(message, 'success' if not not_found else 'warning')
            return redirect(url_for('admin_groups', id=group.id))

        flash(_lang_text('Geçersiz grup işlemi.', 'Invalid group action.'), 'warning')
        return redirect(url_for('admin_groups', id=group.id))

    groups = StudentGroup.query.order_by(StudentGroup.code.asc()).all()
    selected_id = request.args.get('id', type=int)
    selected = next((g for g in groups if g.id == selected_id), None) or (groups[0] if groups else None)
    members, candidates = [], []
    if selected is not None:
        member_ids = {s.id for s in selected.students}
        members = sorted(selected.students, key=lambda s: (s.student_number or ''))
        candidates = [s for s in Student.query.order_by(Student.student_number.asc()).all() if s.id not in member_ids]
    teachers = User.query.filter_by(role='teacher').order_by(User.name.asc()).all()
    return render_template('admin_groups.html', groups=groups, selected=selected, members=members, candidates=candidates, advisor_options=teachers)

# ------------------ TEACHER DASHBOARD ------------------
def _build_teacher_dashboard_context(user: User):
    _auto_finalize_expired_sessions()

    teacher_classes = _accessible_classes(user)

    class_ids = [cls.id for cls in teacher_classes]
    active_session_by_class = {}
    confirmed_sessions_by_class = defaultdict(list)
    student_counts = {}
    records_by_session = defaultdict(list)
    live_present_counts = {}

    if class_ids:
        active_sessions = (
            AttendanceSession.query
            .filter(AttendanceSession.class_id.in_(class_ids), AttendanceSession.active.is_(True))
            .order_by(AttendanceSession.class_id.asc(), AttendanceSession.date.desc())
            .all()
        )
        for sess in active_sessions:
            active_session_by_class.setdefault(sess.class_id, sess)

        confirmed_sessions = (
            AttendanceSession.query
            .filter(AttendanceSession.class_id.in_(class_ids), AttendanceSession.confirmed.is_(True))
            .order_by(AttendanceSession.class_id.asc(), AttendanceSession.date.desc())
            .all()
        )
        for sess in confirmed_sessions:
            confirmed_sessions_by_class[sess.class_id].append(sess)

        student_count_rows = (
            db.session.query(Class.id, func.count(Student.id))
            .outerjoin(Class.students)
            .filter(Class.id.in_(class_ids))
            .group_by(Class.id)
            .all()
        )
        student_counts = {class_id: count for class_id, count in student_count_rows}

        active_session_ids = [sess.id for sess in active_session_by_class.values()]
        if active_session_ids:
            present_rows = (
                db.session.query(AttendanceRecord.session_id, func.count(AttendanceRecord.id))
                .filter(AttendanceRecord.session_id.in_(active_session_ids), AttendanceRecord.present.is_(True))
                .group_by(AttendanceRecord.session_id)
                .all()
            )
            live_present_counts = {session_id: count for session_id, count in present_rows}

            session_records = AttendanceRecord.query.filter(AttendanceRecord.session_id.in_(active_session_ids)).all()
            for record in session_records:
                records_by_session[record.session_id].append(record)

    for cls in teacher_classes:
        active_session = active_session_by_class.get(cls.id)
        cls.active_session = active_session
        cls.live_total_count = student_counts.get(cls.id, 0)
        if active_session:
            cls.attendance_by_student = {r.student_id: r for r in records_by_session.get(active_session.id, [])}
            cls.live_present_count = live_present_counts.get(active_session.id, 0)
            cls.deadline = _session_deadline(active_session)
        else:
            cls.attendance_by_student = {}
            cls.live_present_count = 0
            cls.deadline = None

        cls.confirmed_sessions = confirmed_sessions_by_class.get(cls.id, [])
        cls.last_confirmed = cls.confirmed_sessions[0] if cls.confirmed_sessions else None

    dashboard_summary = {
        'class_count': len(teacher_classes),
        'active_session_count': sum(1 for cls in teacher_classes if cls.active_session),
        'student_total': sum(len(cls.students) for cls in teacher_classes),
        'confirmed_session_total': sum(len(cls.confirmed_sessions) for cls in teacher_classes),
    }
    return teacher_classes, dashboard_summary


@app.route('/teacher_dashboard', methods=['GET', 'POST'])
@role_required('teacher')
@permission_required()
def teacher_dashboard():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    if request.method == 'POST':
        return _handle_user_calendar_dashboard_post(user, PERMISSIONS['TEACHER_DASHBOARD_CALENDAR_UPDATE'], 'teacher_dashboard')

    teacher_classes, _ = _build_teacher_dashboard_context(user)
    tp = _teacher_profile_summary(user, teacher_classes)
    profile = {
        'name': tp['name'],
        'title': tp['title'],
        'initials': tp['initials'],
        'photo_url': url_for('teacher_photo', user_id=user.id) if tp['has_photo'] else None,
        'facts': [
            (_lang_text('E-posta', 'E-mail'), tp['email']),
            (_lang_text('Fakülte', 'Faculty'), tp['faculty']),
            (_lang_text('Bölüm', 'Department'), tp['department']),
            (_lang_text('Ofis', 'Office'), tp['office']),
            (_lang_text('Ofis saatleri', 'Office hours'), tp['office_hours']),
        ],
        'stats': [
            (tp['class_count'], _lang_text('Ders', 'Classes')),
            (tp['student_count'], _lang_text('Öğrenci', 'Students')),
            (tp['advisee_count'], _lang_text('Danışan', 'Advisees')),
            (tp['assistant_count'], _lang_text('Asistan', 'Assistants')),
        ],
        'edit_url': url_for('teacher_account'),
    }
    _, upcoming, events = _serialize_user_calendar_events(user.id)
    return _render_home_dashboard(user, 'teacher', 'teacher_dashboard', _t('teacher_dashboard'), profile, events, len(upcoming))


@app.route('/teacher/attendance-hub', methods=['GET'])
@role_required(('teacher', 'assistant'))
@permission_required()
def teacher_attendance_hub():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    teacher_classes, dashboard_summary = _build_teacher_dashboard_context(user)
    assistant_options = (User.query.filter_by(role='assistant', is_locked=False).order_by(User.name.asc()).all()
                         if session.get('role') == 'teacher' else [])

    return render_template(
        'teacher_attendance_hub.html',
        teacher_classes=teacher_classes,
        user_name=user.name,
        dashboard_summary=dashboard_summary,
        assistant_options=assistant_options,
    )


@app.route('/teacher/grade-entry', methods=['GET', 'POST'])
@role_required(('teacher', 'assistant'))
@permission_required()
def teacher_grade_entry():
    def _letter_from_ratio(ratio_percent: float) -> str:
        if ratio_percent >= 90:
            return 'AA'
        if ratio_percent >= 85:
            return 'BA'
        if ratio_percent >= 80:
            return 'BB'
        if ratio_percent >= 75:
            return 'CB'
        if ratio_percent >= 70:
            return 'CC'
        if ratio_percent >= 65:
            return 'DC'
        if ratio_percent >= 60:
            return 'DD'
        if ratio_percent >= 50:
            return 'FD'
        return 'FF'

    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    teacher_classes = _accessible_classes(user)
    class_lookup = {cls.id: cls for cls in teacher_classes}

    if request.method == 'POST':
        class_id_raw = _normalize_form_text(request.form.get('class_id'), max_length=20)
        student_id_raw = _normalize_form_text(request.form.get('student_id'), max_length=20)
        assessment_name = _normalize_form_text(request.form.get('assessment_name'), max_length=120)
        max_points_raw = _normalize_form_text(request.form.get('max_points'), max_length=20)
        score_raw = _normalize_form_text(request.form.get('score'), max_length=20)

        if not class_id_raw.isdigit() or not student_id_raw.isdigit():
            flash('Geçersiz ders veya öğrenci seçimi.', 'danger')
            return redirect(url_for('teacher_grade_entry'))

        class_id = int(class_id_raw)
        student_id = int(student_id_raw)
        cls = class_lookup.get(class_id)
        if not cls:
            flash('Bu derse not girme yetkiniz yok.', 'danger')
            return redirect(url_for('teacher_grade_entry'))

        allowed_student_ids = {s.id for s in cls.students}
        if student_id not in allowed_student_ids:
            flash('Seçilen öğrenci bu derse kayıtlı değil.', 'danger')
            return redirect(url_for('teacher_grade_entry'))

        if not assessment_name:
            flash('Değerlendirme adı zorunludur.', 'warning')
            return redirect(url_for('teacher_grade_entry'))

        try:
            max_points = float(max_points_raw)
            score = float(score_raw)
        except (TypeError, ValueError):
            flash('Sınav puan alanları sayısal olmalıdır.', 'warning')
            return redirect(url_for('teacher_grade_entry'))

        if max_points <= 0:
            flash("Sınav kaç üzerinden bilgisi 0'dan büyük olmalıdır.", 'warning')
            return redirect(url_for('teacher_grade_entry'))

        if score < 0 or score > max_points:
            flash('Sınav notu 0 ile toplam puan arasında olmalıdır.', 'warning')
            return redirect(url_for('teacher_grade_entry'))

        ratio_percent = (score / max_points) * 100.0
        letter_grade = _letter_from_ratio(ratio_percent)

        try:
            course_code = f'CLS{cls.id:04d}'
            course = Course.query.filter_by(code=course_code).first()
            if not course:
                course = Course(
                    code=course_code,
                    title=cls.name[:120],
                    credit=DEFAULT_COURSE_CREDIT, ects=DEFAULT_COURSE_ECTS,
                    teacher_id=cls.teacher_id,
                    term_id=None,
                )
                db.session.add(course)
                db.session.flush()

            enrollment = CourseEnrollment.query.filter_by(
                student_id=student_id,
                course_id=course.id,
            ).first()
            if not enrollment:
                enrollment = CourseEnrollment(student_id=student_id, course_id=course.id, approved=True)
                db.session.add(enrollment)
                db.session.flush()

            grade = GradeRecord(
                enrollment_id=enrollment.id,
                assessment_name=f"{assessment_name} ({score:g}/{max_points:g})",
                letter_grade=letter_grade,
                grade_point=GradeRecord.point_from_letter(letter_grade),
                entered_by_id=user.id,
            )
            db.session.add(grade)
            db.session.commit()
            flash(f'Not basariyla kaydedildi. Otomatik harf notu: {letter_grade}', 'success')
        except Exception:
            db.session.rollback()
            flash('Not kaydı sırasında hata oluştu.', 'danger')

        return redirect(url_for('teacher_grade_entry', class_id=class_id))

    selected_class_id_raw = (request.args.get('class_id') or '').strip()
    selected_class = None
    if selected_class_id_raw.isdigit():
        selected_class = class_lookup.get(int(selected_class_id_raw))
    if not selected_class and teacher_classes:
        selected_class = teacher_classes[0]

    students_for_class = []
    recent_grades = []
    if selected_class:
        students_for_class = sorted(
            selected_class.students,
            key=lambda s: ((s.user.name if s.user else '').lower(), s.student_number or ''),
        )
        student_ids = [s.id for s in students_for_class]
        if student_ids:
            recent_grades = (
                GradeRecord.query
                .join(CourseEnrollment, GradeRecord.enrollment_id == CourseEnrollment.id)
                .join(Course, CourseEnrollment.course_id == Course.id)
                .filter(Course.code.in_([f"CLS{c.id:04d}" for c in teacher_classes]), CourseEnrollment.student_id.in_(student_ids))
                .order_by(GradeRecord.created_at.desc())
                .limit(50)
                .all()
            )

    return render_template(
        'teacher_grade_entry.html',
        user_name=user.name,
        teacher_classes=teacher_classes,
        selected_class=selected_class,
        students_for_class=students_for_class,
        recent_grades=recent_grades,
    )


@app.route('/teacher/course-roster', methods=['GET'])
@role_required(('teacher', 'assistant'))
@permission_required()
def teacher_course_roster():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    teacher_courses = _accessible_courses(user)
    selected_course = None
    selected_course_id_raw = (request.args.get('course_id') or '').strip()
    if selected_course_id_raw.isdigit():
        selected_course = next((c for c in teacher_courses if c.id == int(selected_course_id_raw)), None)
    if not selected_course and teacher_courses:
        selected_course = teacher_courses[0]

    search_q = (request.args.get('q') or '').strip()
    export_csv = request.args.get('export') == 'csv'

    enrolled_rows = []
    if selected_course:
        query = (
            CourseEnrollment.query
            .join(Student, CourseEnrollment.student_id == Student.id)
            .join(User, Student.id == User.id)
            .filter(CourseEnrollment.course_id == selected_course.id, CourseEnrollment.approved.is_(True))
        )
        if search_q:
            like = f'%{search_q}%'
            query = query.filter(
                db.or_(User.name.ilike(like), Student.student_number.ilike(like))
            )
        enrolled_rows = query.order_by(User.name.asc()).all()

    if export_csv and selected_course:
        import io
        si = io.StringIO()
        writer = csv.writer(si)
        writer.writerow(['Öğrenci No', 'Ad Soyad', 'E-posta', 'Kayıt Tarihi'])
        for en in enrolled_rows:
            sn = en.student.student_number if en.student else ''
            name = en.student.user.name if en.student and en.student.user else ''
            email = en.student.user.email if en.student and en.student.user else ''
            date = en.created_at.strftime('%d-%m-%Y %H:%M') if en.created_at else ''
            writer.writerow([sn, name, email, date])
        output = si.getvalue()
        from flask import Response
        filename = f'roster_{selected_course.code}.csv'
        return Response(
            output,
            mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename="{filename}"'},
        )

    return render_template(
        'teacher_course_roster.html',
        user_name=user.name,
        teacher_courses=teacher_courses,
        selected_course=selected_course,
        enrolled_rows=enrolled_rows,
        search_q=search_q,
    )


# ------------------ TEACHER CLASS HISTORY ------------------
@app.route('/teacher/history/<int:class_id>')
@role_required(('teacher', 'assistant'))
@permission_required()
def teacher_class_history(class_id):
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    cls = _accessible_class_or_404(class_id)

    week_filter = request.args.get('week', '').strip()
    start_date = request.args.get('start', '').strip()
    end_date = request.args.get('end', '').strip()

    query = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True)
    if week_filter:
        query = query.filter(AttendanceSession.week.ilike(f"%{week_filter}%"))
    if start_date:
        try:
            start_dt = datetime.datetime.strptime(start_date, '%Y-%m-%d')
            query = query.filter(AttendanceSession.date >= start_dt)
        except ValueError:
            flash('Geçersiz başlangıç tarihi formatı.', 'warning')
    if end_date:
        try:
            end_dt = datetime.datetime.strptime(end_date, '%Y-%m-%d') + datetime.timedelta(days=1)
            query = query.filter(AttendanceSession.date < end_dt)
        except ValueError:
            flash('Geçersiz bitiş tarihi formatı.', 'warning')

    sessions = query.order_by(AttendanceSession.date.desc()).all()
    session_ids = [sess.id for sess in sessions]
    present_counts = {}
    if session_ids:
        present_rows = (
            db.session.query(AttendanceRecord.session_id, func.count(AttendanceRecord.id))
            .filter(AttendanceRecord.session_id.in_(session_ids), AttendanceRecord.present.is_(True))
            .group_by(AttendanceRecord.session_id)
            .all()
        )
        present_counts = {session_id: count for session_id, count in present_rows}

    total_students = len(cls.students)
    sessions_data = []
    for sess in sessions:
        present_count = present_counts.get(sess.id, 0)
        sessions_data.append({
            'id': sess.id,
            'date': sess.date,
            'week': sess.week,
            'present_count': present_count,
            'total_students': total_students,
        })

    return render_template('teacher_history.html', cls=cls, sessions=sessions_data, week_filter=week_filter, start_date=start_date, end_date=end_date)


@app.route('/teacher/history/<int:class_id>/export')
@role_required(('teacher', 'assistant'))
@permission_required()
def export_teacher_class_history(class_id):
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    cls = _accessible_class_or_404(class_id)

    week_filter = request.args.get('week', '').strip()
    start_date = request.args.get('start', '').strip()
    end_date = request.args.get('end', '').strip()

    query = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True)
    if week_filter:
        query = query.filter(AttendanceSession.week.ilike(f"%{week_filter}%"))
    if start_date:
        try:
            start_dt = datetime.datetime.strptime(start_date, '%Y-%m-%d')
            query = query.filter(AttendanceSession.date >= start_dt)
        except ValueError:
            pass
    if end_date:
        try:
            end_dt = datetime.datetime.strptime(end_date, '%Y-%m-%d') + datetime.timedelta(days=1)
            query = query.filter(AttendanceSession.date < end_dt)
        except ValueError:
            pass

    sessions = query.order_by(AttendanceSession.date.desc()).all()
    output = StringIO()
    writer = csv.writer(output)

    def padded(value):
        """Add visual breathing room inside cells without extra CSV separators."""
        return f" {value} "

    writer.writerow(['Session Date', 'Week', 'Student Number', 'Student Name', 'Student Email', 'Attendance'])

    sorted_students = sorted(cls.students, key=lambda s: (s.student_number or '', s.id))
    for idx, sess in enumerate(sessions):
        present_ids = {record.student_id for record in sess.records if record.present}

        if not sorted_students:
            writer.writerow([
                padded(sess.date.strftime('%d-%m-%Y %H:%M')),
                padded(sess.week or '-'),
                padded('-'),
                padded('-'),
                padded('-'),
                padded('No students in class'),
            ])
            continue

        for student in sorted_students:
            student_name = student.user.name if student.user else f"Student {student.id}"
            student_email = student.user.email if student.user else '-'
            status = 'Present' if student.id in present_ids else 'Absent'
            writer.writerow([
                padded(sess.date.strftime('%d-%m-%Y %H:%M')),
                padded(sess.week or '-'),
                padded(student.student_number),
                padded(student_name),
                padded(student_email),
                padded(status),
            ])

        # Add visual separation between session/week blocks for Excel readability.
        if idx < len(sessions) - 1:
            writer.writerow([])
            writer.writerow([])

    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename={cls.name}_attendance_history.csv'}
    )


@app.route('/teacher/session/<int:session_id>/stats')
@role_required(('teacher', 'assistant'))
@permission_required()
def session_stats(session_id):
    session_obj = AttendanceSession.query.get_or_404(session_id)
    denied = ensure_teacher_session_ownership(session_obj, on_fail='json')
    if denied:
        return denied
    cls = session_obj.class_obj

    present_count = AttendanceRecord.query.filter_by(session_id=session_obj.id, present=True).count()
    return jsonify({
        'present_count': present_count,
        'total_count': len(cls.students),
        'active': session_obj.active,
    })


# Keep legacy history route for backward compatibility
@app.route('/teacher/history')
@role_required('teacher')
@permission_required()
def teacher_history_redirect():
    return redirect(url_for('teacher_dashboard'))


# ------------------ SESSION DETAIL ------------------
@app.route('/teacher/session/<int:session_id>')
@role_required(('teacher', 'assistant'))
@permission_required()
def session_detail(session_id):
    session_obj = AttendanceSession.query.get_or_404(session_id)
    denied = ensure_teacher_session_ownership(session_obj, on_fail='login')
    if denied:
        return denied
    cls = session_obj.class_obj

    attendance_ids = {r.student_id for r in session_obj.records if r.present}
    confirmed_sessions = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True).all()
    confirmed_session_ids = [s.id for s in confirmed_sessions]
    total_confirmed_sessions = len(confirmed_session_ids)
    attended_counts = {}

    if confirmed_session_ids:
        attended_rows = (
            db.session.query(AttendanceRecord.student_id, func.count(AttendanceRecord.id))
            .filter(
                AttendanceRecord.session_id.in_(confirmed_session_ids),
                AttendanceRecord.present == True,
            )
            .group_by(AttendanceRecord.student_id)
            .all()
        )
        attended_counts = {student_id: count for student_id, count in attended_rows}

    absence_stats = {}
    for student in cls.students:
        if total_confirmed_sessions == 0:
            absence_stats[student.id] = {
                'absent_count': 0,
                'total_sessions': 0,
                'absence_pct': 0.0,
            }
            continue

        attended_count = int(attended_counts.get(student.id, 0))
        absent_count = total_confirmed_sessions - attended_count
        absence_stats[student.id] = {
            'absent_count': absent_count,
            'total_sessions': total_confirmed_sessions,
            'absence_pct': (absent_count / total_confirmed_sessions) * 100,
        }

    return render_template(
        'teacher_session_detail.html',
        cls=cls,
        session=session_obj,
        attendance_ids=attendance_ids,
        absence_stats=absence_stats,
    )


# ------------------ DELETE SESSION ------------------
@app.route('/teacher/session/<int:session_id>/delete', methods=['POST'])
@role_required('teacher')
@permission_required()
def delete_session(session_id):
    session_obj = AttendanceSession.query.get_or_404(session_id)
    denied = ensure_teacher_session_ownership(session_obj, on_fail='login')
    if denied:
        return denied
    cls = session_obj.class_obj

    # Only confirmed sessions can be deleted.
    if not session_obj.confirmed:
        flash('Bu oturumun silinmeden önce tamamlanması gerekiyor.', 'warning')
        return redirect(url_for('teacher_class_history', class_id=cls.id))

    AttendanceRecord.query.filter_by(session_id=session_id).delete()
    db.session.delete(session_obj)
    db.session.commit()

    flash('Yoklama kaydı silindi.', 'success')
    return redirect(url_for('teacher_class_history', class_id=cls.id))


# ------------------ TEACHER ACCOUNT ------------------
@app.route('/teacher/account', methods=['GET', 'POST'])
@role_required('teacher')
@permission_required()
def teacher_account():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    if request.method == 'POST':
        permission_error = ensure_permission(PERMISSIONS['TEACHER_ACCOUNT_UPDATE'], 'teacher_account')
        if permission_error:
            return permission_error

        action = _normalize_form_text(request.form.get('action'), max_length=40).lower()
        if action not in {'change_password', 'delete_account', 'update_profile', 'upload_photo', 'remove_photo'}:
            flash('Geçersiz işlem.', 'warning')
            return redirect(url_for('teacher_account'))

        if action == 'update_profile':
            user.academic_title = _normalize_form_text(request.form.get('academic_title'), max_length=60) or None
            user.faculty = _normalize_form_text(request.form.get('faculty'), max_length=140) or None
            user.department = _normalize_form_text(request.form.get('department'), max_length=140) or None
            user.office = _normalize_form_text(request.form.get('office'), max_length=80) or None
            user.office_hours = _normalize_form_text(request.form.get('office_hours'), max_length=160) or None
            db.session.commit()
            flash(_lang_text('Profil bilgileri güncellendi.', 'Profile updated.'), 'success')
            return redirect(url_for('teacher_account'))

        if action == 'remove_photo':
            _remove_profile_photo_file(user.photo_filename)
            user.photo_filename = None
            db.session.commit()
            flash(_lang_text('Fotoğraf kaldırıldı.', 'Photo removed.'), 'success')
            return redirect(url_for('teacher_account'))

        if action == 'upload_photo':
            upload = request.files.get('photo')
            error = (_save_profile_photo(user, upload, prefix='t') if upload and upload.filename
                     else _lang_text('Bir fotoğraf seçin.', 'Choose a photo.'))
            if error:
                flash(error, 'warning')
            else:
                db.session.commit()
                flash(_lang_text('Fotoğraf güncellendi.', 'Photo updated.'), 'success')
            return redirect(url_for('teacher_account'))

        if action == 'change_password':
            current_password = request.form.get('current_password') or ''
            new_password = request.form.get('new_password') or ''
            new_password_confirm = request.form.get('new_password_confirm') or ''

            if not current_password or not new_password or not new_password_confirm:
                flash('Tüm şifre alanları zorunludur.', 'warning')
                return redirect(url_for('teacher_account'))

            if not user.check_password(current_password):
                flash('Mevcut şifre hatalı.', 'danger')
                return redirect(url_for('teacher_account'))

            if new_password != new_password_confirm:
                flash('Yeni şifre ve onay eşleşmiyor.', 'warning')
                return redirect(url_for('teacher_account'))

            password_error = validate_password_policy(new_password, user_name=user.name, user_email=user.email)
            if password_error:
                flash(password_error, 'danger')
                return redirect(url_for('teacher_account'))

            user.set_password(new_password)
            db.session.commit()
            flash('Şifreniz başarıyla değiştirildi.', 'success')
            return redirect(url_for('teacher_account'))

        if action == 'delete_account':
            password = request.form.get('password') or ''
            if not password:
                flash('Şifre zorunludur.', 'warning')
                return redirect(url_for('teacher_account'))
            if not user.check_password(password):
                flash('Şifre hatalı.', 'danger')
                return redirect(url_for('teacher_account'))

            # Delete all related user data.
            for cls in user.teacher_classes:
                # Delete class sessions and records.
                for sess in cls.sessions:
                    AttendanceRecord.query.filter_by(session_id=sess.id).delete()
                AttendanceSession.query.filter_by(class_id=cls.id).delete()
                # Remove class-student links.
                cls.students = []
                db.session.delete(cls)
            # Delete student profile if present.
            for student in user.students:
                db.session.delete(student)

            _remove_profile_photo_file(user.photo_filename)
            db.session.delete(user)
            db.session.commit()

            session.clear()
            flash('Hesabınız silindi.', 'success')
            return redirect(url_for('login'))

    return render_template('account.html', **_account_context(
        user, _lang_text('Öğretmen', 'Teacher'), 'teacher_account', show_profile=True, photo_enabled=True,
        photo_url=url_for('teacher_photo', user_id=user.id) if user.photo_filename else None, delete_mode='enabled'))


@app.route('/admin/account', methods=['GET', 'POST'])
@role_required('admin')
@permission_required()
@csrf_protect
def admin_account():
    """Admin profile summary and password change (no self-delete: the system must keep an admin)."""
    return _staff_account_page('admin_account', _lang_text('Yönetici', 'Administrator'))


@app.route('/registrar/account', methods=['GET', 'POST'])
@role_required('registrar')
@permission_required()
@csrf_protect
def registrar_account():
    return _staff_account_page('registrar_account', _lang_text('Öğrenci İşleri', 'Student Affairs'))


def _account_context(user, role_label, endpoint, show_profile=False, photo_url=None, photo_enabled=False,
                     delete_mode=None):
    """Context for the shared account.html page (every role renders the same template)."""
    return {
        'user': user,
        'role_label': role_label,
        'acc_endpoint': endpoint,
        'show_profile': show_profile,
        'photo_enabled': photo_enabled,
        'photo_url': photo_url,
        'delete_mode': delete_mode,
        'password_min_length': int(app.config.get('PASSWORD_MIN_LENGTH', 10) or 10),
    }


def _staff_account_page(endpoint, role_label):
    """Shared profile + password change page for staff roles (admin, Student Affairs)."""
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        return redirect(url_for('login'))

    if request.method == 'POST':
        current_password = request.form.get('current_password') or ''
        new_password = request.form.get('new_password') or ''
        new_password_confirm = request.form.get('new_password_confirm') or ''
        if not current_password or not new_password or not new_password_confirm:
            flash(_lang_text('Tüm şifre alanları zorunludur.', 'All password fields are required.'), 'warning')
        elif not user.check_password(current_password):
            flash(_lang_text('Mevcut şifre hatalı.', 'Current password is incorrect.'), 'danger')
        elif new_password != new_password_confirm:
            flash(_lang_text('Yeni şifre ve onay eşleşmiyor.', 'New password and confirmation do not match.'), 'warning')
        else:
            password_error = validate_password_policy(new_password, user_name=user.name, user_email=user.email)
            if password_error:
                flash(password_error, 'danger')
            else:
                user.set_password(new_password)
                db.session.commit()
                _log_admin_operation('password_change_self', user.id, None, None)
                flash(_lang_text('Şifreniz başarıyla değiştirildi.', 'Your password has been changed.'), 'success')
        return redirect(url_for(endpoint))

    return render_template('account.html', **_account_context(user, role_label, endpoint))


@app.route('/registrar/dashboard', methods=['GET', 'POST'])
@role_required('registrar')
@permission_required()
def registrar_dashboard():
    user = db.session.get(User, session.get('user_id'))
    if request.method == 'POST':
        return _handle_user_calendar_dashboard_post(user, PERMISSIONS['REGISTRAR_DASHBOARD_CALENDAR_UPDATE'], 'registrar_dashboard')

    groups = StudentGroup.query.all()
    student_count = Student.query.count()
    grouped_ids = {s.id for g in groups for s in g.students}
    profile = _staff_profile(user, _lang_text('Öğrenci İşleri', 'Student Affairs'), 'registrar_account', [
        (student_count, _lang_text('Öğrenci', 'Students')),
        (len(groups), _lang_text('Bölüm grubu', 'Dept. groups')),
        (max(student_count - len(grouped_ids), 0), _lang_text('Grupsuz', 'Ungrouped')),
    ])
    _, upcoming, events = _serialize_user_calendar_events(user.id)
    return _render_home_dashboard(user, 'registrar', 'registrar_dashboard',
                                  _lang_text('Öğrenci İşleri Paneli', 'Student Affairs Dashboard'), profile, events, len(upcoming))


@app.route('/teacher/class/<int:class_id>/assistants', methods=['POST'])
@role_required('teacher')
@permission_required()
@csrf_protect
def teacher_class_assistants(class_id):
    """The class's teacher adds or removes teaching assistants (users with the 'assistant' role)."""
    cls = db.session.get(Class, class_id)
    if not cls or cls.teacher_id != session.get('user_id'):
        flash(_lang_text('Bu derse asistan atama yetkiniz yok.', 'You cannot assign assistants to this class.'), 'danger')
        return redirect(url_for('teacher_attendance_hub'))

    action = _normalize_form_text(request.form.get('action'), max_length=20).lower()
    if action == 'remove':
        raw = _normalize_form_text(request.form.get('assistant_id'), max_length=20)
        assistant = db.session.get(User, int(raw)) if raw.isdigit() else None
        if assistant is not None and assistant in cls.assistants:
            cls.assistants.remove(assistant)
            db.session.commit()
            flash(_lang_text(f'{assistant.name} bu dersten çıkarıldı.', f'{assistant.name} removed from this class.'), 'success')
        return redirect(url_for('teacher_attendance_hub', open=cls.id))

    assistant = _resolve_recipient(request.form.get('assistant'))
    if assistant is None or assistant.role != 'assistant':
        flash(_lang_text('Asistan bulunamadı. Listeden asistan rolündeki bir kullanıcı seçin.',
                         'Assistant not found. Pick a user with the assistant role from the list.'), 'warning')
    elif assistant in cls.assistants:
        flash(_lang_text('Bu kişi zaten bu dersin asistanı.', 'This person already assists this class.'), 'info')
    else:
        cls.assistants.append(assistant)
        db.session.commit()
        flash(_lang_text(f'{assistant.name} derse asistan olarak eklendi.', f'{assistant.name} added as an assistant.'), 'success')
    return redirect(url_for('teacher_attendance_hub', open=cls.id))


@app.route('/assistant/dashboard', methods=['GET', 'POST'])
@role_required('assistant')
@permission_required()
def assistant_dashboard():
    user = db.session.get(User, session.get('user_id'))
    if request.method == 'POST':
        return _handle_user_calendar_dashboard_post(user, PERMISSIONS['ASSISTANT_DASHBOARD_CALENDAR_UPDATE'], 'assistant_dashboard')

    classes = user.assisted_classes
    profile = _staff_profile(user, _lang_text('Asistan', 'Assistant'), 'assistant_account', [
        (len(classes), _lang_text('Ders', 'Classes')),
        (len({st.id for cls in classes for st in cls.students}), _lang_text('Öğrenci', 'Students')),
    ])
    _, upcoming, events = _serialize_user_calendar_events(user.id)
    return _render_home_dashboard(user, 'assistant', 'assistant_dashboard',
                                  _lang_text('Asistan Paneli', 'Assistant Dashboard'), profile, events, len(upcoming))


# ------------------ ADVISOR ------------------
def _advisees_of(teacher_id):
    """Students whose effective advisor is this teacher (own assignment or department group)."""
    group_ids = [g.id for g in StudentGroup.query.filter_by(advisor_id=teacher_id).all()]
    candidates = {s.id: s for s in Student.query.filter_by(advisor_id=teacher_id).all()}
    if group_ids:
        for s in Student.query.filter(Student.groups.any(StudentGroup.id.in_(group_ids))).all():
            candidates.setdefault(s.id, s)
    return [s for s in candidates.values() if (s.advisor_info()[0] or User()).id == teacher_id]


def _advisee_or_none(student_id):
    """The student only if the logged-in teacher is their advisor (admin may see any)."""
    student = db.session.get(Student, student_id)
    if not student:
        return None
    advisor = student.advisor_info()[0]
    if session.get('role') != 'admin' and (advisor is None or advisor.id != session.get('user_id')):
        return None
    return student


def _student_academic_summary(student):
    """Per-course grades and per-class attendance for one student (advisor view)."""
    courses, weighted_points, weighted_credits = [], 0.0, 0
    for enrollment in sorted(student.course_enrollments, key=lambda e: e.course.code if e.course else ''):
        course = enrollment.course
        if not course:
            continue
        avg = (sum(g.grade_point for g in enrollment.grades) / len(enrollment.grades)) if enrollment.grades else None
        if avg is not None and (course.ects or 0) > 0:
            weighted_points += avg * course.ects
            weighted_credits += course.ects
        courses.append({'course': course, 'grades': enrollment.grades, 'average': avg})

    attendance = []
    for cls in sorted(student.classes, key=lambda c: c.name or ''):
        total = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True).count()
        attended = (
            db.session.query(func.count(AttendanceRecord.id))
            .join(AttendanceSession, AttendanceRecord.session_id == AttendanceSession.id)
            .filter(AttendanceSession.class_id == cls.id, AttendanceSession.confirmed == True,
                    AttendanceRecord.student_id == student.id, AttendanceRecord.present == True)
            .scalar() or 0
        )
        pct = (attended / total * 100) if total else 0
        attendance.append({'cls': cls, 'total': total, 'attended': attended, 'percentage': pct,
                           'risk': total > 0 and pct < 70})
    gpa = (weighted_points / weighted_credits) if weighted_credits else None
    return {'courses': courses, 'attendance': attendance, 'gpa': gpa, 'credits': weighted_credits}


STUDENT_PHOTO_DIR = os.path.join(app.root_path, 'uploads', 'student_photos')
STUDENT_PHOTO_MAX_BYTES = 2 * 1024 * 1024


def _save_profile_photo(owner, upload, prefix=''):
    """Validate, re-encode (strips metadata and any embedded payload) and store a profile photo for a
    Student or User (anything with id and photo_filename). Returns an error message, or None on success."""
    from PIL import Image, ImageOps
    data = upload.read(STUDENT_PHOTO_MAX_BYTES + 1)
    if not data:
        return _lang_text('Bir fotoğraf seçin.', 'Choose a photo.')
    if len(data) > STUDENT_PHOTO_MAX_BYTES:
        return _lang_text('Fotoğraf en fazla 2 MB olabilir.', 'The photo can be at most 2 MB.')
    try:
        probe = Image.open(BytesIO(data))
        probe.verify()
        image = Image.open(BytesIO(data))
        if image.format not in ('JPEG', 'PNG', 'WEBP') or image.width * image.height > 25_000_000:
            raise ValueError('unsupported image')
        image = ImageOps.fit(ImageOps.exif_transpose(image).convert('RGB'), (400, 400))
    except Exception:
        return _lang_text('Geçerli bir JPG, PNG veya WEBP fotoğraf yükleyin.', 'Upload a valid JPG, PNG or WEBP photo.')
    os.makedirs(STUDENT_PHOTO_DIR, exist_ok=True)
    filename = f'{prefix}{owner.id}_{secrets.token_hex(8)}.jpg'
    image.save(os.path.join(STUDENT_PHOTO_DIR, filename), 'JPEG', quality=88)
    _remove_profile_photo_file(owner.photo_filename)
    owner.photo_filename = filename
    return None


def _remove_profile_photo_file(filename):
    if filename:
        try:
            os.remove(os.path.join(STUDENT_PHOTO_DIR, os.path.basename(filename)))
        except OSError:
            pass


@app.route('/student-photo/<int:student_id>')
@login_required_session
def student_photo(student_id):
    """A student's own photo, or any photo for teachers, registrars and admins."""
    role = session.get('role')
    if not (role in ('teacher', 'registrar', 'admin') or (role == 'student' and session.get('user_id') == student_id)):
        abort(404)
    student = db.session.get(Student, student_id)
    if not student or not student.photo_filename:
        abort(404)
    response = send_from_directory(STUDENT_PHOTO_DIR, os.path.basename(student.photo_filename), mimetype='image/jpeg')
    response.headers['Cache-Control'] = 'private, max-age=300'
    return response


@app.route('/teacher-photo/<int:user_id>')
@login_required_session
def teacher_photo(user_id):
    """A teacher's profile photo; staff photos are visible to every signed-in user."""
    teacher = db.session.get(User, user_id)
    if not teacher or teacher.role != 'teacher' or not teacher.photo_filename:
        abort(404)
    response = send_from_directory(STUDENT_PHOTO_DIR, os.path.basename(teacher.photo_filename), mimetype='image/jpeg')
    response.headers['Cache-Control'] = 'private, max-age=300'
    return response


def _teacher_profile_summary(user, teacher_classes):
    not_set = _lang_text('Tanımlı değil', 'Not set')
    return {
        'id': user.id,
        'name': user.name,
        'title': (user.academic_title or '').strip(),
        'email': user.email,
        'faculty': (user.faculty or '').strip() or not_set,
        'department': (user.department or '').strip() or not_set,
        'office': (user.office or '').strip() or not_set,
        'office_hours': (user.office_hours or '').strip() or not_set,
        'class_count': len(teacher_classes),
        'student_count': len({st.id for cls in teacher_classes for st in cls.students}),
        'advisee_count': len(_advisees_of(user.id)),
        'assistant_count': len({a.id for cls in teacher_classes for a in cls.assistants}),
        'has_photo': bool(user.photo_filename),
        'initials': ''.join(part[0] for part in (user.name or '?').split()[:2]).upper(),
    }


def _student_class_level(student):
    """Year of study, derived from the entry year in the student number (e.g. K20240362 -> 2024)
    and the current academic year; None when it cannot be determined."""
    match = re.search(r'(20\d{2})', student.student_number or '')
    if not match:
        return None
    academic = re.search(r'(20\d{2})', student.university_academic_year or '')
    if academic:
        start_year = int(academic.group(1))
    else:
        today = datetime.date.today()
        start_year = today.year if today.month >= 9 else today.year - 1
    level = start_year - int(match.group(1)) + 1
    return level if 1 <= level <= 8 else None


def _student_profile_summary(student):
    not_set = _lang_text('Tanımlı değil', 'Not set')
    advisor = student.advisor_info()[0]
    level = _student_class_level(student)
    department = (student.university_department or '').strip() or ', '.join(sorted(g.code for g in student.groups))
    academic_year = (student.university_academic_year or '').strip()
    term = (student.university_term or '').strip()
    return {
        'name': student.user.name if student.user else '-',
        'student_number': student.student_number,
        'email': student.user.email if student.user else '-',
        'faculty': (student.university_faculty or '').strip() or not_set,
        'department': department or not_set,
        'class_level': _lang_text(f'{level}. sınıf', f'Year {level}') if level else not_set,
        'period': ' · '.join(part for part in (academic_year, term) if part) or not_set,
        'advisor': advisor.name if advisor else not_set,
        'has_photo': bool(student.photo_filename),
        'initials': ''.join(part[0] for part in (student.user.name if student.user else '?').split()[:2]).upper(),
    }


@app.route('/teacher/advisees')
@role_required('teacher')
@permission_required()
def teacher_advisees():
    """Student directory for teachers: every student, with the teacher's own advisees marked."""
    user = db.session.get(User, session.get('user_id'))
    mine = {s.id for s in _advisees_of(user.id)} if user else set()
    rows = []
    for student in sorted(Student.query.all(), key=lambda s: (s.user.name if s.user else '')):
        advisor = student.advisor_info()[0]
        row = {
            'student': student,
            'name': student.user.name if student.user else '-',
            'groups': ', '.join(sorted(g.code for g in student.groups)),
            'advisor_name': advisor.name if advisor else None,
            'is_mine': student.id in mine,
            'gpa': None, 'course_count': None, 'risk_count': 0,
        }
        if row['is_mine']:
            summary = _student_academic_summary(student)
            row.update(gpa=summary['gpa'], course_count=len(summary['courses']),
                       risk_count=sum(1 for a in summary['attendance'] if a['risk']))
        rows.append(row)
    return render_template('teacher_advisees.html', user=user, rows=rows)


# Read-only "view as student": the teacher's session is swapped for the student's, and only the
# student's read pages stay reachable until they exit (see restrict_view_as_student).
VIEW_AS_ALLOWED_ENDPOINTS = {
    'static', 'view_as_student_exit', 'logout', 'set_language', 'student_photo', 'student_dashboard', 'student_absence',
    'student_identity_info', 'student_education_info', 'student_family_info', 'student_documents_info',
    'student_contact_info', 'student_term_courses',
    'student_transcript', 'student_academic_calendar', 'student_exams', 'student_class_history', 'reports',
}  # finance pages (current account, payments) are deliberately not viewable by teachers


def _end_view_as():
    """Swap the teacher's own identity back into the session."""
    origin = session.pop('view_as_origin', None)
    session.pop('view_as_name', None)
    if origin:
        session['user_id'] = origin['user_id']
        session['role'] = origin['role']
    return origin


def _view_as_can_edit():
    """True while a teacher views one of their own advisees (effective advisor), who may edit courses."""
    origin = session.get('view_as_origin')
    if not origin:
        return False
    student = db.session.get(Student, session.get('user_id'))
    advisor = student.advisor_info()[0] if student else None
    return bool(advisor and advisor.id == origin['user_id'])


@app.context_processor
def inject_view_as_flags():
    if not session.get('view_as_origin'):
        return {'view_as_can_edit': False, 'view_as_profile': None}
    student = db.session.get(Student, session.get('user_id'))
    return {
        'view_as_can_edit': _view_as_can_edit(),
        'view_as_profile': _student_profile_summary(student) if student else None,
    }


@app.before_request
def restrict_view_as_student():
    if not session.get('view_as_origin'):
        return None
    endpoint = request.endpoint
    if endpoint in ('student_term_courses_update', 'view_as_term_plan_approve') and _view_as_can_edit():
        return None
    if endpoint in VIEW_AS_ALLOWED_ENDPOINTS and (request.method == 'GET' or endpoint in (
            'view_as_student_exit', 'logout', 'set_language')):
        return None
    if request.method == 'GET':
        # A teacher page (top navigation): leave the student view and open it as the teacher.
        _end_view_as()
        return None
    flash(_t('view_as_read_only'), 'warning')
    return redirect(url_for('student_term_courses'))


@app.route('/teacher/students/<int:student_id>/view', methods=['POST'])
@role_required('teacher')
@permission_required()
@csrf_protect
def teacher_view_student(student_id):
    student = db.session.get(Student, student_id)
    if not student or not student.user:
        flash(_t('advisee_not_found'), 'warning')
        return redirect(url_for('teacher_advisees'))
    session['view_as_origin'] = {'user_id': session.get('user_id'), 'role': session.get('role')}
    session['view_as_name'] = student.user.name
    session['user_id'] = student.id
    session['role'] = 'student'
    return redirect(url_for('student_term_courses'))


@app.route('/view-as/exit', methods=['POST'])
@login_required_session
@csrf_protect
def view_as_student_exit():
    if not _end_view_as():
        return redirect(url_for('login'))
    return redirect(url_for('teacher_advisees'))


@app.route('/teacher/advisees/<int:student_id>')
@role_required('teacher')
@permission_required()
def teacher_advisee_detail(student_id):
    student = _advisee_or_none(student_id)
    if not student:
        flash(_t('advisee_not_found'), 'warning')
        return redirect(url_for('teacher_advisees'))
    return render_template('teacher_advisee_detail.html', student=student,
                           name=student.user.name if student.user else '-',
                           email=student.user.email if student.user else '-',
                           summary=_student_academic_summary(student))


@app.route('/assistant/account', methods=['GET', 'POST'])
@role_required('assistant')
@permission_required()
@csrf_protect
def assistant_account():
    return _staff_account_page('assistant_account', _lang_text('Asistan', 'Teaching Assistant'))


# ------------------ STUDENT ACCOUNT ------------------
@app.route('/student/account', methods=['GET', 'POST'])
@role_required('student')
@permission_required()
def student_account():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    if request.method == 'POST':
        permission_error = ensure_permission(PERMISSIONS['STUDENT_ACCOUNT_UPDATE'], 'student_account')
        if permission_error:
            return permission_error

        action = _normalize_form_text(request.form.get('action'), max_length=40).lower()
        if action not in {'change_password', 'delete_account', 'upload_photo', 'remove_photo'}:
            flash('Geçersiz işlem.', 'warning')
            return redirect(url_for('student_account'))

        if action in ('upload_photo', 'remove_photo'):
            student = user.students[0] if user.students else None
            if student is None:
                flash('Öğrenci profili bulunamadı.', 'warning')
                return redirect(url_for('student_account'))
            if action == 'remove_photo':
                _remove_profile_photo_file(student.photo_filename)
                student.photo_filename = None
                db.session.commit()
                flash(_lang_text('Fotoğraf kaldırıldı.', 'Photo removed.'), 'success')
                return redirect(url_for('student_account'))
            upload = request.files.get('photo')
            error = _save_profile_photo(student, upload) if upload and upload.filename else _lang_text('Bir fotoğraf seçin.', 'Choose a photo.')
            if error:
                flash(error, 'warning')
            else:
                db.session.commit()
                flash(_lang_text('Fotoğraf güncellendi.', 'Photo updated.'), 'success')
            return redirect(url_for('student_account'))

        if action == 'change_password':
            current_password = request.form.get('current_password') or ''
            new_password = request.form.get('new_password') or ''
            new_password_confirm = request.form.get('new_password_confirm') or ''

            if not current_password or not new_password or not new_password_confirm:
                flash('Tüm şifre alanları zorunludur.', 'warning')
                return redirect(url_for('student_account'))

            if not user.check_password(current_password):
                flash('Mevcut şifre hatalı.', 'danger')
                return redirect(url_for('student_account'))

            if new_password != new_password_confirm:
                flash('Yeni şifre ve onay eşleşmiyor.', 'warning')
                return redirect(url_for('student_account'))

            password_error = validate_password_policy(new_password, user_name=user.name, user_email=user.email)
            if password_error:
                flash(password_error, 'danger')
                return redirect(url_for('student_account'))

            user.set_password(new_password)
            db.session.commit()
            flash('Şifreniz başarıyla değiştirildi.', 'success')
            return redirect(url_for('student_account'))

        if action == 'delete_account':
            flash('Hesap silme devre dışı. Lütfen sistem yöneticisiyle iletişime geçin.', 'warning')
            return redirect(url_for('student_account'))

    profile = user.students[0] if user.students else None
    return render_template('account.html', **_account_context(
        user, _lang_text('Öğrenci', 'Student'), 'student_account', photo_enabled=profile is not None,
        photo_url=url_for('student_photo', student_id=user.id) if profile and profile.photo_filename else None,
        delete_mode='disabled'))


# ------------------ STUDENT DASHBOARD ------------------
@app.route('/student_dashboard', methods=['GET', 'POST'])
@role_required('student')
@permission_required()
def student_dashboard():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    if not user.students:
        flash('Öğrenci profili bulunamadı. Lütfen destek ekibiyle iletişime geçin.', 'danger')
        return redirect(url_for('logout'))
    student = user.students[0]

    if request.method == 'POST':
        permission_error = ensure_permission(PERMISSIONS['STUDENT_DASHBOARD_CALENDAR_UPDATE'], 'student_dashboard')
        if permission_error:
            return permission_error

        action = _normalize_form_text(request.form.get('action'), max_length=40).lower()
        if action not in {'add_calendar_event', 'delete_calendar_event'}:
            flash('Geçersiz panel işlemi.', 'warning')
            return redirect(url_for('student_dashboard'))

        if action == 'add_calendar_event':
            title = (request.form.get('title') or '').strip()
            event_type = (request.form.get('event_type') or 'activity').strip().lower()
            event_date_raw = (request.form.get('event_date') or '').strip()
            note = (request.form.get('note') or '').strip()

            if len(title) > 120:
                flash('Etkinlik başlığı çok uzun.', 'warning')
                return redirect(url_for('student_dashboard'))

            if not title or not event_date_raw:
                flash('Etkinlik başlığı ve tarihi zorunludur.', 'warning')
                return redirect(url_for('student_dashboard'))

            if event_type not in ('exam', 'activity'):
                event_type = 'activity'

            try:
                event_date = datetime.datetime.strptime(event_date_raw, '%Y-%m-%d').date()
            except ValueError:
                flash('Geçersiz etkinlik tarihi formatı.', 'warning')
                return redirect(url_for('student_dashboard'))

            event = StudentCalendarEvent(
                student_id=student.id,
                title=title,
                event_type=event_type,
                event_date=event_date,
                note=note[:240] if note else None,
            )
            db.session.add(event)
            db.session.commit()
            flash('Takvim etkinliği eklendi.', 'success')
            return redirect(url_for('student_dashboard'))

        if action == 'delete_calendar_event':
            event_id_raw = _normalize_form_text(request.form.get('event_id'), max_length=20)
            if not event_id_raw.isdigit():
                flash('Geçersiz etkinlik kimliği.', 'warning')
                return redirect(url_for('student_dashboard'))
            event_id = int(event_id_raw) if event_id_raw and event_id_raw.isdigit() else None
            event = db.session.get(StudentCalendarEvent, event_id) if event_id is not None else None
            denied = ensure_student_event_ownership(
                student,
                event,
                on_fail='student_dashboard',
                fail_message='Event not found.',
            )
            if denied:
                return denied

            db.session.delete(event)
            db.session.commit()
            flash('Takvim etkinliği silindi.', 'success')
            return redirect(url_for('student_dashboard'))

    today = datetime.date.today()
    all_events = StudentCalendarEvent.query.filter_by(student_id=student.id).order_by(StudentCalendarEvent.event_date.asc()).all()
    upcoming_count = len([e for e in all_events if e.event_date >= today][:12])
    calendar_events = [
        {
            'id': e.id,
            'title': e.title,
            'event_type': e.event_type,
            'event_date': e.event_date.isoformat(),
            'note': e.note or '',
        }
        for e in all_events
    ]

    sp = _student_profile_summary(student)
    profile = {
        'name': sp['name'],
        'initials': sp['initials'],
        'photo_url': url_for('student_photo', student_id=student.id) if sp['has_photo'] else None,
        'facts': [
            (_lang_text('Öğrenci no', 'Student no'), sp['student_number']),
            (_lang_text('E-posta', 'E-mail'), sp['email']),
            (_lang_text('Fakülte', 'Faculty'), sp['faculty']),
            (_lang_text('Bölüm', 'Department'), sp['department']),
            (_lang_text('Sınıf', 'Year'), sp['class_level']),
            (_lang_text('Danışman', 'Advisor'), sp['advisor']),
        ],
        'stats': [
            (len(student.classes), _lang_text('Ders', 'Classes')),
            (len(student.course_enrollments), _lang_text('Kayıtlı ders', 'Enrolled')),
            (upcoming_count, _lang_text('Yaklaşan', 'Upcoming')),
        ],
        'edit_url': url_for('student_account'),
    }
    return _render_home_dashboard(user, 'student', 'student_dashboard', _lang_text('Öğrenci Paneli', 'Student Dashboard'), profile,
                                  calendar_events, upcoming_count)


@app.route('/announcements')
@login_required_session
def announcements_inbox():
    """Mailbox: announcements addressed to the user plus personal messages (inbox), and what they sent."""
    user = db.session.get(User, session.get('user_id'))
    role = session.get('role')
    if not user or role not in ROLE_DASHBOARDS:
        return redirect(url_for('login'))

    def _newest_first(rows):
        return sorted(rows, key=lambda r: r['created_at'] or datetime.datetime.min, reverse=True)

    received_announcements = (_announcements_for_user(user, role, include_expired=True)
                              .filter(Announcement.author_id != user.id)
                              .order_by(Announcement.created_at.desc()).limit(200).all())
    received_messages = Message.query.filter_by(recipient_id=user.id).order_by(Message.sent_at.desc()).limit(200).all()
    sent_announcements = Announcement.query.filter_by(author_id=user.id).order_by(Announcement.created_at.desc()).limit(200).all()
    sent_messages = Message.query.filter_by(sender_id=user.id).order_by(Message.sent_at.desc()).limit(200).all()

    received = _newest_first([_mailbox_item(o) for o in received_announcements + received_messages])
    sent = _newest_first([_mailbox_item(o) for o in sent_announcements + sent_messages])
    can_send = role in {'admin', 'registrar', 'teacher'} or bool(sent)

    # Advisors: term plans their advisees submitted, waiting for approval (oldest first).
    approvals = []
    if role == 'teacher':
        approvals = sorted((s.term_plan for s in _advisees_of(user.id) if s.term_plan and s.term_plan.status == 'pending'),
                           key=lambda p: p.submitted_at or datetime.datetime.min)

    box = request.args.get('box', 'inbox')
    if box not in {'inbox', 'sent', 'approvals'} or (box == 'sent' and not can_send) or (box == 'approvals' and role != 'teacher'):
        box = 'inbox'
    if box == 'approvals':
        raw_id = (request.args.get('id') or '').strip()
        selected_plan = next((p for p in approvals if str(p.student_id) == raw_id), None) or (approvals[0] if approvals else None)
        return render_template(
            'announcements_inbox.html',
            box=box,
            approvals=approvals,
            selected_plan=selected_plan,
            received_count=len(received),
            sent_count=len(sent),
            approval_count=len(approvals),
            new_count=0,
            can_send=can_send,
            show_approvals=True,
        )
    items = sent if box == 'sent' else received

    seen_at = user.announcements_seen_at
    new_keys = {f'a{a.id}' for a in received_announcements if seen_at is None or (a.created_at and a.created_at > seen_at)}
    new_keys |= {f'm{m.id}' for m in received_messages if not m.is_read}

    # ?id= accepts "a12" / "m5"; a bare number means an announcement (older links).
    raw_id = (request.args.get('id') or '').strip()
    selected_key = f'a{raw_id}' if raw_id.isdigit() else raw_id
    selected = next((i for i in items if i['key'] == selected_key), None) or (items[0] if items else None)

    changed = False
    if box == 'inbox' and selected and selected['kind'] == 'message' and not selected['obj'].is_read:
        selected['obj'].is_read = True  # opening a message marks it read, like e-mail
        new_keys.discard(selected['key'])
        changed = True
    if box == 'inbox' and any(k.startswith('a') for k in new_keys):
        user.announcements_seen_at = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
        changed = True
    if changed:
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()

    return render_template(
        'announcements_inbox.html',
        box=box,
        items=items,
        selected=selected,
        new_keys=new_keys,
        received_count=len(received),
        sent_count=len(sent),
        new_count=len(new_keys),
        can_send=can_send,
        show_approvals=role == 'teacher',
        approval_count=len(approvals),
    )


@app.route('/messages/<int:message_id>/reply', methods=['POST'])
@login_required_session
@csrf_protect
def message_reply(message_id):
    """Reply to a personal message the current user received (users can't start threads with arbitrary people)."""
    user = db.session.get(User, session.get('user_id'))
    original = db.session.get(Message, message_id)
    if not user or not original or original.recipient_id != user.id:
        flash(_lang_text('Mesaj bulunamadı.', 'Message not found.'), 'warning')
        return redirect(url_for('announcements_inbox'))

    body = _normalize_form_text(request.form.get('body'), max_length=2000)
    if not body:
        flash(_lang_text('Yanıt boş olamaz.', 'The reply cannot be empty.'), 'warning')
        return redirect(url_for('announcements_inbox', id=f'm{original.id}'))

    subject = original.subject or ''
    if not subject.lower().startswith('re:'):
        subject = f'Re: {subject}'.strip()
    reply = Message(sender_id=user.id, recipient_id=original.sender_id, subject=subject[:150], body=body)
    db.session.add(reply)
    db.session.commit()
    flash(_lang_text('Yanıtınız gönderildi.', 'Your reply has been sent.'), 'success')
    return redirect(url_for('announcements_inbox', box='sent', id=f'm{reply.id}'))


@app.route('/reports')
@login_required_session
def reports():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    role = session.get('role')
    if role == 'teacher':
        permission_error = ensure_permission(PERMISSIONS['TEACHER_REPORT_READ'], 'teacher_dashboard')
        if permission_error:
            return permission_error

        classes = Class.query.filter_by(teacher_id=user.id).order_by(Class.name.asc()).all()
        rows = []
        for cls in classes:
            confirmed_sessions = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True).all()
            session_ids = [sess.id for sess in confirmed_sessions]
            total_sessions = len(session_ids)
            total_students = len(cls.students)
            present_total = 0
            if session_ids:
                present_total = AttendanceRecord.query.filter(
                    AttendanceRecord.session_id.in_(session_ids),
                    AttendanceRecord.present.is_(True),
                ).count()

            expected_total = total_sessions * total_students
            attendance_rate = (present_total / expected_total * 100.0) if expected_total else 0.0
            rows.append(
                {
                    'class_name': cls.name,
                    'student_count': total_students,
                    'confirmed_session_count': total_sessions,
                    'present_total': present_total,
                    'expected_total': expected_total,
                    'attendance_rate': round(attendance_rate, 2),
                }
            )

        totals = {
            'class_count': len(rows),
            'student_total': sum(row['student_count'] for row in rows),
            'session_total': sum(row['confirmed_session_count'] for row in rows),
            'present_total': sum(row['present_total'] for row in rows),
            'expected_total': sum(row['expected_total'] for row in rows),
        }
        totals['attendance_rate'] = round((totals['present_total'] / totals['expected_total'] * 100.0), 2) if totals['expected_total'] else 0.0

        return render_template('reports.html', report_role='teacher', user=user, rows=rows, totals=totals)

    if role == 'student':
        permission_error = ensure_permission(PERMISSIONS['STUDENT_REPORT_READ'], 'student_dashboard')
        if permission_error:
            return permission_error

        if not user.students:
            flash('Öğrenci profili bulunamadı. Lütfen destek ekibiyle iletişime geçin.', 'danger')
            return redirect(url_for('logout'))

        student = user.students[0]
        rows = []
        for cls in sorted(student.classes, key=lambda c: (c.name or '').lower()):
            confirmed_sessions = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True).all()
            session_ids = [sess.id for sess in confirmed_sessions]
            total_sessions = len(session_ids)
            present_count = 0
            if session_ids:
                present_count = AttendanceRecord.query.filter(
                    AttendanceRecord.student_id == student.id,
                    AttendanceRecord.session_id.in_(session_ids),
                    AttendanceRecord.present.is_(True),
                ).count()
            absence_count = total_sessions - present_count
            attendance_rate = (present_count / total_sessions * 100.0) if total_sessions else 0.0
            rows.append(
                {
                    'class_name': cls.name,
                    'total_sessions': total_sessions,
                    'present_count': present_count,
                    'absence_count': absence_count,
                    'attendance_rate': round(attendance_rate, 2),
                }
            )

        totals = {
            'class_count': len(rows),
            'total_sessions': sum(row['total_sessions'] for row in rows),
            'present_total': sum(row['present_count'] for row in rows),
            'absence_total': sum(row['absence_count'] for row in rows),
        }
        totals['attendance_rate'] = round((totals['present_total'] / totals['total_sessions'] * 100.0), 2) if totals['total_sessions'] else 0.0

        return render_template('reports.html', report_role='student', user=user, rows=rows, totals=totals)

    flash('Raporlar ekranı yalnızca öğretmen ve öğrenci hesapları içindir.', 'warning')
    return redirect(url_for('admin_dashboard'))


@app.route('/student/absence', methods=['GET', 'POST'])
@role_required('student')
@permission_required()
def student_absence():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    _auto_finalize_expired_sessions()

    if not user.students:
        flash('Öğrenci profili bulunamadı. Lütfen destek ekibiyle iletişime geçin.', 'danger')
        return redirect(url_for('logout'))
    student = user.students[0]

    if request.method == 'POST':
        permission_error = ensure_permission(PERMISSIONS['STUDENT_ABSENCE_UPDATE'], 'student_absence')
        if permission_error:
            return permission_error

        class_id = _normalize_form_text(request.form.get('class_id'), max_length=20)
        if not class_id.isdigit():
            flash('Geçersiz sınıf kimliği.', 'warning')
            return redirect(url_for('student_absence'))

        cls = db.session.get(Class, int(class_id))
        if not cls:
            flash('Sınıf bulunamadı.', 'warning')
            return redirect(url_for('student_absence'))

        if cls not in student.classes:
            student.classes.append(cls)
            db.session.commit()
            flash(f"Derse başarıyla katıldınız: {cls.name}", "success")
        else:
            flash("Derse katılım başarısız veya zaten kayıtlısınız.", "warning")
        return redirect(url_for('student_absence'))

    search_query = request.args.get('search')
    search_message = None
    filtered_classes = student.classes
    if search_query:
        normalized_query = search_query.strip().lower()
        if normalized_query:
            already_joined = any(normalized_query in cls.name.lower() for cls in student.classes)
            if already_joined:
                search_message = _t('already_enrolled_named', name=search_query)
            else:
                search_message = _t('no_class_found_named', name=search_query)

    classes = list(student.classes)
    class_ids = [cls.id for cls in classes]

    active_session_map = {}
    if class_ids:
        active_sessions = (
            AttendanceSession.query
            .filter(AttendanceSession.class_id.in_(class_ids), AttendanceSession.active == True)
            .order_by(AttendanceSession.class_id.asc(), AttendanceSession.date.desc())
            .all()
        )
        for active_session in active_sessions:
            if active_session.class_id not in active_session_map:
                active_session_map[active_session.class_id] = active_session

    confirmed_totals_by_class = {}
    if class_ids:
        confirmed_total_rows = (
            db.session.query(AttendanceSession.class_id, func.count(AttendanceSession.id))
            .filter(AttendanceSession.class_id.in_(class_ids), AttendanceSession.confirmed == True)
            .group_by(AttendanceSession.class_id)
            .all()
        )
        confirmed_totals_by_class = {class_id: total for class_id, total in confirmed_total_rows}

    attended_by_class = {}
    if class_ids:
        attended_rows = (
            db.session.query(AttendanceSession.class_id, func.count(AttendanceRecord.id))
            .join(AttendanceRecord, AttendanceRecord.session_id == AttendanceSession.id)
            .filter(
                AttendanceSession.class_id.in_(class_ids),
                AttendanceSession.confirmed == True,
                AttendanceRecord.student_id == student.id,
                AttendanceRecord.present == True,
            )
            .group_by(AttendanceSession.class_id)
            .all()
        )
        attended_by_class = {class_id: total for class_id, total in attended_rows}

    marked_in_active_session_ids = set()
    active_session_ids = [sess.id for sess in active_session_map.values()]
    if active_session_ids:
        marked_rows = (
            db.session.query(AttendanceRecord.session_id)
            .filter(
                AttendanceRecord.student_id == student.id,
                AttendanceRecord.session_id.in_(active_session_ids),
            )
            .all()
        )
        marked_in_active_session_ids = {session_id for session_id, in marked_rows}

    active_announcements = []
    for cls in classes:
        active_session = active_session_map.get(cls.id)
        if active_session:
            deadline = _session_deadline(active_session)
            active_announcements.append({
                'class_name': cls.name,
                'week': active_session.week,
                'deadline': deadline,
                'token': cls.qr_token,
            })

    class_data = []
    for cls in filtered_classes:
        total = int(confirmed_totals_by_class.get(cls.id, 0))
        attended = int(attended_by_class.get(cls.id, 0))

        active_session = active_session_map.get(cls.id)
        active_week = active_session.week if active_session else None
        already_marked = False
        if active_session:
            already_marked = active_session.id in marked_in_active_session_ids

        class_data.append({
            'class_id': cls.id,
            'class_name': cls.name,
            'total': total,
            'attended': attended,
            'percentage': (attended / total * 100) if total else 0,
            'status': _t('pass') if (total > 0 and (attended / total) >= 0.7) else _t('fail'),
            'risk': (total > 0 and (attended / total) < 0.7),
            'active': bool(active_session),
            'active_week': active_week,
            'already_marked': already_marked,
            'qr_token': cls.qr_token,
        })

    # Hide integration-test fixtures from join list (example.com test users, Integration Class names).
    all_classes = (
        Class.query
        .join(User, Class.teacher_id == User.id)
        .filter(~Class.name.ilike('Integration Class %'))
        .filter(~User.email.ilike('%@example.com'))
        .all()
    )
    joined_class_ids = {cls.id for cls in classes}
    available_classes = [cls for cls in all_classes if cls.id not in joined_class_ids]

    absence_summary = {
        'joined_class_count': len(student.classes),
        'active_session_count': len(active_announcements),
        'total_confirmed_sessions': sum(item['total'] for item in class_data),
        'average_absence_pct': (sum(item['percentage'] for item in class_data) / len(class_data)) if class_data else 0,
    }

    return render_template(
        'student_absence.html',
        user_name=user.name,
        class_data=class_data,
        available_classes=available_classes,
        search_message=search_message,
        active_announcements=active_announcements,
        absence_summary=absence_summary,
    )


@app.route('/student/identity')
@role_required('student')
@permission_required()
def student_identity_info():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))
    student = user.students[0]

    first_name = (student.first_name or '').strip() if hasattr(student, 'first_name') else ''
    last_name = (student.last_name or '').strip() if hasattr(student, 'last_name') else ''
    if not first_name and not last_name:
        name_parts = [part for part in (user.name or '').strip().split(' ') if part]
        first_name = name_parts[0] if name_parts else _t('not_set_yet')
        last_name = ' '.join(name_parts[1:]) if len(name_parts) > 1 else _t('not_set_yet')

    gender = (student.gender or '').strip() if hasattr(student, 'gender') else ''
    birth_place = (student.birth_place or '').strip() if hasattr(student, 'birth_place') else ''
    birth_date_text = _t('not_set_yet')
    if getattr(student, 'birth_date', None):
        birth_date_text = student.birth_date.strftime('%d-%m-%Y')

    identity_details = [
        (_t('first_name'), first_name or _t('not_set_yet')),
        (_t('last_name'), last_name or _t('not_set_yet')),
        (_t('gender'), gender or _t('not_set_yet')),
        (_t('birth_date'), birth_date_text),
        (_t('birth_place'), birth_place or _t('not_set_yet')),
    ]

    identity_number = (student.identity_number or '').strip() if hasattr(student, 'identity_number') else ''
    nationality = (student.nationality or '').strip() if hasattr(student, 'nationality') else ''
    registered_city = (student.registered_city or '').strip() if hasattr(student, 'registered_city') else ''
    registered_district = (student.registered_district or '').strip() if hasattr(student, 'registered_district') else ''
    passport_number = (student.passport_number or '').strip() if hasattr(student, 'passport_number') else ''
    passport_active = bool(getattr(student, 'passport_active', False))
    passport_issue_place = (student.passport_issue_place or '').strip() if hasattr(student, 'passport_issue_place') else ''

    passport_issue_date_text = _t('not_set_yet')
    if getattr(student, 'passport_issue_date', None):
        passport_issue_date_text = student.passport_issue_date.strftime('%d-%m-%Y')

    passport_expiry_date_text = _t('not_set_yet')
    if getattr(student, 'passport_expiry_date', None):
        passport_expiry_date_text = student.passport_expiry_date.strftime('%d-%m-%Y')

    marital_status = (student.marital_status or '').strip() if hasattr(student, 'marital_status') else ''
    blood_type = (student.blood_type or '').strip() if hasattr(student, 'blood_type') else ''
    is_veteran_martyr_relative = bool(getattr(student, 'is_veteran_martyr_relative', False))
    is_disabled = bool(getattr(student, 'is_disabled', False))
    disability_type = (student.disability_type or '').strip() if hasattr(student, 'disability_type') else ''
    disability_rate = (student.disability_rate or '').strip() if hasattr(student, 'disability_rate') else ''
    is_employed = bool(getattr(student, 'is_employed', False))
    is_group_company = bool(getattr(student, 'is_group_company', False))
    company_name = (student.company_name or '').strip() if hasattr(student, 'company_name') else ''
    work_type = (student.work_type or '').strip() if hasattr(student, 'work_type') else ''

    employment_start_date_text = _t('not_set_yet')
    if getattr(student, 'employment_start_date', None):
        employment_start_date_text = student.employment_start_date.strftime('%d-%m-%Y')

    citizenship_details = [
        (_t('identity_number'), identity_number or _t('not_set_yet')),
        (_t('nationality'), nationality or _t('not_set_yet')),
        (_t('registered_city'), registered_city or _t('not_set_yet')),
        (_t('registered_district'), registered_district or _t('not_set_yet')),
    ]

    passport_summary = passport_number or _t('not_set_yet')
    passport_details = [
        (_t('active_status_box'), passport_active),
        (_t('issued_date'), passport_issue_date_text),
        (_t('issued_place'), passport_issue_place or _t('not_set_yet')),
        (_t('valid_until'), passport_expiry_date_text),
    ]

    other_details = [
        (_t('marital_status'), marital_status or _t('not_set_yet')),
        (_t('blood_type'), blood_type or _t('not_set_yet')),
        (_t('veteran_martyr_relative_q'), _t('yes') if is_veteran_martyr_relative else _t('no')),
        (_t('disabled_q'), _t('yes') if is_disabled else _t('no')),
        (_t('disability_type'), disability_type or _t('not_set_yet')),
        (_t('disability_rate'), disability_rate or _t('not_set_yet')),
    ]

    working_details = [
        (_t('is_employed_q'), _t('yes') if is_employed else _t('no')),
        (_t('is_group_company_q'), _t('yes') if is_group_company else _t('no')),
        (_t('company_name'), company_name or _t('not_set_yet')),
        (_t('work_type'), work_type or _t('not_set_yet')),
        (_t('employment_start_date'), employment_start_date_text),
    ]

    return render_template(
        'student_identity.html',
        title=_t('identity_information'),
        identity_details=identity_details,
        citizenship_title=_t('citizenship'),
        citizenship_details=citizenship_details,
        passport_title=_t('passport_number'),
        passport_summary=passport_summary,
        passport_details=passport_details,
        other_title=_t('other_information'),
        other_details=other_details,
        working_title=_t('working_information'),
        working_details=working_details,
    )


@app.route('/student/education')
@role_required('student')
@permission_required()
def student_education_info():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))
    student = user.students[0]

    details = [
        (_t('entry_place'), (student.university_entry_place or '').strip() or _t('not_set_yet')),
        (_t('entry_type'), (student.university_entry_type or '').strip() or _t('not_set_yet')),
        (_t('academic_year'), (student.university_academic_year or '').strip() or _t('not_set_yet')),
        (_t('term'), (student.university_term or '').strip() or _t('not_set_yet')),
        (_t('faculty'), (student.university_faculty or '').strip() or _t('not_set_yet')),
        (_t('department'), (student.university_department or '').strip() or _t('not_set_yet')),
        (_t('scholarship_type'), (student.university_scholarship_type or '').strip() or _t('not_set_yet')),
        (_t('placement_type'), (student.university_placement_type or '').strip() or _t('not_set_yet')),
        (_t('score_type'), (student.university_score_type or '').strip() or _t('not_set_yet')),
        (_t('achievement_score'), (student.university_achievement_score or '').strip() or _t('not_set_yet')),
        (_t('placement_score'), (student.university_placement_score or '').strip() or _t('not_set_yet')),
        (_t('preference_order'), (student.university_preference_order or '').strip() or _t('not_set_yet')),
    ]

    highschool_graduation_date = (
        student.highschool_graduation_date.strftime('%Y-%m-%d')
        if student.highschool_graduation_date
        else _t('not_set_yet')
    )
    highschool_details = [
        (_t('highschool_name'), (student.highschool_name or '').strip() or _t('not_set_yet')),
        (_t('highschool_info'), (student.highschool_info or '').strip() or _t('not_set_yet')),
        (_t('highschool_graduation_date'), highschool_graduation_date),
    ]

    return render_template(
        'student_education.html',
        title=_t('education_information'),
        section_title=_t('university_information'),
        highschool_title=_t('highschool_information'),
        details=details,
        highschool_details=highschool_details,
    )


@app.route('/student/family')
@role_required('student')
@permission_required()
def student_family_info():
    details = [
        (_t('guardian_full_name'), _t('not_set_yet')),
        (_t('relationship'), _t('not_set_yet')),
        (_t('emergency_phone'), _t('not_set_yet')),
        (_t('note'), _t('family_note')),
    ]
    return render_template('student_module_page.html', title=_t('family_information'), details=details)


@app.route('/student/documents')
@role_required('student')
@permission_required()
def student_documents_info():
    details = [
        (_t('document_status'), _t('no_uploaded_document')),
        (_t('registration_document'), _t('pending')),
        (_t('student_certificate'), _t('can_be_requested')),
        (_t('note'), _t('documents_note')),
    ]
    return render_template('student_module_page.html', title=_t('documents'), details=details)


@app.route('/student/contact')
@role_required('student')
@permission_required()
def student_contact_info():
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))
    details = [
        (_t('email'), user.email),
        (_t('phone'), _t('not_set_yet')),
        (_t('address'), _t('not_set_yet')),
        (_t('contact_preference'), _t('email')),
    ]
    return render_template('student_module_page.html', title=_t('contact'), details=details)


@app.route('/student/current-account')
@role_required('student')
@permission_required()
def student_current_account():
    details = [
        (_t('total_debt'), '0.00'),
        (_t('total_paid'), '0.00'),
        (_t('current_balance'), '0.00'),
        (_t('note'), _t('current_account_note')),
    ]
    return render_template('student_module_page.html', title=_t('current_account'), details=details)


@app.route('/student/payment')
@role_required('student')
@permission_required()
def student_payment_info():
    details = [
        (_t('last_payment'), _t('no_record_found')),
        (_t('payment_method'), _t('undefined')),
        (_t('scheduled_payment'), _t('none')),
        (_t('note'), _t('payment_note')),
    ]
    return render_template('student_module_page.html', title=_t('payments'), details=details)


def _term_plan_of(student):
    """The student's term plan; an unsaved draft until they first submit."""
    return student.term_plan or TermPlan(student_id=student.id, status='draft')


def _term_plan_student_can_edit(student, plan):
    """Students edit only an unsubmitted (or rejected) plan, and only once they have an advisor."""
    return plan.status in ('draft', 'rejected') and student.advisor_info()[0] is not None


def _add_drop_window_open(policy):
    now = datetime.datetime.now()
    return bool(policy) and not policy.is_locked \
        and (policy.add_drop_start is None or now >= policy.add_drop_start) \
        and (policy.add_drop_end is None or now <= policy.add_drop_end)


@app.route('/student/term-courses')
@role_required('student')
@permission_required()
def student_term_courses():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    def _get_or_create_registration_policy():
        policy = CourseRegistrationPolicy.query.first()
        if policy:
            return policy
        policy = CourseRegistrationPolicy(
            add_drop_start=None,
            add_drop_end=None,
            min_credits=0,
            max_credits=45,
            is_locked=False,
        )
        db.session.add(policy)
        db.session.commit()
        return policy

    def _is_add_drop_open(policy):
        if not policy or policy.is_locked:
            return False
        now = datetime.datetime.now()
        if policy.add_drop_start and now < policy.add_drop_start:
            return False
        if policy.add_drop_end and now > policy.add_drop_end:
            return False
        return True

    def _ensure_course_catalog_from_classes():
        created = False
        for cls in Class.query.all():
            code = f'CLS{cls.id:04d}'
            existing = Course.query.filter_by(code=code).first()
            if existing:
                continue
            db.session.add(Course(
                code=code,
                title=(cls.name or 'Course')[:120],
                credit=DEFAULT_COURSE_CREDIT, ects=DEFAULT_COURSE_ECTS,
                capacity=60,
                schedule_slot=None,
                teacher_id=cls.teacher_id,
                term_id=None,
            ))
            created = True
        if created:
            db.session.commit()

    student = user.students[0]
    _ensure_course_catalog_from_classes()
    policy = _get_or_create_registration_policy()
    is_open = _is_add_drop_open(policy) or _view_as_can_edit()

    enrollments = (
        CourseEnrollment.query
        .join(Course, CourseEnrollment.course_id == Course.id)
        .filter(CourseEnrollment.student_id == student.id)
        .order_by(Course.code.asc())
        .all()
    )
    enrolled_course_ids = {en.course_id for en in enrollments}
    enrolled_credits = sum((en.course.ects or 0) for en in enrollments if en.course)

    available_courses = (
        Course.query
        .order_by(Course.code.asc())
        .all()
    )

    course_student_counts = {
        course_id: cnt
        for course_id, cnt in (
            db.session.query(CourseEnrollment.course_id, func.count(CourseEnrollment.id))
            .group_by(CourseEnrollment.course_id)
            .all()
        )
    }

    plan = _term_plan_of(student)
    advisor_mode = _view_as_can_edit()
    return render_template(
        'student_term_courses.html',
        student=student,
        policy=policy,
        is_open=is_open,
        plan=plan,
        advisor=student.advisor_info()[0],
        advisor_mode=advisor_mode,
        can_edit=is_open and (advisor_mode or _term_plan_student_can_edit(student, plan)),
        enrollments=enrollments,
        enrolled_course_ids=enrolled_course_ids,
        enrolled_credits=enrolled_credits,
        available_courses=available_courses,
        course_student_counts=course_student_counts,
    )


@app.route('/student/term-courses/update', methods=['POST'])
@role_required('student')
@permission_required()
def student_term_courses_update():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    student = user.students[0]
    # An advisor editing their own advisee's courses from the teacher's student view.
    advisor_mode = _view_as_can_edit()
    actor_id = session['view_as_origin']['user_id'] if advisor_mode else user.id
    actor_label = 'advisor' if advisor_mode else 'student-self'
    action = _normalize_form_text(request.form.get('action'), max_length=20)
    course_id_raw = _normalize_form_text(request.form.get('course_id'), max_length=20)

    if action not in {'add', 'drop'} or not course_id_raw.isdigit():
        flash('Geçersiz ders güncelleme isteği.', 'danger')
        return redirect(url_for('student_term_courses'))

    course = db.session.get(Course, int(course_id_raw))
    if not course:
        flash('Ders bulunamadı.', 'danger')
        return redirect(url_for('student_term_courses'))

    policy = CourseRegistrationPolicy.query.first()
    if not policy:
        policy = CourseRegistrationPolicy(min_credits=0, max_credits=45, is_locked=False)
        db.session.add(policy)
        db.session.commit()

    if not (advisor_mode or _add_drop_window_open(policy)):
        flash('Ders ekleme-bırakma penceresi şu an kapalı.', 'warning')
        return redirect(url_for('student_term_courses'))

    plan = _term_plan_of(student)
    if not advisor_mode:
        if student.advisor_info()[0] is None:
            flash(_lang_text('Danışmanınız atanmadığı için ders kaydı yapamazsınız. Öğrenci işleriyle iletişime geçin.',
                             'You cannot register for courses until an advisor is assigned. Contact Student Affairs.'), 'warning')
            return redirect(url_for('student_term_courses'))
        if not _term_plan_student_can_edit(student, plan):
            flash(_lang_text('Programınız danışmanınıza gönderildi; değişiklikleri artık danışmanınız yapabilir.',
                             'Your plan has been sent to your advisor; only your advisor can change it now.'), 'warning')
            return redirect(url_for('student_term_courses'))

    enrollment = CourseEnrollment.query.filter_by(student_id=student.id, course_id=course.id).first()

    if action == 'add':
        if enrollment:
            flash('Bu derse zaten kayıtlısınız.', 'warning')
            return redirect(url_for('student_term_courses'))

        current_credits = (
            db.session.query(func.coalesce(func.sum(Course.ects), 0))
            .join(CourseEnrollment, Course.id == CourseEnrollment.course_id)
            .filter(CourseEnrollment.student_id == student.id)
            .scalar() or 0
        )
        if current_credits + (course.ects or 0) > (policy.max_credits or 45):
            flash(f'Maksimum AKTS limitini ({policy.max_credits or 45} AKTS) aştınız.', 'warning')
            return redirect(url_for('student_term_courses'))

        if course.schedule_slot:
            clash_exists = (
                db.session.query(Course.id)
                .join(CourseEnrollment, Course.id == CourseEnrollment.course_id)
                .filter(
                    CourseEnrollment.student_id == student.id,
                    Course.schedule_slot == course.schedule_slot,
                )
                .first()
            )
            if clash_exists:
                flash(f'Ders programı çakışması: {course.schedule_slot} saatinde başka bir dersiniz var.', 'warning')
                return redirect(url_for('student_term_courses'))

        if course.capacity and course.capacity > 0:
            enrolled_count = CourseEnrollment.query.filter_by(course_id=course.id).count()
            if enrolled_count >= course.capacity:
                flash(f'{course.code} dersinin kontenjanı doldu ({course.capacity} kişi).', 'warning')
                return redirect(url_for('student_term_courses'))

        try:
            # An advisor adding to an already approved plan makes the course official right away.
            db.session.add(CourseEnrollment(student_id=student.id, course_id=course.id,
                                            approved=advisor_mode and plan.status == 'approved'))
            db.session.flush()
            db.session.add(CourseEnrollmentAudit(
                actor_user_id=actor_id,
                student_id=student.id,
                course_id=course.id,
                action='add',
                detail=f'{actor_label} add {course.code}',
            ))
            db.session.commit()
            flash(f'{course.code} dersi başarıyla eklendi.', 'success')
        except Exception:
            db.session.rollback()
            flash('Ders eklenirken bir hata oluştu. Lütfen tekrar deneyin.', 'danger')
        return redirect(url_for('student_term_courses'))

    if not enrollment:
        flash('Bu derse kayıtlı değilsiniz.', 'warning')
        return redirect(url_for('student_term_courses'))

    current_credits = (
        db.session.query(func.coalesce(func.sum(Course.ects), 0))
        .join(CourseEnrollment, Course.id == CourseEnrollment.course_id)
        .filter(CourseEnrollment.student_id == student.id)
        .scalar() or 0
    )
    remaining_credits = current_credits - (course.ects or 0)
    if remaining_credits < (policy.min_credits or 0):
        flash(f'Bu dersi bırakırsanız minimum AKTS limitinin ({policy.min_credits or 0} AKTS) altına düşersiniz.', 'warning')
        return redirect(url_for('student_term_courses'))

    try:
        db.session.delete(enrollment)
        db.session.flush()
        db.session.add(CourseEnrollmentAudit(
            actor_user_id=actor_id,
            student_id=student.id,
            course_id=course.id,
            action='drop',
            detail=f'{actor_label} drop {course.code}',
        ))
        db.session.commit()
        flash(f'{course.code} dersi başarıyla bırakıldı.', 'success')
    except Exception:
        db.session.rollback()
        flash('Ders bırakılırken bir hata oluştu. Lütfen tekrar deneyin.', 'danger')
    return redirect(url_for('student_term_courses'))


@app.route('/student/term-plan/submit', methods=['POST'])
@role_required('student')
@permission_required()
@csrf_protect
def student_term_plan_submit():
    """The student sends their course selection to the advisor; the plan is locked until a decision."""
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        return redirect(url_for('login'))
    student = user.students[0]
    plan = _term_plan_of(student)
    policy = CourseRegistrationPolicy.query.first()
    credits = sum((e.course.ects or 0) for e in student.course_enrollments if e.course)

    if student.advisor_info()[0] is None:
        flash(_lang_text('Danışmanınız atanmadığı için programınızı gönderemezsiniz.',
                         'You cannot submit your plan until an advisor is assigned.'), 'warning')
    elif plan.status not in ('draft', 'rejected'):
        flash(_lang_text('Programınız zaten gönderildi.', 'Your plan has already been submitted.'), 'info')
    elif not _add_drop_window_open(policy):
        flash('Ders ekleme-bırakma penceresi şu an kapalı.', 'warning')
    elif not student.course_enrollments:
        flash(_lang_text('Göndermeden önce en az bir ders seçin.', 'Select at least one course before submitting.'), 'warning')
    elif policy and credits < (policy.min_credits or 0):
        flash(_lang_text(f'Minimum AKTS ({policy.min_credits}) sağlanmadan program gönderilemez.',
                         f'Your plan needs at least {policy.min_credits} ECTS.'), 'warning')
    else:
        plan.status = 'pending'
        plan.submitted_at = datetime.datetime.now()
        plan.decided_at = plan.decided_by_id = plan.rejection_reason = None
        db.session.add(plan)
        db.session.commit()
        flash(_lang_text('Programınız danışman onayına gönderildi. Danışmanınız onayladığında ya da reddettiğinde mesaj kutunuza bildirim gelecek.',
                         'Your plan was sent to your advisor for approval. You will get a message in your inbox when they approve or reject it.'),
              'popup-success')
    return redirect(url_for('student_term_courses'))


def _approve_term_plan(student, plan, advisor):
    """Make the plan's enrollments official and tell the student (commits)."""
    plan.status = 'approved'
    plan.rejection_reason = None
    plan.decided_at = datetime.datetime.now()
    plan.decided_by_id = advisor.id
    for enrollment in student.course_enrollments:
        enrollment.approved = True
    db.session.add(Message(
        sender_id=advisor.id, recipient_id=student.id,
        subject=_lang_text('Ders programınız onaylandı', 'Your course plan was approved'),
        body=_lang_text('Danışmanınız dönem ders programınızı onayladı.', 'Your advisor approved your course plan for the term.'),
    ))
    db.session.commit()


@app.route('/view-as/term-plan/approve', methods=['POST'])
@login_required_session
@csrf_protect
def view_as_term_plan_approve():
    """Advisor approves the plan from the advisee's course page, then returns to the approvals box."""
    if not _view_as_can_edit():
        flash(_t('view_as_read_only'), 'warning')
        return redirect(url_for('student_term_courses'))
    student = db.session.get(Student, session.get('user_id'))
    plan = student.term_plan if student else None
    _end_view_as()
    if not plan or plan.status != 'pending':
        flash(_lang_text('Onay bekleyen program bulunamadı.', 'No plan awaiting approval was found.'), 'warning')
        return redirect(url_for('announcements_inbox', box='approvals'))
    _approve_term_plan(student, plan, db.session.get(User, session.get('user_id')))
    flash(_lang_text(f'{student.user.name} adlı öğrencinin ders programı onaylandı; öğrenciye bildirim gönderildi.',
                     f"{student.user.name}'s course plan was approved and the student was notified."), 'popup-success')
    return redirect(url_for('announcements_inbox', box='approvals'))


@app.route('/teacher/term-plans/<int:student_id>/decision', methods=['POST'])
@role_required('teacher')
@permission_required()
@csrf_protect
def teacher_term_plan_decision(student_id):
    """Advisor approves (enrollments become official) or rejects (with a reason) a submitted plan."""
    student = _advisee_or_none(student_id)
    plan = student.term_plan if student else None
    if not plan or plan.status != 'pending':
        flash(_lang_text('Onay bekleyen program bulunamadı.', 'No plan awaiting approval was found.'), 'warning')
        return redirect(url_for('announcements_inbox', box='approvals'))

    action = _normalize_form_text(request.form.get('action'), max_length=20)
    reason = (request.form.get('reason') or '').strip()[:500]
    advisor = db.session.get(User, session.get('user_id'))
    if action == 'approve':
        _approve_term_plan(student, plan, advisor)
        flash(_lang_text('Karar öğrenciye bildirildi.', 'The student has been notified.'), 'success')
        return redirect(url_for('announcements_inbox', box='approvals'))
    elif action == 'reject':
        if not reason:
            flash(_lang_text('Reddetmek için bir sebep yazın.', 'Write a reason to reject.'), 'warning')
            return redirect(url_for('announcements_inbox', box='approvals', id=student.id))
        plan.status = 'rejected'
        plan.rejection_reason = reason
        subject = _lang_text('Ders programınız reddedildi', 'Your course plan was rejected')
        body = _lang_text('Danışmanınız dönem ders programınızı reddetti. Sebep:\n\n',
                          'Your advisor rejected your course plan for the term. Reason:\n\n') + reason + \
            _lang_text('\n\nDüzeltip yeniden gönderebilirsiniz.', '\n\nYou can correct it and submit again.')
    else:
        flash(_lang_text('Geçersiz işlem.', 'Invalid action.'), 'danger')
        return redirect(url_for('announcements_inbox', box='approvals'))

    plan.decided_at = datetime.datetime.now()
    plan.decided_by_id = advisor.id
    db.session.add(Message(sender_id=advisor.id, recipient_id=student.id, subject=subject, body=body))
    db.session.commit()
    flash(_lang_text('Karar öğrenciye bildirildi.', 'The student has been notified.'), 'success')
    return redirect(url_for('announcements_inbox', box='approvals'))


@app.route('/admin/course-registration-window', methods=['GET', 'POST'])
@role_required('admin')
@permission_required()
def admin_course_registration_window():
    created = False
    for cls in Class.query.all():
        code = f'CLS{cls.id:04d}'
        if Course.query.filter_by(code=code).first():
            continue
        db.session.add(Course(
            code=code,
            title=(cls.name or 'Course')[:120],
            credit=DEFAULT_COURSE_CREDIT, ects=DEFAULT_COURSE_ECTS,
            capacity=60,
            schedule_slot=None,
            teacher_id=cls.teacher_id,
            term_id=None,
        ))
        created = True
    if created:
        db.session.commit()

    policy = CourseRegistrationPolicy.query.first()
    if not policy:
        policy = CourseRegistrationPolicy(min_credits=0, max_credits=45, is_locked=False)
        db.session.add(policy)
        db.session.commit()

    if request.method == 'POST':
        action = _normalize_form_text(request.form.get('action'), max_length=30)
        if action == 'save_policy':
            start_raw = _normalize_form_text(request.form.get('add_drop_start'), max_length=25)
            end_raw = _normalize_form_text(request.form.get('add_drop_end'), max_length=25)
            min_credits_raw = _normalize_form_text(request.form.get('min_credits'), max_length=10)
            max_credits_raw = _normalize_form_text(request.form.get('max_credits'), max_length=10)
            is_locked = request.form.get('is_locked') == 'on'

            try:
                policy.add_drop_start = datetime.datetime.fromisoformat(start_raw) if start_raw else None
                policy.add_drop_end = datetime.datetime.fromisoformat(end_raw) if end_raw else None
                policy.min_credits = int(min_credits_raw or 0)
                policy.max_credits = int(max_credits_raw or 45)
                policy.is_locked = is_locked
                if policy.max_credits < policy.min_credits:
                    flash('Maksimum AKTS, minimum AKTS değerinden az olamaz.', 'warning')
                else:
                    db.session.commit()
                    flash('Ders kayıt penceresi politikası güncellendi.', 'success')
            except Exception:
                db.session.rollback()
                flash('Politika güncellenirken hata oluştu.', 'danger')

    students = User.query.filter_by(role='student').order_by(User.name.asc()).limit(500).all()
    courses = Course.query.order_by(Course.code.asc()).limit(500).all()
    recent_audits = CourseEnrollmentAudit.query.order_by(CourseEnrollmentAudit.created_at.desc()).limit(80).all()

    return render_template(
        'admin_course_registration_window.html',
        policy=policy,
        students=students,
        courses=courses,
        recent_audits=recent_audits,
    )


@app.route('/admin/course-registration-window/override', methods=['POST'])
@role_required('admin')
@permission_required()
def admin_course_registration_override():
    admin_user = db.session.get(User, session.get('user_id'))
    student_id_raw = _normalize_form_text(request.form.get('student_id'), max_length=20)
    course_id_raw = _normalize_form_text(request.form.get('course_id'), max_length=20)
    action = _normalize_form_text(request.form.get('action'), max_length=20)

    if action not in {'add', 'drop'} or not student_id_raw.isdigit() or not course_id_raw.isdigit():
        flash('Geçersiz override isteği.', 'danger')
        return redirect(url_for('admin_course_registration_window'))

    student = db.session.get(Student, int(student_id_raw))
    course = db.session.get(Course, int(course_id_raw))
    if not student or not course:
        flash('Öğrenci veya ders bulunamadı.', 'danger')
        return redirect(url_for('admin_course_registration_window'))

    enrollment = CourseEnrollment.query.filter_by(student_id=student.id, course_id=course.id).first()

    try:
        if action == 'add':
            if not enrollment:
                db.session.add(CourseEnrollment(student_id=student.id, course_id=course.id, approved=True))
                db.session.flush()
            db.session.add(CourseEnrollmentAudit(
                actor_user_id=admin_user.id,
                student_id=student.id,
                course_id=course.id,
                action='add_override',
                detail=f'admin override add {course.code}',
            ))
            db.session.commit()
            flash(f'Admin ekleme tamamlandı: {course.code}.', 'success')
            return redirect(url_for('admin_course_registration_window'))

        if enrollment:
            db.session.delete(enrollment)
            db.session.flush()
        db.session.add(CourseEnrollmentAudit(
            actor_user_id=admin_user.id,
            student_id=student.id,
            course_id=course.id,
            action='drop_override',
            detail=f'admin override drop {course.code}',
        ))
        db.session.commit()
        flash(f'Admin bırakma tamamlandı: {course.code}.', 'success')
    except Exception:
        db.session.rollback()
        flash('Override işlemi başarısız. Lütfen tekrar deneyin.', 'danger')
    return redirect(url_for('admin_course_registration_window'))


@app.route('/student/transcript')
@role_required('student')
@permission_required()
def student_transcript():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    student = user.students[0]
    gpa_4_scale = 0.0
    weighted_credits = 0
    transcript_note = _t('no_graded_data')

    try:
        enrollments = CourseEnrollment.query.filter_by(student_id=student.id).all()
        total_weighted_points = 0.0
        graded_course_count = 0

        for enrollment in enrollments:
            if not enrollment.grades:
                continue
            course_credit = enrollment.course.ects if enrollment.course and enrollment.course.ects else 0
            if course_credit <= 0:
                continue
            avg_grade_point = sum(g.grade_point for g in enrollment.grades) / len(enrollment.grades)
            total_weighted_points += avg_grade_point * course_credit
            weighted_credits += course_credit
            graded_course_count += 1

        if weighted_credits > 0:
            gpa_4_scale = total_weighted_points / weighted_credits
            transcript_note = _t('calculated_from_courses', count=graded_course_count)
    except Exception:
        db.session.rollback()
        transcript_note = _t('transcript_table_not_initialized')

    details = [
        (_t('student_label'), user.name),
        (_t('gpa_4'), f'{gpa_4_scale:.2f}'),
        (_t('weighted_credits'), str(weighted_credits)),
        (_t('note'), transcript_note),
    ]
    return render_template('student_module_page.html', title=_t('transcript'), details=details)


@app.route('/student/academic-calendar')
@role_required('student')
@permission_required()
def student_academic_calendar():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    student = user.students[0]
    now = datetime.datetime.now()
    class_ids = [cls.id for cls in student.classes]

    upcoming_count = 0
    if class_ids:
        upcoming_count = AttendanceSession.query.filter(
            AttendanceSession.class_id.in_(class_ids),
            AttendanceSession.date >= now,
        ).count()

    details = [
        (_t('today'), now.strftime('%d-%m-%Y')),
        (_t('upcoming_academic_events'), str(upcoming_count)),
        (_t('term_window'), _t('configured_by_admin')),
        (_t('note'), _t('academic_calendar_note')),
    ]
    return render_template('student_module_page.html', title=_t('academic_calendar'), details=details)


@app.route('/student/exams')
@role_required('student')
@permission_required()
def student_exams():
    user = db.session.get(User, session.get('user_id'))
    if not user or not user.students:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))

    details = [
        (_t('upcoming_exams'), '0'),
        (_t('completed_exams'), '0'),
        (_t('next_exam_date'), _t('not_scheduled')),
        (_t('note'), _t('exam_module_note')),
    ]
    return render_template('student_module_page.html', title=_t('exams'), details=details)


@app.route('/student/history/<int:class_id>')
@role_required('student')
@permission_required()
def student_class_history(class_id):
    user = db.session.get(User, session.get('user_id'))
    if not user:
        session.clear()
        flash('Oturum süresi doldu. Lütfen tekrar giriş yapın.', 'warning')
        return redirect(url_for('login'))
    if not user.students:
        flash('Öğrenci profili bulunamadı. Lütfen destek ekibiyle iletişime geçin.', 'danger')
        return redirect(url_for('logout'))
    student = user.students[0]
    cls = Class.query.get_or_404(class_id)
    denied = ensure_student_class_membership(
        student,
        cls,
        on_fail='student_absence',
        fail_message='You are not authorized to view this class history.',
    )
    if denied:
        return denied

    sessions = AttendanceSession.query.filter_by(class_id=cls.id, confirmed=True).order_by(AttendanceSession.date.desc()).all()
    history_rows = []
    for sess in sessions:
        present = AttendanceRecord.query.filter_by(session_id=sess.id, student_id=student.id, present=True).first() is not None
        history_rows.append({
            'date': sess.date,
            'week': sess.week,
            'present': present,
        })

    return render_template('student_class_history.html', cls=cls, history_rows=history_rows)

# ------------------ CREATE CLASS ------------------
@app.route('/teacher/create_class', methods=['GET', 'POST'])
@role_required('teacher')
@permission_required()
def create_class():
    if request.method == 'POST':
        permission_error = ensure_permission(PERMISSIONS['TEACHER_CLASS_CREATE'], 'create_class')
        if permission_error:
            return permission_error

        class_name = (request.form.get('class_name') or '').strip()
        if not class_name:
            flash("Class name cannot be empty.", "danger")
            return redirect(url_for('create_class'))

        if len(class_name) > 120:
            flash("Class name is too long.", "danger")
            return redirect(url_for('create_class'))

        existing = Class.query.filter_by(name=class_name).first()
        if existing:
            flash("A class with this name already exists.", "warning")
            return redirect(url_for('create_class'))

        # There is no separate teacher model, so we use user.id.
        # Create one fixed QR code per class.
        qr_token = secrets.token_hex(4)
        qr_filename = f"{class_name}_qr_{qr_token}.png"

        qr_folder = os.path.join(app.root_path, "static", "qrcodes")
        if not os.path.exists(qr_folder):
            os.makedirs(qr_folder)

        qr_path = os.path.join(qr_folder, qr_filename)
        img = qrcode.make(_attendance_url(qr_token))
        img.save(qr_path)

        new_class = Class(name=class_name, teacher_id=session['user_id'], qr_token=qr_token, qr_filename=qr_filename)
        db.session.add(new_class)
        db.session.commit()
        flash("Class created.", "success")
        return redirect(url_for('teacher_dashboard'))

    return render_template('create_class.html')

# ------------------ START ATTENDANCE SESSION ------------------
@app.route('/create_session', methods=['POST'])
@role_required(('teacher', 'assistant'))
@permission_required()
def create_session():
    class_id_raw = _normalize_form_text(request.form.get('class_id'), max_length=20)
    week = _normalize_form_text(request.form.get('week'), max_length=20)
    action = _normalize_form_text(request.form.get('action', 'start'), max_length=30)

    if not class_id_raw.isdigit():
        flash("Geçersiz sınıf kimliği.", "danger")
        return redirect(url_for('teacher_dashboard'))

    if not week:
        # For QR-open flow, allow quick launch by assigning today's label.
        if action == 'start_and_show_qr':
            week = datetime.datetime.now().strftime('%Y-%m-%d')
        else:
            flash("Week is required.", "danger")
            return redirect(url_for('teacher_dashboard'))

    if not re.match(r'^[A-Za-z0-9._ -]+$', week):
        flash("Week contains invalid characters.", "danger")
        return redirect(url_for('teacher_dashboard'))

    if action not in {'start', 'start_and_show_qr'}:
        action = 'start'

    class_id = int(class_id_raw) if class_id_raw and class_id_raw.isdigit() else None
    cls = db.session.get(Class, class_id) if class_id is not None else None
    if not cls:
        flash("Sınıf bulunamadı.", "danger")
        return redirect(url_for('teacher_dashboard'))
    denied = ensure_teacher_class_ownership(
        cls,
        on_fail='teacher_dashboard',
        fail_message='You are not allowed to start attendance for this class.',
    )
    if denied:
        return denied

    # Deactivate previous active sessions for this class.
    AttendanceSession.query.filter_by(class_id=class_id, active=True).update({'active': False})

    new_session = AttendanceSession(
        class_id=class_id,
        date=datetime.datetime.now(),
        qr_token=secrets.token_hex(4),
        qr_filename='',
        name=None,
        week=week,
        active=True
    )
    db.session.add(new_session)
    db.session.commit()

    flash("Attendance session started.", "success")
    if action == 'start_and_show_qr':
        return redirect(url_for('view_qr', token=cls.qr_token))
    return redirect(url_for('teacher_dashboard'))

# ------------------ VIEW QR ------------------
@app.route('/view_qr/<token>')
@role_required(('teacher', 'assistant'))
@permission_required()
def view_qr(token):
    # Token is the class-level fixed QR token.
    cls = Class.query.filter_by(qr_token=token).first()
    if not cls:
        flash("QR not found.", "danger")
        return redirect(url_for('teacher_dashboard'))
    denied = ensure_teacher_class_ownership(
        cls,
        on_fail='teacher_dashboard',
        fail_message='You are not allowed to view this QR.',
    )
    if denied:
        return denied

    return render_template("view_qr.html", qr_filename=cls.qr_filename, token=token, attendance_url=_attendance_url(token))

# ------------------ CLASS DETAIL ------------------
@app.route('/teacher/class/<int:class_id>')
@role_required(('teacher', 'assistant'))
@permission_required()
def class_detail(class_id):
    cls = Class.query.get_or_404(class_id)
    denied = ensure_teacher_class_ownership(
        cls,
        on_fail='teacher_dashboard',
        fail_message='You are not allowed to view this class.',
    )
    if denied:
        return denied

    students_data = []
    session_ids = [session.id for session in cls.sessions]
    total_sessions = len(session_ids)
    attended_by_student = {}

    if session_ids:
        attended_rows = (
            db.session.query(AttendanceRecord.student_id, func.count(AttendanceRecord.id))
            .filter(
                AttendanceRecord.session_id.in_(session_ids),
                AttendanceRecord.present == True,
            )
            .group_by(AttendanceRecord.student_id)
            .all()
        )
        attended_by_student = {student_id: count for student_id, count in attended_rows}

    for student in cls.students:
        attended = int(attended_by_student.get(student.id, 0))

        absence = total_sessions - attended
        passed = (attended / total_sessions) >= 0.7 if total_sessions > 0 else False

        students_data.append({
            'name': student.user.name,
            'attended': attended,
            'absence': absence,
            'passed': passed
        })

    return render_template('class_detail.html', class_obj=cls, students_data=students_data)

# ------------------ ATTENDANCE ------------------
@app.route('/attendance/<token>')
@role_required('student')
@permission_required()
def mark_attendance(token):
    user_id = session['user_id']
    student = Student.query.get_or_404(user_id)

    # Token is the class-level fixed QR token.
    cls = Class.query.filter_by(qr_token=token).first()
    if not cls:
        flash('Geçersiz QR kodu!', 'danger')
        return redirect(url_for('student_absence'))
    denied = ensure_student_class_membership(
        student,
        cls,
        on_fail='student_absence',
        fail_message='You are not enrolled in this class.',
    )
    if denied:
        return denied

    active_session = AttendanceSession.query.filter_by(class_id=cls.id, active=True).order_by(AttendanceSession.date.desc()).first()
    if not active_session:
        flash('Bu sınıf için aktif yoklama oturumu yok. Öğretmeninizden başlatmasını isteyin.', 'warning')
        return redirect(url_for('student_absence'))

    # Reject scans outside the allowed time window (QR security layer).
    if datetime.datetime.now() > _session_deadline(active_session):
        active_session.active = False
        active_session.confirmed = True
        db.session.commit()
        flash('Yoklama süresi doldu. Oturum kapatıldı.', 'warning')
        return redirect(url_for('student_absence'))

    existing_record = AttendanceRecord.query.filter_by(
        student_id=student.id,
        session_id=active_session.id
    ).first()

    if existing_record:
        flash('Yoklamanız zaten işaretlendi!', 'info')
    else:
        new_record = AttendanceRecord(
            student_id=student.id,
            session_id=active_session.id,
            present=True
        )
        db.session.add(new_record)
        db.session.commit()
        return render_template('attendance_success.html', cls=cls, active_session=active_session)

    return redirect(url_for('student_absence'))

# ------------------ UPDATE ATTENDANCE ------------------
@app.route('/teacher/session/<int:session_id>/update_attendance', methods=['POST'])
@role_required(('teacher', 'assistant'))
@permission_required()
def update_attendance(session_id):
    session_obj = AttendanceSession.query.get_or_404(session_id)
    denied = ensure_teacher_session_ownership(session_obj, on_fail='login')
    if denied:
        return denied
    cls = session_obj.class_obj

    raw_present_ids = request.form.getlist('present')
    present_ids = set()
    for raw_id in raw_present_ids:
        raw_id = (raw_id or '').strip()
        if raw_id.isdigit():
            present_ids.add(int(raw_id))

    valid_student_ids = {student.id for student in cls.students}
    present_ids = {student_id for student_id in present_ids if student_id in valid_student_ids}

    for student in cls.students:
        record = AttendanceRecord.query.filter_by(student_id=student.id, session_id=session_obj.id).first()
        if student.id in present_ids:
            if not record:
                db.session.add(AttendanceRecord(student_id=student.id, session_id=session_obj.id, present=True))
            else:
                record.present = True
        else:
            if record:
                db.session.delete(record)

    # If stop is requested, close and confirm the session.
    stop = request.form.get('stop')
    if stop == '1':
        session_obj.active = False
        session_obj.confirmed = True
        flash('Yoklama tamamlandı ve kaydedildi.', 'success')
    else:
        flash('Yoklama güncellendi.', 'success')

    db.session.commit()
    return redirect(url_for('teacher_dashboard'))


# ------------------ STOP ATTENDANCE SESSION ------------------
@app.route('/teacher/session/<int:session_id>/stop')
@role_required(('teacher', 'assistant'))
@permission_required()
def stop_session(session_id):
    session_obj = AttendanceSession.query.get_or_404(session_id)
    denied = ensure_teacher_session_ownership(session_obj, on_fail='login')
    if denied:
        return denied
    cls = session_obj.class_obj

    session_obj.active = False
    session_obj.confirmed = True
    db.session.commit()
    flash('Yoklama oturumu tamamlandı ve kaydedildi.', 'success')
    return redirect(url_for('teacher_dashboard'))


# ------------------ ERROR HANDLERS ------------------
@app.errorhandler(404)
def handle_not_found(error):
    return render_template('404.html'), 404


@app.errorhandler(500)
def handle_internal_error(error):
    # Roll back pending transactions so the app can continue serving requests.
    db.session.rollback()
    app.logger.exception('Unhandled server error: %s', error)
    return render_template('500.html'), 500


# ------------------ RUN ------------------
if __name__ == "__main__":
    debug_mode = os.getenv('FLASK_DEBUG', '0') == '1'
    print("=" * 60)
    print("QR Attendance app is starting...")
    print("Open: http://127.0.0.1:5000/login")
    print("Or: http://localhost:5000/login")
    print(f"Debug mode: {debug_mode}")
    print("=" * 60)
    app.run(debug=debug_mode, host='0.0.0.0', port=5000)




